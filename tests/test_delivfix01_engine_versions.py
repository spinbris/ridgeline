"""DELIVFIX-01 / R-3: no sprint code reaches a client document.

The audit trail's "Engine versions" printed ``SP18-line``, ``SP19 (SR5)``, ``SP21 (SR8)`` and
``SP22 (SR9 v3.5+mean)`` - this repository's sprint history - and the Word document's properties and
visible footer took ``SP22`` as a hard-coded fallback. R-3 (Director, 22/09/2026): the package version
and the engine SHA instead.

Why this is a live render and not an assertion on source: the only earlier tests of the Word document's
properties and footer, ``test_sr9_v32_d6_doc_properties`` (retired by SNAPSHOTCUT-01), **skipped on
every run** ("v31 legacy state-cache pickle not available") - they pinned ``SP22`` as correct and could
never fail. This renders
a real document through the real composer and reads it back, including the ``docProps/`` parts where
the old fallback lived.
"""
import zipfile

import pytest

from src import __version__

SPRINT_CODES = ("SP18", "SP19", "SP21", "SP22")


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    from tests.test_sr9_v3_word_generator import _generate_v3_doc
    return _generate_v3_doc(tmp_path_factory.mktemp("delivfix01_r3"))


def _everything_a_reader_can_see(path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    for s in d.sections:
        parts += [p.text for p in s.footer.paragraphs] + [p.text for p in s.header.paragraphs]
    with zipfile.ZipFile(path) as z:  # core and custom document properties
        parts += [z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.startswith("docProps/")]
    return "\n".join(parts)


def test_the_composed_engine_versions_are_the_package_version(rendered):
    _, composed = rendered
    ev = composed["appendix_audit_trail"]["engine_versions"]
    assert ev == {"package": __version__}, f"the composer's default emits {ev}"


def test_no_sprint_code_reaches_the_word_document(rendered):
    path, _ = rendered
    found = [c for c in SPRINT_CODES if c in _everything_a_reader_can_see(path)]
    assert not found, f"sprint codes reach the client's Word document: {found}"


def test_the_word_document_names_the_package_version(rendered):
    path, _ = rendered
    assert __version__ in _everything_a_reader_can_see(path)
