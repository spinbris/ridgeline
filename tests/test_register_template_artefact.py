"""The shipped register template is what the scaffolder produces. Nothing else.

D-4 (SPEC_product.md §8, ruled 02/09/2026): an empty register workbook ships as a
tracked artefact, `templates/register_template.xlsx`, so that obtaining one does
not require Python. The 21/08 product-path inventory found the directory absent.

WHY A TEST AND NOT JUST THE FILE
    A tracked artefact that nothing regenerates is stale the first time the
    scaffolder changes, and it stays stale silently: the file is present, the
    contract version inside it is old, and a generalist fills in a workbook the
    engine will reject. That is the shape of the SHOWCASE-RENDER-01 defect
    applied to the one artefact the Tier-1 user actually touches.

    So the test regenerates the template through the shipped entry point and
    compares STRUCTURE - sheet names, every defined name and its target, and
    every cell value and formula on every sheet. Not bytes: xlsx bytes carry
    timestamps and zip ordering that differ run to run without meaning anything.

WHY NOT A SKIP IF THE FILE IS MISSING
    It is a tracked artefact. Its absence is a defect, not an environmental
    condition, and a skip would trade coverage for a green lane.

WHY IN-PROCESS AND NOT A SUBPROCESS
    The first draft shelled out to `python tools/register_scaffold.py`, and the
    conftest shell-out guard refused it: per the Director (25/08) a shell-out
    over one second is marked `slow`, never exempted, and this one is ~5s. But
    `slow` would move the drift check to the increment lane, and a stale
    template is exactly what the per-stage gate should catch. The property is
    about CONTENT - tracked artefact equals scaffolder output - not about the
    CLI wrapper, so calling `build_template()` directly tests the property in
    the fast lane at a fraction of the cost. (Contrast test_showcase_smoke.py,
    where the entry point itself is the thing under test.)
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from tools.register_scaffold import build_template

REPO_ROOT = Path(__file__).resolve().parents[1]
TRACKED = REPO_ROOT / "templates" / "register_template.xlsx"


def _structure(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)  # formulas preserved, not data_only
    names = {
        name: sorted(str(dest) for dest in dn.destinations)
        for name, dn in wb.defined_names.items()
    }
    cells = {
        ws.title: {
            cell.coordinate: cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }
        for ws in wb.worksheets
    }
    return {"sheets": wb.sheetnames, "names": names, "cells": cells}


def test_tracked_template_matches_scaffolder_output(tmp_path):
    assert TRACKED.is_file(), (
        "templates/register_template.xlsx is a tracked artefact (D-4) and is missing."
    )
    fresh = tmp_path / "regenerated.xlsx"
    build_template(fresh)
    assert fresh.is_file(), "build_template returned but wrote nothing"

    tracked, regenerated = _structure(TRACKED), _structure(fresh)

    assert tracked["sheets"] == regenerated["sheets"], (
        "sheet list differs from what the scaffolder now emits - regenerate the "
        "tracked template with `python tools/register_scaffold.py` and commit it."
    )
    assert tracked["names"] == regenerated["names"], (
        "defined names differ from the scaffolder's - the tracked template is stale."
    )
    for sheet in regenerated["sheets"]:
        diff = {
            k: (tracked["cells"][sheet].get(k), v)
            for k, v in regenerated["cells"][sheet].items()
            if tracked["cells"][sheet].get(k) != v
        }
        extra = set(tracked["cells"][sheet]) - set(regenerated["cells"][sheet])
        assert not diff and not extra, (
            f"sheet {sheet!r}: {len(diff)} cell(s) differ, {len(extra)} extra in tracked. "
            f"First differences: {dict(list(diff.items())[:5])}. The tracked template is "
            "stale; regenerate and commit."
        )
