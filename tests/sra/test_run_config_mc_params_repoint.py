"""REG-01 §12.4 lock: read_run_config reads config from either layout.

The consolidated showcase register dissolves the strawman and folds
Run_Config_Preview into MC_Params. read_run_config must read the folded
config from MC_Params when Run_Config_Preview is absent, while leaving the
RefProject strawman path (Run_Config_Preview) byte-unchanged.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.sra.strawman_reader import read_run_config


def _run_config_preview_workbook(path: Path) -> None:
    """RefProject-shape strawman: Run_Config_Preview (label col A, value col B)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Run_Config_Preview"
    rows = [
        ("iterations", 50000),
        ("programme_seasons", 1),
        ("day_rate_default ($/wd)", 100000.0),
        ("cp_hybrid_float_threshold_days", 20),
        ("p-level_reporting", 80),
    ]
    for i, (label, value) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
    wb.save(str(path))


def _mc_params_workbook(path: Path) -> None:
    """Consolidated-shape register: MC_Params (label col B, value col C)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Register"  # analyst sheet; irrelevant to read_run_config
    mp = wb.create_sheet("MC_Params")
    rows = [
        ("Iterations (n)", 50000),
        ("Random seed (0=random)", 42),
        ("Reporting P-Level (pct)", 80),
        ("day_rate ($/wd)", 35000.0),
        ("programme_seasons", 1),
        ("cp_hybrid_float_threshold_days", 20),
        ("schedule duration (weeks)", 90),
        ("expected_time_yes_count", 10),
    ]
    for i, (label, value) in enumerate(rows, start=2):
        mp.cell(row=i, column=2, value=label)
        mp.cell(row=i, column=3, value=value)
    wb.save(str(path))


def test_reads_run_config_preview_layout(tmp_path: Path):
    """RefProject path unchanged: values come off Run_Config_Preview."""
    wb = tmp_path / "strawman.xlsx"
    _run_config_preview_workbook(wb)

    cfg = read_run_config(wb)

    assert cfg["iterations"] == 50000
    assert cfg["day_rate_default"] == 100000.0
    assert cfg["p_level_reporting"] == 80
    assert cfg["cp_hybrid_float_threshold_days"] == 20


def test_reads_mc_params_layout(tmp_path: Path):
    """Consolidated path: folded config read off MC_Params."""
    wb = tmp_path / "register.xlsx"
    _mc_params_workbook(wb)

    cfg = read_run_config(wb)

    assert cfg["iterations"] == 50000
    assert cfg["day_rate_default"] == 35000.0
    assert cfg["p_level_reporting"] == 80
    assert cfg["cp_hybrid_float_threshold_days"] == 20
    assert cfg["schedule_duration_weeks"] == 90.0
    assert cfg["expected_time_yes_count"] == 10
    assert cfg["programme_seasons"] == 1


def test_neither_sheet_raises(tmp_path: Path):
    """A workbook with no config sheet fails loud, not silent-defaults."""
    wb = tmp_path / "bare.xlsx"
    b = openpyxl.Workbook()
    b.active.title = "Something"
    b.save(str(wb))

    with pytest.raises(ValueError):
        read_run_config(wb)
