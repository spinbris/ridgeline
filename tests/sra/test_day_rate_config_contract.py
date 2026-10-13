"""REG-01 §12.3 lock: day_rate config contract.

`load_sra_config` must source day_rate from the workbook. A failed or absent
config read RAISES; it MUST NOT silently substitute the RefProject-specific
`FIXTURE_DAY_RATE` for another project. This is the same silent-wrong class as
the Mode-zero and per-season-probability traps.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.sra.sra_config import load_sra_config

from tests._fixture_values import FIXTURE_DAY_RATE


def _write_strawman(path: Path, day_rate: float | None) -> None:
    """Minimal strawman with a Run_Config_Preview sheet.

    If ``day_rate`` is None, omit the day_rate_default row entirely (the
    absent-config case the contract must catch).
    """
    wb = openpyxl.Workbook()
    rc = wb.active
    rc.title = "Run_Config_Preview"
    rows = [("iterations", 5000), ("programme_seasons", 1), ("p-level_reporting", 80)]
    if day_rate is not None:
        rows.append(("day_rate_default ($/wd)", day_rate))
    for i, (label, value) in enumerate(rows, start=1):
        rc.cell(row=i, column=1, value=label)
        rc.cell(row=i, column=2, value=value)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _dummy_inputs(tmp_path: Path) -> tuple[str, str]:
    """register + p6 paths: p6 only needs to be non-None; the register is read (RS-D-3b)."""
    # RS-D-3b (D-15): load_sra_config reads the register's Mode multiplier grid, so the
    # register must be a real workbook that carries one - the wind fixture register.
    reg = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"
    p6 = tmp_path / "p6.xlsx"
    p6.write_bytes(b"")
    return str(reg), str(p6)


def test_day_rate_sourced_from_workbook(tmp_path: Path):
    """Positive control: a workbook day_rate flows through, not the default."""
    reg, p6 = _dummy_inputs(tmp_path)
    strawman = tmp_path / "straw.xlsx"
    _write_strawman(strawman, day_rate=35_000.0)

    cfg = load_sra_config(reg, p6, str(strawman))

    assert cfg.day_rate == 35_000.0
    assert cfg.day_rate != FIXTURE_DAY_RATE


def test_absent_day_rate_raises_not_substitutes(tmp_path: Path):
    """§12.3 core: missing config raises; never returns the code default."""
    reg, p6 = _dummy_inputs(tmp_path)
    strawman = tmp_path / "straw_no_rate.xlsx"
    _write_strawman(strawman, day_rate=None)

    with pytest.raises(ValueError) as exc:
        load_sra_config(reg, p6, str(strawman))

    # DAYRATE-02: this used to assert the refusal QUOTED the day rate, which
    # pinned the wrong behaviour - it printed the client's commercial rate to
    # the console on every failure, and typed it into tracked test source in
    # two spellings to check. The refusal must name the INPUT and its expected
    # source, and must not carry the rate. Same rule the DAYRATE-01 contract
    # module was written under.
    message = str(exc.value)
    assert "day_rate_default" in message
    assert "Run_Config_Preview.Day_rate_default" in message
    assert f"{FIXTURE_DAY_RATE:,.0f}" not in message
    assert f"{FIXTURE_DAY_RATE:.0f}" not in message


def test_unreadable_strawman_raises(tmp_path: Path):
    """A workbook read failure raises rather than falling back."""
    reg, p6 = _dummy_inputs(tmp_path)
    strawman = tmp_path / "corrupt.xlsx"
    strawman.write_bytes(b"not a real xlsx")

    with pytest.raises(ValueError):
        load_sra_config(reg, p6, str(strawman))


def test_missing_strawman_file_raises(tmp_path: Path):
    """A resolved-but-absent strawman cannot silently yield the code default."""
    reg, p6 = _dummy_inputs(tmp_path)
    strawman = tmp_path / "does_not_exist.xlsx"

    with pytest.raises(ValueError):
        load_sra_config(reg, p6, str(strawman))
