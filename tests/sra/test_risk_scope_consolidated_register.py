"""REG-01 §3 lock: read_risk_scope reads the consolidated register.

The showcase dissolves the strawman: schedule three-points + Mode live on
the themed "Risk Register" sheet, CPM binding on the "Mapping" sheet. This
test proves read_risk_scope produces the SAME schedule numbers from the
consolidated register as from an equivalent strawman Risk_Scope (reader-level
bit-identity), and that the Mapping join lands.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.sra.strawman_reader import read_risk_scope

# (rid, time_impact, mode, prob_time, opt_d, ml_d, pess_d, prefix, act_count)
RISKS = [
    ("WT-001", "Yes", "Series", 0.45, 5, 12, 24, "20.A", 11),
    ("WT-002", "Yes", "Hybrid", 0.50, 3, 7, 16, "10.G.001", 1),
    ("WT-003", "No", "", 0.0, 0, 0, 0, "", 1),
]


def _build_strawman(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk_Scope"
    headers = ["Risk_ID", "Time_Impact", "Mode_Aggregated", "Probability", "Probability_Time",
               "Opt_Days", "ML_Days", "Pess_Days", "Pattern_ID", "Prefix_List",
               "Activity_Count"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c, value=h)
    for i, r in enumerate(RISKS):
        rid, ti, mode, pt, od, md, pd_, pfx, ac = r
        vals = [rid, ti, mode, pt if ti == "Yes" else 0.30, pt, od, md, pd_, (f"P-{rid}" if pfx else ""), pfx, ac]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=6 + i, column=c, value=v)
    wb.save(str(path))


def _build_consolidated_register(path: Path) -> None:
    """Risk Register (section row 3, header row 4) + Mapping sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Register"
    # Section labels row 3: cost band at col 14, time band at col 20, derived
    # band at col 26 (so the time section stays cols 20-25).
    ws.cell(row=3, column=14, value="Contingent Risk Cost")
    ws.cell(row=3, column=20, value="Time (Working Days)")
    ws.cell(row=3, column=26, value="Analyst / Derived")
    headers = [
        "Risk ID", "Risk Status", "Risk Description", "Package", "RBS LVL 1",
        "RBS LVL 2", "MC Include", "Distribution", "Correlation Group",
        "Recurrence Type", "Causes", "Consequences", "Existing Controls",
        "Cost Impact", "Probability", "Optimistic ($)", "Most Likely ($)",
        "Pessimistic ($)", "Basis of Estimate",
        "Time Impact", "Probability (Time)", "Optimistic (days)",
        "Most Likely (days)", "Pessimistic (days)", "Basis of Estimate",
        "Mode",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    for i, r in enumerate(RISKS):
        rid, ti, mode, pt, od, md, pd_, pfx, ac = r
        row = 5 + i
        # cost side (cols 14-19) — arbitrary, not under test here
        ws.cell(row=row, column=1, value=rid)
        ws.cell(row=row, column=7, value="Yes")
        ws.cell(row=row, column=14, value="Yes")
        # SPEC R-10: a Time Impact = Yes risk has one probability, so the cost column carries it.
        ws.cell(row=row, column=15, value=pt if ti == "Yes" else 0.30)
        ws.cell(row=row, column=16, value=100000)
        ws.cell(row=row, column=17, value=200000)
        ws.cell(row=row, column=18, value=300000)
        # time side (cols 20-24) + Mode (col 26)
        ws.cell(row=row, column=20, value=ti)
        ws.cell(row=row, column=21, value=pt)
        ws.cell(row=row, column=22, value=od)
        ws.cell(row=row, column=23, value=md)
        ws.cell(row=row, column=24, value=pd_)
        ws.cell(row=row, column=26, value=mode)

    mp = wb.create_sheet("Mapping")
    mp.cell(row=1, column=1, value="Risk ID")
    mp.cell(row=1, column=2, value="Pattern_ID")
    mp.cell(row=1, column=3, value="Prefix_List")
    mp.cell(row=1, column=4, value="Activity_Count")
    # Real builder maps ONLY Time=Yes risks (cost-only risks carry no binding).
    mapped = [r for r in RISKS if r[7]]
    for i, r in enumerate(mapped):
        rid, ti, mode, pt, od, md, pd_, pfx, ac = r
        mp.cell(row=2 + i, column=1, value=rid)
        mp.cell(row=2 + i, column=2, value=f"P-{rid}")
        mp.cell(row=2 + i, column=3, value=pfx)
        mp.cell(row=2 + i, column=4, value=ac)

    # MC_Params so it reads as a V5-style consolidated workbook.
    cfg = wb.create_sheet("MC_Params")
    cfg.cell(row=2, column=2, value="Iterations (n)")
    cfg.cell(row=2, column=3, value=50000)
    wb.save(str(path))


@pytest.fixture
def _wt_prefix(monkeypatch):
    """Force the active binding's risk-id prefix to 'WT' for this read."""
    monkeypatch.setattr(
        "src.sra.cpm.bindings.active_risk_id_prefix", lambda default="R": "WT"
    )


def test_consolidated_matches_strawman_schedule_numbers(tmp_path, _wt_prefix):
    straw = tmp_path / "straw.xlsx"
    reg = tmp_path / "register.xlsx"
    _build_strawman(straw)
    _build_consolidated_register(reg)

    s = read_risk_scope(straw).set_index("risk_id")
    r = read_risk_scope(reg).set_index("risk_id")

    assert list(s.index) == list(r.index) == ["WT-001", "WT-002", "WT-003"]
    # Schedule numbers must match exactly (these drive the bit-identity headline).
    for col in ("time_impact", "mode_aggregated", "probability",
                "opt_days", "ml_days", "pess_days"):
        assert list(s[col]) == list(r[col]), f"column {col} differs"
    # Mapping columns feed the sim only for Time=Yes risks; cost-only risks
    # carry no binding in the consolidated design (None) where the strawman
    # left a default. Compare only the mapped (Time=Yes) rows.
    yes = [rid for rid, ti, *_ in RISKS if ti == "Yes"]
    for col in ("prefix_list", "activity_count"):
        assert list(s.loc[yes, col]) == list(r.loc[yes, col]), f"column {col} differs"


def test_consolidated_mapping_join_lands(tmp_path, _wt_prefix):
    reg = tmp_path / "register.xlsx"
    _build_consolidated_register(reg)
    r = read_risk_scope(reg).set_index("risk_id")
    assert r.loc["WT-001", "prefix_list"] == "20.A"
    assert r.loc["WT-001", "activity_count"] == 11
    assert r.loc["WT-002", "prefix_list"] == "10.G.001"
