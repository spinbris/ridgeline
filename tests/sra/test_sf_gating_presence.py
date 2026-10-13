"""SRA-02 S-F — P-F1/P-F2: presence validation for ``Gating_Activity_ID`` (D-7).

GATINGINERT-01: the column was read, stored on live drivers and consulted by
nothing, so a register could carry ``Gating_Activity_ID = TOTALLY-MADE-UP`` on
every mapped risk and every gate in the tree stayed green. The old validator
checked the gate only under ``injection_method == "discrete_event_additive_days"``
(``src/sra/_pre_flight.py:168-169`), and no live path selects that method.

D-7 ruled the fail-closed limb, and it validates on **presence, not method**: a
non-blank gate must resolve to a ``task_code`` in the ``.xer`` the run will use,
whatever ``injection_method`` says, blank included. The method-scoped limb
(a blank gate *under* ``discrete_event_additive_days`` raises) is unchanged.

These tests drive that limb on all three runtime paths that expand drivers
against an ``.xer`` (A11.1 / P-F2):

1. the RefProject binding's driver expansion - retired with the binding at REFRETIRE-01
   (``1f3db94``); kept in this list because the tests below were written against it.
2. ``scripts/run_wind_showcase.run_stage3_cpm``.
3. ``scripts/run_wind_showcase.run_integrated`` — the consolidated path the AC-5
   pin and the S-E schedule calibration check ride on.

P-F4: the golden input is never touched (D-10 Option 2). The wind sites run
against a COPY of ``tests/fixtures/synthetic_wind/wind_register.xlsx`` carrying
``TOTALLY-MADE-UP`` on one Time=Yes row with ``injection_method`` blank — the
exact shape the register row describes. ``test_golden_input_is_never_written``
proves the copy stayed a copy.

Every negative here has a positive control beside it (G6): the unmodified wind
fixture resolves all ten gates, which is A11.1's D-7 precondition probe promoted
from a one-shot script into a standing test.
"""

from __future__ import annotations

import hashlib
import importlib.util
import shutil
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src.sra._pre_flight import (
    MethodologyValidationError,
    assert_injection_method_consistency,
)

BAD_GATE = "TOTALLY-MADE-UP"
FIXTURE_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic_wind"
WIND_REGISTER = FIXTURE_DIR / "wind_register.xlsx"
WIND_XER = FIXTURE_DIR / "wind_schedule.xer"

# The Mapping sheet carries Gating_Activity_ID (header row 1, ten data rows
# 2-11, one per Time=Yes risk). Row 2 is WT-001, asserted before it is broken so
# a fixture reshuffle cannot leave this test editing a cell that means something
# else.
_MAPPING_SHEET = "Mapping"
_FIRST_DATA_ROW = 2
_FIRST_RISK_ID = "WT-001"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _register_copy_with_bad_gate(tmp_path: Path) -> Path:
    """Copy the wind register; break ONE Time=Yes row's gating id.

    Writes only to ``tmp_path``. ``injection_method`` is left blank, which is
    what makes this the GATINGINERT-01 shape rather than the pre-existing
    discrete-event case.
    """
    dst = tmp_path / "wind_register_bad_gate.xlsx"
    shutil.copyfile(WIND_REGISTER, dst)
    wb = openpyxl.load_workbook(dst)
    ws = wb[_MAPPING_SHEET]
    header = {
        ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)
    }
    assert ws.cell(row=_FIRST_DATA_ROW, column=header["Risk ID"]).value == _FIRST_RISK_ID
    ws.cell(row=_FIRST_DATA_ROW, column=header["Gating_Activity_ID"]).value = BAD_GATE
    wb.save(dst)
    return dst


@pytest.fixture
def _wind(monkeypatch):
    """Activate the synthetic-wind binding and hand back the showcase module.

    Function-scoped so it runs after the conftest refproject autouse select (GD-0)
    and overrides it for this module only; the env is restored on teardown so
    the selection cannot leak into a later test.
    """
    import scripts.run_wind_showcase as wind

    spec = importlib.util.spec_from_file_location(
        "wind_binding_sf", FIXTURE_DIR / "wind_binding.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # self-registers 'synthetic_wind'
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")
    return wind


def _wind_graph_codes():
    from src.sra.cpm.graph_builder import build_schedule_graph
    from src.sra.cpm.xer_reader import parse_xer

    graph = build_schedule_graph(parse_xer(str(WIND_XER)))
    return {a.code for a in graph.activities.values()}


# ---------------------------------------------------------------------------
# P-F2 site 1 — the RefProject binding path
# ---------------------------------------------------------------------------


class TestRefProjectBindingSurface:
    """The binding's validation surface, driven with a RefProject-shaped frame.

    The RefProject register, strawman and Stage-3 ``.xer`` are the confidential
    tier and are gitignored, so the fixture path cannot be exercised from the
    committed suite — A11.2 P-F4 allows a RefProject-shaped scope frame instead,
    and this says so. What is exercised is the exact function the binding calls
    in the retired RefProject binding, with the same argument shape: the live
    schedule graph's activity codes.
    """

    def test_non_resolving_gate_raises_with_blank_method(self) -> None:
        df = pd.DataFrame([
            {
                "risk_id": "R013",
                "probability_time": 0.6,
                "injection_method": "",
                "gating_activity_id": BAD_GATE,
            },
        ])
        with pytest.raises(MethodologyValidationError) as excinfo:
            assert_injection_method_consistency(df, {"A1000", "A2000"})
        message = str(excinfo.value)
        assert "R013" in message, message
        assert BAD_GATE in message, message

    def test_resolving_gate_with_blank_method_passes(self) -> None:
        """Positive control: the limb fires on non-resolution, not on presence."""
        df = pd.DataFrame([
            {
                "risk_id": "R013",
                "probability_time": 0.6,
                "injection_method": "",
                "gating_activity_id": "A1000",
            },
        ])
        assert_injection_method_consistency(df, {"A1000", "A2000"})

    def test_a_later_row_is_not_masked_by_a_clean_first_row(self) -> None:
        df = pd.DataFrame([
            {"risk_id": "R013", "injection_method": "", "gating_activity_id": "A1000"},
            {"risk_id": "R067", "injection_method": "", "gating_activity_id": BAD_GATE},
        ])
        with pytest.raises(MethodologyValidationError, match="R067"):
            assert_injection_method_consistency(df, {"A1000"})


# ---------------------------------------------------------------------------
# P-F2 sites 2 and 3 — both wind expansion paths, on a copied register
# ---------------------------------------------------------------------------


class TestWindRuntimePaths:

    def test_stage3_cpm_raises_before_any_iteration(
        self, _wind, tmp_path, monkeypatch
    ) -> None:
        bad = _register_copy_with_bad_gate(tmp_path)
        monkeypatch.setattr(_wind, "WIND_REGISTER", bad)
        with pytest.raises(MethodologyValidationError) as excinfo:
            _wind.run_stage3_cpm()
        message = str(excinfo.value)
        assert _FIRST_RISK_ID in message, message
        assert BAD_GATE in message, message

    def test_consolidated_path_raises_before_any_iteration(
        self, _wind, tmp_path, monkeypatch
    ) -> None:
        bad = _register_copy_with_bad_gate(tmp_path)
        monkeypatch.setattr(_wind, "WIND_REGISTER", bad)
        with pytest.raises(MethodologyValidationError) as excinfo:
            _wind.run_integrated()
        message = str(excinfo.value)
        assert _FIRST_RISK_ID in message, message
        assert BAD_GATE in message, message

    def test_all_ten_wind_gates_resolve_today(self, _wind) -> None:
        """Positive control, and A11.1's D-7 precondition as a standing test.

        Every non-blank ``Gating_Activity_ID`` in the shipped wind register
        resolves against ``wind_schedule.xer``, so adding the presence limb reds
        nothing on the showcase. If this ever fails, the two negatives above
        stop being evidence of anything.
        """
        from src.sra.strawman_reader import read_risk_scope

        scope = read_risk_scope(str(WIND_REGISTER))
        codes = _wind_graph_codes()
        populated = [
            str(row.gating_activity_id).strip()
            for row in scope.itertuples()
            if str(getattr(row, "gating_activity_id", "") or "").strip()
        ]
        assert len(populated) == 10, populated
        assert set(populated) <= codes, sorted(set(populated) - codes)
        assert_injection_method_consistency(scope, codes)

    def test_golden_input_is_never_written(self, tmp_path) -> None:
        """D-10 Option 2: the red is shown on a copy, never on the golden input."""
        before = _sha256(WIND_REGISTER)
        copied = _register_copy_with_bad_gate(tmp_path)
        assert copied.parent == tmp_path
        assert _sha256(WIND_REGISTER) == before
