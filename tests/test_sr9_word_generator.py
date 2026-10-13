"""SR9 Pass A tests for word_generator."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from src.final import disclaimer_text
from src.final.word_generator import (
    CUSTOM_PROPERTY_NAME,
    FAIL_BANNER_LINES,
    _get_custom_property,
    generate_word_summary,
)


def _make_summary_data():
    return {
        "headline_contingency_range_p80": {
            "lower_bound_dollars": 25_000_000.0,
            "lower_bound_scenario": "DP Review baseline",
            "indicative_dollars": 30_000_000.0,
            "indicative_scenario": "Dummy Realistic (" + "S" + "JP indicative)",
            "upper_bound_dollars": 35_000_000.0,
            "upper_bound_scenario": "All Series",
            "all_three_scenarios": {
                "DP Review baseline": 25_000_000.0,
                "Dummy Realistic (" + "S" + "JP indicative)": 30_000_000.0,
                "All Series": 35_000_000.0,
            },
            "p_level": 80,
        },
        "top_3_cost_drivers": [
            {"risk_id": "R067", "title": "Transformer tail risk",
             "metric_value": 5_000_000.0, "mean_value": 100_000.0,
             "percentage_of_total": 0.40, "units": "dollars"},
            {"risk_id": "R017.1", "title": "Planning Series risk",
             "metric_value": 3_000_000.0, "mean_value": 200_000.0,
             "percentage_of_total": 0.25, "units": "dollars"},
            {"risk_id": "R013", "title": "Site establishment",
             "metric_value": 2_000_000.0, "mean_value": 150_000.0,
             "percentage_of_total": 0.15, "units": "dollars"},
        ],
        "top_3_schedule_drivers": [
            {"risk_id": "R017.1", "title": "Planning",
             "metric_value": 12.0, "mean_value": 3.0,
             "percentage_of_total": 0.4, "units": "days",
             "p_level_dollars": 12.0 * 100000.0,
             "mode": "Series", "basis_scenario": "Dummy Realistic"},
            {"risk_id": "R067", "title": "Transformer tail",
             "metric_value": 8.0, "mean_value": 1.0,
             "percentage_of_total": 0.3, "units": "days",
             "p_level_dollars": 8.0 * 100000.0,
             "mode": "Series", "basis_scenario": "Dummy Realistic"},
            {"risk_id": "R090", "title": "Substation",
             "metric_value": 5.0, "mean_value": 2.0,
             "percentage_of_total": 0.2, "units": "days",
             "p_level_dollars": 5.0 * 100000.0,
             "mode": "Series", "basis_scenario": "Dummy Realistic"},
        ],
        "schedule_basis_note": (
            "Schedule drivers reported under Dummy Realistic Mode classification."
        ),
        "key_findings": [
            "Finding 1.", "Finding 2.", "Finding 3.", "Finding 4.",
        ],
        "methodology_caveats": [
            "PM review pending.", "Longest-path subset.", "LD/IDC deferred.",
        ],
        "recommendations": [
            "Adopt P80.", "Complete PM review.", "Confirm LD parameters.",
        ],
        "diversification_summary": {
            "all_series_dollars": 460_000.0,
            "all_series_pct": 0.016,
            "interpretation": (
                "Joint sampling produces $0.46M (1.6%) diversification benefit "
                "at P80 under All Series."
            ),
        },
    }


def _make_audit_trail():
    return {
        "register_path": "/path/to/register.xlsx",
        "strawman_path": "/path/to/strawman.xlsx",
        "run_date": "2026-04-26",
        "iterations": 30000,
        "random_seed": 42,
        "day_rate_dollars_per_workday": 100000.0,
        "reporting_p_level": 80,
        "engine_versions": {
            "cost": "SP18-line", "schedule": "SP19 (SR5)",
            "integrated": "SP21 (SR8)", "final": "SP22 (SR9)",
        },
    }


def _read_all_text(docx_path: Path) -> str:
    """Concatenate all paragraph and table text from a .docx into one string."""
    doc = Document(str(docx_path))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_generate_word_summary_creates_docx_file(tmp_path):
    out = tmp_path / "exec_summary.docx"
    cv = {"all_pass": True}
    result = generate_word_summary(
        summary_data=_make_summary_data(),
        audit_trail_data=_make_audit_trail(),
        disclaimer_paragraphs=disclaimer_text.get_paragraphs(),
        cross_validation=cv,
        output_path=out,
    )
    assert out.exists()
    assert result["output_path"] == str(out.resolve())
    assert result["size_bytes"] > 0


def test_generate_word_summary_returns_metadata(tmp_path):
    out = tmp_path / "exec_summary.docx"
    cv = {"all_pass": True}
    result = generate_word_summary(
        summary_data=_make_summary_data(),
        audit_trail_data=_make_audit_trail(),
        disclaimer_paragraphs=disclaimer_text.get_paragraphs(),
        cross_validation=cv,
        output_path=out,
    )
    for k in ("output_path", "page_count_estimate", "size_bytes", "fail_banner_present"):
        assert k in result
    assert result["fail_banner_present"] is False


def test_generate_word_summary_uses_role_based_language(tmp_path):
    """Document body must NOT contain 'V8' or 'V7' outside the audit trail."""
    out = tmp_path / "exec_summary.docx"
    cv = {"all_pass": True}
    audit = _make_audit_trail()
    audit["register_path"] = "/path/with/no/version/in/it.xlsx"  # avoid false positive
    audit["strawman_path"] = "/path/with/no/version/in/it.xlsx"
    generate_word_summary(
        summary_data=_make_summary_data(),
        audit_trail_data=audit,
        disclaimer_paragraphs=disclaimer_text.get_paragraphs(),
        cross_validation=cv,
        output_path=out,
    )
    text = _read_all_text(out)
    # Role-based language present
    assert "project risk register" in text.lower() or "Risk register" in text
    # No version pin in body narrative
    assert "V8" not in text
    assert "V7" not in text


def test_generate_word_summary_injects_fail_banner_when_cross_validation_fails(tmp_path):
    """Banner text present, fail_banner_present True, custom property FAILED, banner styled red bold 14pt."""
    out = tmp_path / "exec_summary_fail.docx"
    cv = {"all_pass": False}
    result = generate_word_summary(
        summary_data=_make_summary_data(),
        audit_trail_data=_make_audit_trail(),
        disclaimer_paragraphs=disclaimer_text.get_paragraphs(),
        cross_validation=cv,
        output_path=out,
    )
    assert result["fail_banner_present"] is True
    text = _read_all_text(out)
    for line in FAIL_BANNER_LINES:
        assert line in text

    # Verify banner styling on the first table cell's runs
    doc = Document(str(out))
    assert len(doc.tables) > 0, "Banner table should be the first table in the doc"
    banner_cell = doc.tables[0].cell(0, 0)
    found_styled_run = False
    for para in banner_cell.paragraphs:
        for run in para.runs:
            if run.text.strip() and run.text.strip().upper() in [
                line.upper() for line in FAIL_BANNER_LINES
            ]:
                # Confirm red, bold, 14pt
                assert run.bold is True, f"Banner run '{run.text}' should be bold"
                assert run.font.size is not None
                assert run.font.size.pt == 14, (
                    f"Banner font size should be 14pt, got {run.font.size.pt}"
                )
                assert run.font.color.rgb is not None
                rgb_hex = str(run.font.color.rgb)
                assert rgb_hex.upper() == "C00000", (
                    f"Banner colour should be red C00000, got {rgb_hex}"
                )
                found_styled_run = True
    assert found_styled_run, "No banner run with the expected styling was found"

    # Custom property
    assert _get_custom_property(out, CUSTOM_PROPERTY_NAME) == "FAILED"


def test_generate_word_summary_sets_status_passed_when_cross_validation_passes(tmp_path):
    out = tmp_path / "exec_summary_pass.docx"
    cv = {"all_pass": True}
    result = generate_word_summary(
        summary_data=_make_summary_data(),
        audit_trail_data=_make_audit_trail(),
        disclaimer_paragraphs=disclaimer_text.get_paragraphs(),
        cross_validation=cv,
        output_path=out,
    )
    assert result["fail_banner_present"] is False
    text = _read_all_text(out)
    for line in FAIL_BANNER_LINES:
        assert line not in text
    assert _get_custom_property(out, CUSTOM_PROPERTY_NAME) == "PASSED"
