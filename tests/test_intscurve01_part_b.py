"""INTSCURVE-01 Part B: the base beneath the integrated total, moving no event-risk number.

P8   the base is drawn on its own stream and never through ``IntegratedConfig.layer1_config``:
     every event-risk array is byte-identical with it on or off (Reviewer R1).
P9   total = base + event-risk cost + delay cost; contingency is the P-value less the declared
     base; Layer 1 off renders Part A unchanged.
P10  D-4 (a): a ``duration``-driven Layer 1 factor is left out of the base by name and stated;
     the register is not edited.
"""
from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import numpy as np
import pytest

from src.integrated.base_layer import (
    TIME_DRIVER_TAG,
    deliverable_base_layer,
    exclude_time_factors,
    register_layer1,
)
from src.integrated.integrated_dollars import integrated_dollars_array
from src.final.sections.integrated_total import (
    base_exclusion_sentence,
    total_cost_data,
    total_cost_rows,
)

FX = Path(__file__).resolve().parent / "fixtures"
WIND = FX / "synthetic_wind" / "wind_register.xlsx"
ROAD = FX / "synthetic_road" / "road_register.xlsx"
DAY_RATE = 35_000.0


# ---------------------------------------------------------------------------
# P8 - its own stream
# ---------------------------------------------------------------------------

def test_the_layer1_stream_is_none_of_the_run_s_other_streams():
    from src.rng_streams import (
        LAYER1_STREAM,
        SEVERITY_STREAM,
        fire_generator,
        kernel_generator,
        layer1_generator,
        severity_generator,
    )

    assert LAYER1_STREAM not in (0, SEVERITY_STREAM)
    own = layer1_generator(42).random(64)
    for other in (fire_generator, severity_generator, kernel_generator):
        assert not np.array_equal(own, other(42).random(64)), other.__name__
    assert np.array_equal(own, layer1_generator(42).random(64))


@pytest.mark.slow
def test_the_base_moves_no_event_risk_array(monkeypatch):
    """The positive control (property 8): base on and base off, every event-risk array equal."""
    for k, v in dict(QRA_PROJECT="synthetic_wind", REGISTER_FILE=str(WIND),
                     P6_FILE=str(FX / "synthetic_wind" / "wind_schedule.xer"),
                     STRAWMAN_FILE=str(WIND), SR9_ITERATIONS="1000", SR9_SEED="42").items():
        monkeypatch.setenv(k, v)
    from src.final.deliverable import resolve_inputs, run_engines
    from src.integrated import base_layer

    on = run_engines(resolve_inputs())
    assert on.stage3_beta_result["base_layer"] is not None
    assert on.stage3_config.integrated_config.layer1_config is None
    monkeypatch.setattr(base_layer, "register_layer1", lambda _p: None)
    off = run_engines(resolve_inputs())
    assert off.stage3_beta_result["base_layer"] is None

    s_on, s_off = on.stage3_result, off.stage3_result
    assert np.array_equal(s_on.cost_result["cost_dollars_array"],
                          s_off.cost_result["cost_dollars_array"])
    assert np.array_equal(s_on.schedule_days_array, s_off.schedule_days_array)
    assert np.array_equal(on.cost_results["cost_dollars_array"],
                          off.cost_results["cost_dollars_array"])
    for name in on.integrated_by_scenario:
        assert np.array_equal(on.integrated_by_scenario[name]["cost_dollars_array"],
                              off.integrated_by_scenario[name]["cost_dollars_array"]), name
    assert (on.stage3_beta_result["leave_one_out"]["rows"]
            == off.stage3_beta_result["leave_one_out"]["rows"])


def test_the_deliverable_refuses_a_layer1_config_on_the_cost_arm():
    src = (Path(__file__).resolve().parents[1] / "src" / "final" / "deliverable.py").read_text(
        encoding="utf-8")
    assert 'getattr(config, "layer1_config", None) is not None' in src
    assert "deliverable_base_layer(inputs.register_path, n_iter, seed)" in src


# ---------------------------------------------------------------------------
# P10 - D-4 (a), a logged exclusion
# ---------------------------------------------------------------------------

def test_wind_leaves_out_its_duration_factor_and_road_leaves_out_nothing():
    wind = deliverable_base_layer(WIND, 2_000, 42)
    assert [d["name"] for d in wind["excluded_factors"]] == ["Duration"]
    assert wind["excluded_factors"][0]["mean_effect_dollars"] > 0
    assert deliverable_base_layer(ROAD, 2_000, 42)["excluded_factors"] == []
    # The register itself still carries the factor: the Tier 1 owner view is not touched.
    assert "Duration" in [f.name for f in register_layer1(WIND).factors]


def test_the_excluded_factor_is_really_absent_from_the_draw():
    cfg = register_layer1(WIND)
    kept, dropped = exclude_time_factors(cfg)
    assert all(f.driver.lower() != TIME_DRIVER_TAG for f in kept.factors)
    assert {d["name"] for d in dropped} == {"Duration"}
    off = deliverable_base_layer(WIND, 2_000, 42, delay_cost_in_total=False)
    assert "Duration" in off["factor_mean_mult"] and off["excluded_factors"] == []
    on = deliverable_base_layer(WIND, 2_000, 42)
    assert "Duration" not in on["factor_mean_mult"]


def test_a_correlated_time_factor_leaves_the_correlation_block():
    cfg = register_layer1(WIND)
    timed = replace(cfg, factors=tuple(
        replace(f, driver=TIME_DRIVER_TAG) if f.name == "FX" else f for f in cfg.factors))
    kept, dropped = exclude_time_factors(timed)
    names = {d["name"] for d in dropped}
    assert names == {"Duration", "FX"} and "FX" not in kept.corr_names
    k = len(kept.corr_names)
    assert kept.corr_matrix.shape == (k, k) == (len(cfg.corr_names) - 1,) * 2
    i = cfg.corr_names.index("R_Steel")
    j = kept.corr_names.index("R_Steel")
    assert kept.corr_matrix[j, 0] == cfg.corr_matrix[i, 0]


def test_the_exclusion_is_stated_by_name():
    assert base_exclusion_sentence([]) == ""
    wind = deliverable_base_layer(WIND, 2_000, 42)
    text = base_exclusion_sentence(wind["excluded_factors"])
    assert "Duration" in text and "Commissioning" in text and "register is unchanged" in text


# ---------------------------------------------------------------------------
# P9 - the total, and contingency read off it
# ---------------------------------------------------------------------------

def _arrays(n: int = 2_000, seed: int = 5):
    rng = np.random.default_rng(seed)
    return rng.lognormal(15.0, 0.6, n), np.maximum(rng.normal(40.0, 25.0, n), 0.0)


def test_the_total_is_base_plus_the_headline_array_exactly():
    cost, days = _arrays()
    base = deliverable_base_layer(WIND, cost.size, 42)
    total = total_cost_data(base, cost, days, DAY_RATE)
    assert np.array_equal(total["total_dollars"],
                          base["base_array"] + integrated_dollars_array(cost, days, DAY_RATE))
    for p, v in total["percentiles_dollars"].items():
        assert v == float(np.percentile(total["total_dollars"], p))
        assert total["contingency_dollars"][p] == v - 50_000_000.0
    rows = total_cost_rows(total)
    assert [r["Confidence level"] for r in rows] == ["P50", "P80", "P90"]
    assert not any("recommended" in r["Confidence level"] for r in rows)


def test_layer1_off_or_unpaired_arrays_draw_nothing():
    cost, days = _arrays()
    assert total_cost_data(None, cost, days, DAY_RATE) is None
    short = deliverable_base_layer(WIND, cost.size - 1, 42)
    assert total_cost_data(short, cost, days, DAY_RATE) is None


def test_layer1_off_renders_part_a_unchanged(tmp_path, monkeypatch):
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")
    from src.final import composer, commentary_text
    from tests.test_sr9_v3_word_generator import _generate_v3_doc

    _, composed = _generate_v3_doc(tmp_path)
    p2 = composed["page2_distribution"]
    assert p2["total_cost_s_curve_path"] == "" and p2["total_cost_table"] == []
    assert composed["page7_methodology"]["base_layer_exclusion"] == ""
    assert (commentary_text.L1_SYSTEMIC_CAVEAT_BULLET
            in composed["page7_methodology"]["limitations_current"])
    assert composer._apply_total_cost  # the one seam Part B writes through


def test_the_total_cost_section_renders_in_html_and_word(tmp_path, monkeypatch):
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")
    from docx import Document

    from src.final import commentary_text, composer
    from src.final.charts import make_s_curve
    from src.final.html_blocks import blocks
    from src.final.word_generator import generate_word_summary
    from tests.test_sr9_v3_word_generator import _generate_v3_doc

    _, composed = _generate_v3_doc(tmp_path)
    cost, days = _arrays()
    total = total_cost_data(deliverable_base_layer(WIND, cost.size, 42), cost, days, DAY_RATE)
    png = make_s_curve(distributions={"Total": total["total_dollars"] / 1e6},
                       percentile_markers=[50, 80, 90], title="Total cost",
                       output_path=tmp_path / "total.png")
    composer._apply_total_cost(composed, total, {"total_cost_s_curve_path": str(png)})

    p2 = composed["page2_distribution"]
    assert "illustrative" in p2["total_cost_basis"]
    page = blocks(p2, section="page2_distribution")
    assert str(Path(png)) in {b[1] for b in page if b[0] == "image"}
    assert page[0] == ("text", "#### " + p2["total_cost_heading"]), "total cost leads the page"
    p7 = composed["page7_methodology"]
    assert commentary_text.L1_BASE_DRAWN_CAVEAT_BULLET in p7["limitations_current"]
    assert commentary_text.L1_SYSTEMIC_CAVEAT_BULLET not in p7["limitations_current"]

    out = tmp_path / "with_total.docx"
    generate_word_summary(
        summary_data=composed["executive_summary"],
        audit_trail_data=composed["appendix_audit_trail"],
        disclaimer_paragraphs=composed["disclaimer_paragraphs"],
        cross_validation=composed["cross_validation"], output_path=out, composed=composed)
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Total cost - base, event-risk cost and delay cost" in headings
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Duration" in text and p2["total_cost_basis"] in text
    assert any("Contingency, P-value less base ($M)" in c.text
               for t in doc.tables for c in t.rows[0].cells)


def test_illustrative_base_is_a_declared_bool():
    from src.sra.cpm.bindings import get_binding
    from src.sra.cpm.bindings.registry import ProjectBinding

    assert get_binding("synthetic_wind").illustrative_base is True
    assert get_binding("synthetic_road").illustrative_base is True
    assert ProjectBinding(name="x", project_name="X").illustrative_base is False
    with pytest.raises(TypeError, match="illustrative_base"):
        ProjectBinding(name="x", project_name="X", illustrative_base=1)
