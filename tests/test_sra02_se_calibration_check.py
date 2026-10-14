"""SRA-02 S-E P-E4 and P-E8 - the consolidated calibration check.

P-E4: the check runs with its tolerance taken from config, and its result reaches a
named report surface. P-E8: no run-time path reads a comparator from a cached cell -
the check consumes the Python computation only, so the register's ``Summary`` cells
are audit surface and never an input.
"""
from __future__ import annotations

import pathlib

import ast
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.integrated.integrated_simulation import IntegratedConfig
from src.sra.schedule_cross_check import (
    CALIBRATION_REPORT_KEYS,
    compute_schedule_comparators,
    evaluate_schedule_calibration,
    read_register_day_rate,
    read_register_mode_multipliers,
)
from src.sra.strawman_reader import read_risk_scope
from tools.register_scaffold import CALIBRATION_LABELS, CALIBRATION_SHEET

# RS-B step 2: this module asserts against the reference project's register
# shape, RBS map, chain identifiers or narrative templates, so it selects that
# binding explicitly rather than inheriting the suite default.
pytestmark = pytest.mark.project("fixture_r_prefix")   # REFRETIRE-01: the fixtures here carry R-prefixed risk IDs

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE = REPO_ROOT / "tests" / "fixtures" / "sra02_se" / "se_cross_check_register.xlsx"
SHOWCASE = REPO_ROOT / "scripts" / "run_wind_showcase.py"
#: The injection every tracked register uses, and the one Comparator 1 bounds (SCHEDCAL-01).
DISCRETE = ("discrete_event_additive_days",)


def _comparators(path: Path) -> dict:
    return compute_schedule_comparators(
        read_risk_scope(str(path)),
        read_register_mode_multipliers(path),
        read_register_day_rate(path),
    )


# ---------------------------------------------------------------------------
# 1. The check itself
# ---------------------------------------------------------------------------

def test_tolerance_comes_from_the_register():
    """The brief's wiring point, re-pointed by RS-D-2 (D-14): the tolerance is a register
    value with no dataclass default, so it is read off the wind register here and a
    register change moves this test with it."""
    from src.sra.methodology_defaults_reader import read_calibration_tolerances
    import dataclasses
    assert IntegratedConfig.__dataclass_fields__["schedule_calibration_tolerance_days"].default is dataclasses.MISSING
    tol = read_calibration_tolerances(
        FIXTURE.parent / "wind_register.xlsx" if (FIXTURE.parent / "wind_register.xlsx").exists()
        else pathlib.Path("tests/fixtures/synthetic_wind/wind_register.xlsx"),
        caller_module="tests")["schedule_calibration_tolerance_days"]
    comparators = _comparators(FIXTURE)
    out = evaluate_schedule_calibration(
        comparators["comparator_2_finish_delay_days_pert"], comparators, tol, mode_multipliers_applied=True, injection_methods=DISCRETE)
    assert out["absolute_tolerance_days"] == tol
    assert out["passes_calibration"] is True
    assert out["absolute_gap_days"] == 0.0


def test_a_shifted_comparator_2_fails_the_check():
    """A10.4's P-E4 positive control. An MC mean one tolerance-plus-a-bit away from
    Comparator 2 must FAIL; without this the check could be incapable of failing."""
    tol = 1.0
    comparators = _comparators(FIXTURE)
    c2 = comparators["comparator_2_finish_delay_days_pert"]
    assert evaluate_schedule_calibration(c2 + tol * 0.5, comparators, tol, mode_multipliers_applied=True, injection_methods=DISCRETE)[
        "passes_calibration"] is True
    assert evaluate_schedule_calibration(c2 + tol * 1.5, comparators, tol, mode_multipliers_applied=True, injection_methods=DISCRETE)[
        "passes_calibration"] is False


def test_the_check_reports_rather_than_raises_on_a_failure():
    """It must not gate. A deliverable does not come down over a methodology
    disagreement; it reports the gap and lets a human rule on it."""
    comparators = _comparators(FIXTURE)
    out = evaluate_schedule_calibration(9_999.0, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)
    assert out["passes_calibration"] is False
    assert "FAIL" in out["interpretation"]


def test_result_carries_every_key_the_report_surface_renders():
    comparators = _comparators(FIXTURE)
    out = evaluate_schedule_calibration(1.0, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)
    assert set(CALIBRATION_REPORT_KEYS) <= set(out)


def test_bracket_flag_tracks_where_the_mc_mean_sits():
    """The diagnostic half: Comparator 1 and Comparator 2 should span what a network
    produces, so the flag distinguishes 'network absorbed less float than the Mode
    weighting assumed' from 'something is wrong with the binding'."""
    comparators = _comparators(FIXTURE)
    c1 = comparators["comparator_1_emv_days_pert"]
    c2 = comparators["comparator_2_finish_delay_days_pert"]
    assert c1 > c2, "fixture must have a Parallel or Hybrid row for this to mean anything"
    inside = evaluate_schedule_calibration((c1 + c2) / 2.0, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)
    outside = evaluate_schedule_calibration(c1 * 2.0, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)
    assert inside["mc_mean_within_comparator_bracket"] is True
    assert outside["mc_mean_within_comparator_bracket"] is False


# ---------------------------------------------------------------------------
# 2. Interpretation text carries no project identity
# ---------------------------------------------------------------------------

def test_interpretation_names_no_project_or_risk_id():
    """The strawman-path text at sra_analysis.py:151-158 names a risk ID and a
    project-specific figure; A10.3 forbids reusing it. This one must stay generic
    so it is correct on any register."""
    comparators = _comparators(FIXTURE)
    text = evaluate_schedule_calibration(1.0, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)["interpretation"]
    for banned in ("R090", "RefProject", "refproject", "wind", "Wind", "Tropic", "1.75"):
        assert banned not in text, f"interpretation text leaks {banned!r}: {text}"


# ---------------------------------------------------------------------------
# 3. P-E8 - the run-time path never reads a cached Summary cell
# ---------------------------------------------------------------------------

def test_blanking_the_summary_cells_leaves_the_check_unchanged(tmp_path):
    """A10.4's P-E8 control, run against a workbook that actually HAS cached Summary
    values - so the control is capable of failing. Blank all four and the calibration
    result must be byte-identical: the check consumes the Python computation, and the
    Summary cell is audit surface only (a cached value is only as fresh as the last
    Excel save)."""
    before = evaluate_schedule_calibration(30.0, _comparators(FIXTURE), 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)

    blanked = tmp_path / "blanked.xlsx"
    shutil.copyfile(FIXTURE, blanked)
    wb = load_workbook(blanked)
    try:
        ws = wb[CALIBRATION_SHEET]
        for r in range(1, len(CALIBRATION_LABELS) + 1):
            ws.cell(row=r, column=2, value=None)
    finally:
        wb.save(blanked)
        wb.close()

    # The control must be able to fail: confirm the cells really are blank now.
    check = load_workbook(blanked, data_only=True)
    try:
        cells = [check[CALIBRATION_SHEET].cell(row=r, column=2).value
                 for r in range(1, len(CALIBRATION_LABELS) + 1)]
    finally:
        check.close()
    assert all(v is None for v in cells), f"blanking did not take: {cells}"

    after = evaluate_schedule_calibration(30.0, _comparators(blanked), 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)
    assert after == before


def test_the_consolidated_path_dollarises_off_the_register_not_a_literal():
    """Review finding F-3. ``read_register_day_rate``'s docstring promises both sides
    dollarise off one source; the one run-time call site was passing the module's own
    ``DAY_RATE`` literal instead. They agree today, so no number was wrong - but
    nothing enforced the agreement, and an edit to the register's named range would
    have moved the Summary dollars and not the cross-check's.

    Static, because asserting it dynamically means a 50k-iteration engine run. The
    call site must hand `compute_schedule_comparators` the register reader. Parsed
    with ``ast`` rather than by string-splitting - the arguments contain nested
    calls, so a naive split on the first bracket reads the wrong span.
    """
    tree = ast.parse(SHOWCASE.read_text(encoding="utf-8"))
    calls = [
        node for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Name)
        and node.func.id == "compute_schedule_comparators"
    ]
    assert len(calls) == 1, f"expected one call site, found {len(calls)}"
    day_rate_arg = calls[0].args[2]
    assert isinstance(day_rate_arg, ast.Call), (
        f"day-rate argument is {ast.dump(day_rate_arg)}, not a reader call")
    assert day_rate_arg.func.id == "read_register_day_rate", (
        f"the S-E call site must dollarise off the register, got "
        f"{day_rate_arg.func.id!r}")


def test_the_consolidated_path_never_calls_the_cached_summary_reader():
    """Static half of P-E8, so it holds without running the engine: the showcase must
    not import or call ``get_strawman_calibration_target``. A run-time read of the
    cached cells would be invisible to the test above, which only proves the check
    function is clean."""
    source = SHOWCASE.read_text(encoding="utf-8")
    assert "get_strawman_calibration_target" not in source
    assert "compute_schedule_comparators" in source, (
        "the showcase should compute comparators in Python")


# ---------------------------------------------------------------------------
# 4. SCHEDCAL-01 - judged on the run's own Mode basis
# ---------------------------------------------------------------------------

def test_a_mode_off_run_is_bounded_above_by_comparator_1():
    """Stage 3 runs Mode-off: every risk injects its full delay, and the network absorbs some of it
    as float. Its mean is Comparator 1 less that float, so Comparator 1 is an upper bound and the
    check is one-sided (Gate 3, Director 14/10/2026). A mean at Comparator 1 passes; a mean at
    Comparator 2, a network that absorbed float, passes too (the two-sided first cut failed it and
    called the network suspect); a mean above Comparator 1 by more than the tolerance fails."""
    comparators = _comparators(FIXTURE)
    c1 = comparators["comparator_1_emv_days_pert"]
    c2 = comparators["comparator_2_finish_delay_days_pert"]
    assert c1 - c2 > 1.0, "the fixture must separate the comparators by more than the tolerance"
    judge = lambda mean: evaluate_schedule_calibration(  # noqa: E731
        mean, comparators, 1.0, mode_multipliers_applied=False, injection_methods=DISCRETE)
    assert judge(c1)["passes_calibration"] is True
    assert judge(c1)["judged_against"] == "Comparator 1 (upper bound)"
    assert judge(c2)["passes_calibration"] is True
    assert judge(c2 - 5.0)["passes_calibration"] is True
    assert judge(c1 + 0.5)["passes_calibration"] is True
    over = judge(c1 + 1.5)
    assert over["passes_calibration"] is False
    assert "suspect" in over["interpretation"]
    # And the Mode-on rule is unchanged: two-sided about Comparator 2.
    assert evaluate_schedule_calibration(c2, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)[
        "passes_calibration"] is True
    assert evaluate_schedule_calibration(c1, comparators, 1.0, mode_multipliers_applied=True, injection_methods=DISCRETE)[
        "passes_calibration"] is False


def test_the_delay_not_reaching_the_finish_is_always_reported_and_never_decides():
    comparators = _comparators(FIXTURE)
    c1 = comparators["comparator_1_emv_days_pert"]
    c2 = comparators["comparator_2_finish_delay_days_pert"]
    for mean in (c1, c2, (c1 + c2) / 2.0, c2 - 3.0):
        out = evaluate_schedule_calibration(mean, comparators, 1.0, mode_multipliers_applied=False, injection_methods=DISCRETE)
        assert out["passes_calibration"] is True
        assert out["delay_not_reaching_finish_days"] == max(c1 - mean, 0.0)
        assert out["diagnostic_comparator"] == "Comparator 2"
        assert out["diagnostic_gap_days"] == abs(mean - c2)
        assert "Diagnostic, not part of the check" in out["interpretation"]
        assert "suspect" not in out["interpretation"]


def test_the_basis_must_be_stated():
    comparators = _comparators(FIXTURE)
    with pytest.raises(TypeError):
        evaluate_schedule_calibration(1.0, comparators, 1.0)
    with pytest.raises(TypeError, match="own bool setting"):
        evaluate_schedule_calibration(1.0, comparators, 1.0, mode_multipliers_applied=None,
                                      injection_methods=DISCRETE)
    with pytest.raises(TypeError):
        evaluate_schedule_calibration(1.0, comparators, 1.0, mode_multipliers_applied=False)
    for bad in ("discrete_event_additive_days", (), [""]):
        with pytest.raises(TypeError, match="injection"):
            evaluate_schedule_calibration(1.0, comparators, 1.0, mode_multipliers_applied=False,
                                          injection_methods=bad)


@pytest.mark.parametrize("mean", [None, float("nan"), float("inf")])
def test_a_missing_or_non_finite_mean_is_refused_not_read_as_zero(mean):
    """Gate 3 round 2: read as 0.0, a missing mean passed the one-sided rule as total absorption."""
    with pytest.raises(ValueError, match="broken run"):
        evaluate_schedule_calibration(mean, _comparators(FIXTURE), 1.0,
                                      mode_multipliers_applied=False, injection_methods=DISCRETE)


@pytest.mark.parametrize("methods", [
    ("additive_pert_days",),
    ("productivity_multiplicative_factor",),
    ("discrete_event_additive_days", "additive_pert_days"),
])
def test_comparator_1_bounds_only_discrete_injection_so_other_methods_are_not_assessed(methods):
    """Gate 3 round 2 (Director, 14/10/2026, option (a)): additive_pert_days writes one delay to
    every activity a risk binds, so a serial cohort moves the finish by a multiple of it, and a
    multiplicative factor is not days; Comparator 1 bounds neither. Such a run is reported, never
    passed or failed, even at a mean far above Comparator 1."""
    comparators = _comparators(FIXTURE)
    c1 = comparators["comparator_1_emv_days_pert"]
    out = evaluate_schedule_calibration(c1 * 10.0, comparators, 1.0,
                                        mode_multipliers_applied=False, injection_methods=methods)
    assert out["passes_calibration"] is None and out["assessed"] is False
    assert out["judged_against"] == "not assessed"
    assert "NOT ASSESSED" in out["interpretation"] and "suspect" not in out["interpretation"]
    # The same mean, discrete injection: judged, and failed.
    judged = evaluate_schedule_calibration(c1 * 10.0, comparators, 1.0,
                                           mode_multipliers_applied=False, injection_methods=DISCRETE)
    assert judged["assessed"] is True and judged["passes_calibration"] is False


def test_the_bound_is_inclusive_at_comparator_1_plus_the_tolerance():
    comparators = _comparators(FIXTURE)
    c1 = comparators["comparator_1_emv_days_pert"]
    tol = (c1 + 1.0) - c1  # exactly the excess a mean of c1 + 1.0 carries, in floating point
    judge = lambda mean: evaluate_schedule_calibration(  # noqa: E731
        mean, comparators, tol, mode_multipliers_applied=False, injection_methods=DISCRETE)
    assert judge(c1 + 1.0)["passes_calibration"] is True
    assert judge(c1 + 1.0 + 1e-6)["passes_calibration"] is False


def test_the_showcase_passes_its_runs_own_mode_setting():
    """Static: the one caller hands the check the setting its Stage 3 run used, not a literal."""
    tree = ast.parse(SHOWCASE.read_text(encoding="utf-8"))
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and isinstance(n.func, ast.Name) and n.func.id == "evaluate_schedule_calibration"]
    assert len(calls) == 1
    kw = {k.arg: k.value for k in calls[0].keywords}
    assert "mode_multipliers_applied" in kw
    # Exactly the setting the Stage 3 config ran with: an inverted, defaulted or wrong-field
    # expression passed a looser "names apply_mode_multipliers" check (Gate 3, 14/10/2026).
    assert ast.unparse(kw["mode_multipliers_applied"]) == (
        "bool(s3_cfg.schedule_mc_config.apply_mode_multipliers)"), (
        "the showcase must pass its Stage 3 run's own schedule_mc_config setting")
    assert ast.unparse(kw["injection_methods"]) == (
        "{d.injection_method for d in drivers.values()}"), (
        "the showcase must pass the injection methods of the drivers its engine ran")
