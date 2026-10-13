"""REPORT01 F-B part 2: the model-validation section reads what it claims to read.

The section exists to show the sampler did what the register asked. Its failure modes
are all quiet ones - claiming the arms fired together on a run where they did not,
measuring the sampled magnitude over all iterations instead of the fired ones, or
reporting a never-fired risk as a zero-magnitude one. Every test below targets one of
those rather than the shape of the output.

METHFIRE-01 (SPEC R-10) replaced the two firing tables with one: a risk has one
probability and both arms fire on one draw, so the schedule record is checked against the
cost record rather than tabled against a second column.
"""
from __future__ import annotations

from types import SimpleNamespace

import numpy as np
import pytest

from src.final.sections.model_validation import (
    NEVER_FIRED,
    build_model_validation_section,
    elicited_mean,
)

# Asymmetric probabilities, so a table that swapped risks is visible in the numbers.
COST_P = {"R1": 0.20, "R2": 0.80}
N = 10_000


def _risk(rid, *, opt=1e6, ml=2e6, pess=6e6, distribution="Triangular"):
    # Triangular, not PERT: an all-PERT fixture is what let the PERT-for-all defect
    # through every test in this file until the numerics panel found it.
    return SimpleNamespace(
        risk_id=rid, probability=COST_P[rid], distribution=distribution,
        opt_cost=opt, ml_cost=ml, pess_cost=pess,
    )


def _mask(p, seed):
    return np.random.default_rng(seed).random(N) < p


def _result(**over):
    ids = ["R1", "R2"]
    fired = np.stack([_mask(COST_P[r], i) for i, r in enumerate(ids)], axis=1)
    per_risk = np.zeros((N, 2))
    for i, r in enumerate(ids):
        # Magnitude 3.0M exactly, on fired iterations only.
        per_risk[fired[:, i], i] = 3.0e6
    base = {
        "fired_array": fired,
        "risk_ids_all": ids,
        # R-10: the schedule arm fired on the cost arm's draw. Reversed column order, so
        # a check that compared by position rather than by risk would fail on a good run.
        "schedule_fired_array": fired[:, ::-1].copy(),
        "schedule_fired_risk_ids": ids[::-1],
        "per_risk_cost_array": per_risk,
    }
    base.update(over)
    return base


def _section(**over):
    return build_model_validation_section(
        _result(**over), [_risk("R1"), _risk("R2")], {"R1": "First", "R2": "Second"},
    )["page_model_validation"]


# ---------------------------------------------------------------------------
# Numerics review BLOCKER 2, 12/09/2026: the elicited mean assumed PERT for all
# ---------------------------------------------------------------------------

def test_the_elicited_mean_follows_the_risk_s_own_distribution():
    """The engine dispatches on the Distribution column; this table must too.

    Both shipped registers carry Lognormal and Triangular risks alongside PERT, so
    applying the PERT formula to all of them printed a phantom sampler deviation in the
    table built to detect a real one. Same three-point input, four different means.
    """
    o, m, p = 1e6, 2e6, 6e6
    assert elicited_mean(o, m, p, "PERT") == pytest.approx((o + 4 * m + p) / 6, rel=1e-6)
    assert elicited_mean(o, m, p, "Triangular") == pytest.approx((o + m + p) / 3, rel=1e-6)
    assert elicited_mean(o, m, p, "Binary") == pytest.approx(m, rel=1e-6)
    # Triangular and PERT differ by 0.5M on this input: a PERT-for-all implementation
    # reports the triangular risk as sampling 20% high when it is sampling correctly.
    assert elicited_mean(o, m, p, "Triangular") != pytest.approx(
        elicited_mean(o, m, p, "PERT"), rel=1e-3)


def test_an_unknown_distribution_follows_the_engine_s_own_default():
    """`sample_cost` defaults to triangular; the elicited column must agree with it."""
    o, m, p = 1e6, 2e6, 6e6
    assert elicited_mean(o, m, p, "no-such-distribution") == pytest.approx(
        elicited_mean(o, m, p, "Triangular"), rel=1e-9)


def test_the_elicited_mean_is_deterministic_and_seed_free():
    """Quadrature, not sampling. Two calls must agree to the last bit."""
    a = elicited_mean(1e6, 2e6, 6e6, "Lognormal")
    b = elicited_mean(1e6, 2e6, 6e6, "Lognormal")
    assert a.hex() == b.hex()


def test_the_firing_table_is_compared_against_probability():
    rows = {r["Risk"]: r for r in _section()["firing_table"]}
    assert rows["R1"]["Elicited probability"] == "20.0%"
    assert rows["R2"]["Elicited probability"] == "80.0%"
    # Seeded masks at n=10,000: within a percentage point of their own probability.
    assert abs(float(rows["R1"]["Realised firing"].rstrip("%")) - 20.0) < 1.0


def test_there_is_one_firing_table():
    """R-10: one probability, one draw, one table. A second table is a second claim."""
    sec = _section()
    assert "schedule_firing_table" not in sec and "cost_firing_table" not in sec


def test_a_schedule_record_matching_the_cost_record_is_confirmed_by_risk_not_position():
    assert _section()["schedule_fires_match_cost"] is True


def test_a_schedule_record_that_fired_on_its_own_draw_is_caught():
    """Red first for the check: the pre-METHFIRE-01 shape, a second independent draw."""
    own = np.stack([_mask(p, 10 + i) for i, p in enumerate((0.20, 0.80))], axis=1)
    assert _section(schedule_fired_array=own, schedule_fired_risk_ids=["R1", "R2"])[
        "schedule_fires_match_cost"] is False


def test_a_schedule_driver_with_no_cost_record_is_a_mismatch_not_a_skip():
    base = _result()
    sched = np.concatenate([base["fired_array"], base["fired_array"][:, :1]], axis=1)
    assert _section(schedule_fired_array=sched, schedule_fired_risk_ids=["R1", "R2", "R9"])[
        "schedule_fires_match_cost"] is False


def test_the_difference_column_is_signed():
    """A reader needs to see which way a risk missed, not just by how much."""
    for row in _section()["firing_table"]:
        assert row["Difference"].startswith(("+", "-")), row
        assert row["Difference"].endswith(" pp")


def test_sampled_magnitude_is_conditioned_on_firing():
    """The whole point of the magnitude check.

    Every fired iteration carries exactly 3.0M, so the mean over fired iterations is
    3.000. Averaging over all iterations instead would give 0.6M for R1 (p=0.2) - the
    elicited comparison would then be wrong by a factor of five and look like a badly
    calibrated input rather than a bug in the report.
    """
    rows = {r["Risk"]: r for r in _section()["cost_magnitude_table"]}
    assert rows["R1"]["Sampled mean when fired ($M)"] == "3.000"
    assert rows["R2"]["Sampled mean when fired ($M)"] == "3.000"


def test_elicited_magnitude_uses_the_declared_distribution_and_names_it():
    rows = {r["Risk"]: r for r in _section()["cost_magnitude_table"]}
    # The fixture declares Triangular: (1 + 2 + 6) / 3 = 3.0M, not PERT's 2.5M.
    assert rows["R1"]["Elicited mean ($M)"] == "3.000"
    assert rows["R1"]["Distribution"] == "Triangular"


def test_two_risks_with_one_input_and_different_distributions_read_differently():
    """The whole point of naming the distribution on the row."""
    section = build_model_validation_section(
        _result(), [_risk("R1"), _risk("R2", distribution="PERT")], None,
    )["page_model_validation"]
    rows = {r["Risk"]: r for r in section["cost_magnitude_table"]}
    assert rows["R1"]["Elicited mean ($M)"] == "3.000"
    assert rows["R2"]["Elicited mean ($M)"] == "2.500"


def test_a_risk_that_never_fired_says_so_rather_than_reporting_zero():
    """Zero magnitude and no observation are different claims.

    A never-fired risk has nothing to compare; printing 0.000 would tell a reader the
    sampler drew a zero impact, which it did not.
    """
    ids = ["R1", "R2"]
    fired = np.zeros((N, 2), dtype=bool)
    fired[:, 1] = _mask(COST_P["R2"], 1)          # R1 never fires
    per_risk = np.zeros((N, 2))
    per_risk[fired[:, 1], 1] = 3.0e6
    sec = _section(fired_array=fired, per_risk_cost_array=per_risk)
    rows = {r["Risk"]: r for r in sec["cost_magnitude_table"]}
    assert rows["R1"]["Sampled mean when fired ($M)"] == NEVER_FIRED
    assert rows["R1"]["Difference"] == NEVER_FIRED
    assert rows["R2"]["Sampled mean when fired ($M)"] == "3.000"
    # It still appears in the firing table, at 0.0%.
    firing = {r["Risk"]: r for r in sec["firing_table"]}
    assert firing["R1"]["Realised firing"] == "0.0%"


def test_the_post_chain_caveat_and_the_one_draw_note_are_on_the_page():
    """The table misleads without them, so their absence is a defect."""
    sec = _section()
    assert "after chain logic" in sec["firing_basis_note"]
    assert "same firing decisions" in sec["one_fire_draw_note"]
    assert "Probability (Time)" not in sec["one_fire_draw_note"]
    assert "the register declares for that risk" in sec["magnitude_note"], (
        "the note must not assert one formula for every distribution")


def test_iteration_count_is_reported():
    """A firing rate without its sample size is not evidence."""
    assert _section()["n_iterations"] == N


def test_absent_arrays_yield_empty_tables_rather_than_raising():
    """A run that somehow carries no firing record must not take the report down.

    D-11 writes the artefact and fails afterwards; a section builder that raised here
    would abort composition before anything reached disk.
    """
    sec = build_model_validation_section({}, [_risk("R1")], None)["page_model_validation"]
    assert sec["firing_table"] == []
    assert sec["schedule_fires_match_cost"] is None
    assert sec["one_fire_draw_note"] == "", "an unchecked run must not assert one draw"
    assert sec["cost_magnitude_table"] == []
    assert sec["n_iterations"] == 0


# ---------------------------------------------------------------------------
# The render side. A section that is built and never rendered is not delivered.
# ---------------------------------------------------------------------------

def _doc_text(composed):
    from docx import Document

    from src.final.word_generator import _add_model_validation_section

    doc = Document()
    _add_model_validation_section(doc, composed)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_the_section_reaches_the_rendered_document():
    """Headings, both caveats, and real numbers."""
    # A PERT risk, so the elicited mean (2.500) and the sampled mean (3.000) are
    # different numbers and each can be seen to have reached the page on its own.
    section = build_model_validation_section(
        _result(), [_risk("R1", distribution="PERT"), _risk("R2", distribution="PERT")],
        {"R1": "First", "R2": "Second"},
    )["page_model_validation"]
    text = _doc_text({"page_model_validation": section})
    assert "Model validation" in text
    assert "Realised firing against Probability" in text
    assert "Probability (Time)" not in text
    assert "same firing decisions" in text
    assert "WARNING" not in text
    assert "after chain logic" in text
    assert "Sampled magnitude against the elicited estimate" in text
    # The tables' own content, not just their headings.
    assert "20.0%" in text and "80.0%" in text   # elicited probability, both risks
    assert "2.500" in text                        # elicited mean, PERT
    assert "3.000" in text                        # sampled mean when fired
    assert "PERT" in text                         # the distribution is named on the row


def test_a_run_whose_arms_did_not_fire_together_says_so_on_the_page():
    """The one-draw note must not stand unqualified over a run it is untrue of."""
    own = np.stack([_mask(p, 10 + i) for i, p in enumerate((0.20, 0.80))], axis=1)
    section = build_model_validation_section(
        _result(schedule_fired_array=own, schedule_fired_risk_ids=["R1", "R2"]),
        [_risk("R1"), _risk("R2")], None)["page_model_validation"]
    assert "WARNING" in _doc_text({"page_model_validation": section})


def test_a_run_with_no_schedule_record_does_not_assert_one_draw():
    """The note is a claim the section checks. Where there is nothing to check it is not made."""
    base = _result()
    del base["schedule_fired_array"], base["schedule_fired_risk_ids"]
    section = build_model_validation_section(base, [_risk("R1"), _risk("R2")], None)["page_model_validation"]
    assert section["schedule_fires_match_cost"] is None
    assert section["one_fire_draw_note"] == ""
    assert "same firing decisions" not in _doc_text({"page_model_validation": section})


def test_an_executive_summary_without_the_key_renders_nothing():
    """D-2: the executive summary stays an executive summary.

    The Stage 2.5 composer never sets this key. Rendering must be a no-op rather than
    an empty heading, or every exec summary grows a stub section.
    """
    assert _doc_text({}).strip() == ""
    assert _doc_text({"page_model_validation": {}}).strip() == ""


def test_empty_tables_do_not_produce_an_empty_heading():
    """A heading over nothing tells a reader the run had no risks. It did not."""
    empty = {"n_iterations": 0, "firing_table": [], "cost_magnitude_table": []}
    assert _doc_text({"page_model_validation": empty}).strip() == ""


def test_the_renderer_is_actually_called_by_the_document_builder():
    """Built, rendered, and *reached*.

    Every other test here can pass while the section is never invoked during a real
    render - which is precisely how F-G's diagnostic came to be a working script no
    caller ran. A source check is narrow, but it catches the regression that matters:
    the call being dropped.
    """
    import inspect

    from src.final.word_generator import _render_v3_executive_summary

    src = inspect.getsource(_render_v3_executive_summary)
    assert "_add_model_validation_section(doc, composed)" in src, (
        "the document builder no longer calls the model-validation renderer")


def test_the_elicited_mean_is_accurate_on_an_unbounded_heavy_tail():
    """Numerics re-review WARN, 12/09/2026: the fixed grid truncated the lognormal tail.

    Lognormal is the shape reserved for catastrophic-tail risks, and it is unbounded, so
    a uniform midpoint grid loses mass past its last midpoint. Measured against a
    2,000,000-point reference the old 10,000-point grid was 0.51% low at a 100:1
    pessimistic-to-optimistic ratio and 2.64% low at 500:1 - printed as a phantom
    positive Difference in the table built to detect a real one.

    The reference here is deliberately computed the OLD way at very high N: if the
    adaptive result ever drifts from a brute-force integral of the same sampler, one of
    them is wrong and this says so.
    """
    import numpy as np

    from src.distributions import sample_cost

    def brute_force(o, m, p, dist, n=2_000_000):
        u = (np.arange(n, dtype=float) + 0.5) / n
        return float(np.asarray(sample_cost(u, o, m, p, dist)).mean())

    for ratio, tol_pct in ((6, 0.05), (100, 0.10), (500, 0.50)):
        o, m, p = 1e6, 2e6, 1e6 * ratio
        got, ref = elicited_mean(o, m, p, "Lognormal"), brute_force(o, m, p, "Lognormal")
        err = abs(got - ref) / ref * 100
        assert err < tol_pct, (
            f"lognormal {ratio}:1 - adaptive {got/1e6:.4f}M vs reference "
            f"{ref/1e6:.4f}M, {err:.3f}% off (tolerance {tol_pct}%)")


def test_bounded_distributions_stay_exact():
    """The fix must not trade the easy cases away for the hard one."""
    o, m, p = 1e6, 2e6, 100e6
    assert elicited_mean(o, m, p, "Triangular") == pytest.approx((o + m + p) / 3, rel=1e-5)
    assert elicited_mean(o, m, p, "PERT") == pytest.approx((o + 4 * m + p) / 6, rel=1e-5)
