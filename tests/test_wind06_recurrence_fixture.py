"""WIND-06 lock: WT-002 is carried as ``Single`` in the wind showcase fixture.

The integrated engine (which produces the wind showcase headline) has no
recurrence handling; the standalone ``MonteCarloEngine`` honours it. WT-002 was
recoded ``BINOMIAL`` -> ``Single`` (WIND-06) so the showcase is internally
coherent with the integrated engine. This locks that accommodation against a
silent revert on fixture regeneration. The engine capability gap is COST-01.

The falsifiable WIND-06 experiment (flipping WT-002 leaves every integrated
figure bit-identical -> the integrated engine ignores ``Recurrence Type``) was
run at 100k as a build-time assertion and recorded in the WIND-06 stage_log
entry; that run is not repeated here (it is a slow showcase run).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

_FIXTURE = (
    Path(__file__).resolve().parents[1]
    / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"
)
_HEADER_ROW = 4


def _header_cols(ws) -> dict[str, int]:
    return {
        ws.cell(_HEADER_ROW, c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(_HEADER_ROW, c).value
    }


def test_wt002_recurrence_type_is_single():
    wb = openpyxl.load_workbook(_FIXTURE, data_only=False)
    ws = wb["Risk Register"]
    cols = _header_cols(ws)
    id_col = cols["Risk ID"]
    rec_col = cols["Recurrence Type"]
    wt002_rows = [
        r
        for r in range(_HEADER_ROW + 1, ws.max_row + 1)
        if ws.cell(r, id_col).value == "WT-002"
    ]
    assert len(wt002_rows) == 1, f"expected exactly one WT-002 row, got {wt002_rows}"
    assert ws.cell(wt002_rows[0], rec_col).value == "Single"


def test_no_binomial_rows_in_wind_register():
    """The showcase register must carry no recurring rows (integrated-engine coherent)."""
    wb = openpyxl.load_workbook(_FIXTURE, data_only=False)
    ws = wb["Risk Register"]
    cols = _header_cols(ws)
    rec_col = cols["Recurrence Type"]
    values = {
        ws.cell(r, rec_col).value
        for r in range(_HEADER_ROW + 1, ws.max_row + 1)
        if ws.cell(r, cols["Risk ID"]).value
    }
    assert "BINOMIAL" not in values
    assert "POISSON" not in values
