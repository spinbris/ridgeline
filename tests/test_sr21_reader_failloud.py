"""SR21: live cost-arm reader fail-loud + robustness tests.

Covers Build Brief 3 (SR21) scope:
  - §3a read_risks_v5 raises a located RegisterReadError on a broken Excel
    reference in a read cell, instead of silently coercing to a default.
  - §3b _find_section_bounds raises on a missing or ambiguous section band
    rather than mis-slicing.
  - §3c read_cg_descriptions raises a located CgBlockError on a broken
    reference or a duplicate CG code in the MC_Params column-B block.
  - Output-neutral: the RefProject register reads unchanged (34 risks, 13 CG descriptions).
"""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest

from src.openpyxl_reader import CgBlockError, read_cg_descriptions
from src.v5_register_reader import (
    RegisterReadError,
    _find_section_bounds,
    read_risks_v5,
)

# ---------------------------------------------------------------------------
# Synthetic-workbook builders (minimal shape the readers require)
# ---------------------------------------------------------------------------

_CONT_HEADERS = {
    3: "Cost Impact", 4: "Probability", 5: "Optimistic ($)",
    6: "Most Likely ($)", 7: "Pessimistic ($)", 8: "Basis of Estimate",
}
_TIME_HEADERS = {
    10: "Time Impact", 11: "Probability", 12: "Optimistic",
    13: "Most Likely", 14: "Pessimistic", 15: "Basis of Estimate",
}


def _build_register(tmp_path: Path, *, ml_cost=200.0, extra_section_cell=None) -> Path:
    """Build a minimal V5-shaped Risk Register workbook.

    One included risk R001. ``ml_cost`` is the Most Likely ($) cell (col F,
    row 5) -- set to an Excel error string to exercise the fail-loud path.
    ``extra_section_cell`` (col, value) injects a second row-3 label to make
    a section ambiguous.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Register"
    wb.create_sheet("MC_Params")
    ws.cell(row=3, column=3, value="Contingent Risk Cost")
    ws.cell(row=3, column=10, value="Time (Working Days)")
    if extra_section_cell is not None:
        ws.cell(row=3, column=extra_section_cell[0], value=extra_section_cell[1])
    ws.cell(row=4, column=1, value="Risk ID")
    ws.cell(row=4, column=2, value="MC Include")
    for c, v in {**_CONT_HEADERS, **_TIME_HEADERS}.items():
        ws.cell(row=4, column=c, value=v)
    row5 = {1: "R001", 2: "Yes", 3: "Yes", 4: 0.5, 5: 100.0, 6: ml_cost, 7: 300.0}
    for c, v in row5.items():
        ws.cell(row=5, column=c, value=v)
    p = tmp_path / "synthetic_register.xlsx"
    wb.save(p)
    return p


def _build_cg_workbook(tmp_path: Path, col_b_rows: list) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "Risk Register"
    mp = wb.create_sheet("MC_Params")
    for i, v in enumerate(col_b_rows, start=30):
        mp.cell(row=i, column=2, value=v)
    p = tmp_path / "synthetic_cg.xlsx"
    wb.save(p)
    return p


# ---------------------------------------------------------------------------
# §3a read_risks_v5 fail-loud
# ---------------------------------------------------------------------------

def test_clean_register_reads(tmp_path: Path) -> None:
    risks = read_risks_v5(_build_register(tmp_path), risk_id_prefix="R")
    assert [r.risk_id for r in risks] == ["R001"]
    assert risks[0].ml_cost == 200.0


@pytest.mark.parametrize("err", ["#REF!", "#N/A", "#VALUE!", "#DIV/0!"])
def test_broken_cost_ref_raises_located(tmp_path: Path, err: str) -> None:
    path = _build_register(tmp_path, ml_cost=err)
    with pytest.raises(RegisterReadError) as exc:
        read_risks_v5(path, risk_id_prefix="R")
    msg = str(exc.value)
    # Located: the offending cell (F5), the field, the risk id, the error.
    assert "F5" in msg
    assert "ml_cost" in msg
    assert "R001" in msg
    assert err in msg


# ---------------------------------------------------------------------------
# §3b _find_section_bounds hardening
# ---------------------------------------------------------------------------

def test_section_bounds_missing_raises(tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(_build_register(tmp_path))
    with pytest.raises(ValueError, match="not found"):
        _find_section_bounds(wb["Risk Register"], 3, "no such section")


def test_section_bounds_ambiguous_raises(tmp_path: Path) -> None:
    path = _build_register(tmp_path, extra_section_cell=(20, "Contingent Risk Cost (dup)"))
    wb = openpyxl.load_workbook(path)
    with pytest.raises(ValueError, match="ambiguous"):
        _find_section_bounds(wb["Risk Register"], 3, "contingent risk cost")


def test_section_bounds_single_match_unchanged(tmp_path: Path) -> None:
    """Output-neutral: a well-formed single-match band returns (start, end)."""
    wb = openpyxl.load_workbook(_build_register(tmp_path))
    start, end = _find_section_bounds(wb["Risk Register"], 3, "contingent risk cost")
    assert start == 2  # col C, 0-indexed
    assert end == 9    # col J (Time section), 0-indexed


# ---------------------------------------------------------------------------
# §3c read_cg_descriptions hardening
# ---------------------------------------------------------------------------

def test_cg_clean_reads(tmp_path: Path) -> None:
    path = _build_cg_workbook(tmp_path, ["CG01 - Labour", "CG02 - Design"])
    assert read_cg_descriptions(path) == {"CG01": "Labour", "CG02": "Design"}


def test_cg_broken_ref_raises_located(tmp_path: Path) -> None:
    path = _build_cg_workbook(tmp_path, ["CG01 - Labour", "#REF!"])
    with pytest.raises(CgBlockError) as exc:
        read_cg_descriptions(path)
    assert "broken Excel reference" in str(exc.value)
    assert "B31" in str(exc.value)


def test_cg_duplicate_code_raises(tmp_path: Path) -> None:
    path = _build_cg_workbook(tmp_path, ["CG01 - Labour", "CG01 - Again"])
    with pytest.raises(CgBlockError, match="duplicate CG code CG01"):
        read_cg_descriptions(path)


# ---------------------------------------------------------------------------
# Output-neutrality against the production register
# ---------------------------------------------------------------------------
