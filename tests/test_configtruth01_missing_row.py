"""CONFIGTRUTH-01 P2 as ruled at task -3: a missing run-settings row resolves from the binding, or is refused.

``read_run_config`` used to fill 100,000 iterations and P80 when the run-settings sheet had no such row,
so a defaulted value reached the deliverable looking like a chosen one. The Director's ruling (08/10/2026):
a project declares its default iteration count and reporting P-level in its binding. For each setting,
the register's row wins when present; otherwise the binding's declared value is used and the manifest
says so; otherwise the run is refused naming the row, the sheet and that the binding declares none. A
row that is present but blank stays refused, as it was.

The stubs are the shipped wind register with one MC_Params row removed (label and value cleared, so the
reader's walk does not see it), and a separate strawman whose Run_Config_Preview lacks the row, which is
the other layout the reader accepts. The test bindings declare values the registers do not carry (7000,
85), so a resolved value can only have come from the binding.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest

from src.final.deliverable_preflight import run_preflight
from src.integrated.stage3_config import build_stage3_config, stage3_config_summary
from src.sra.cpm.bindings import registry
from src.sra.cpm.bindings.registry import ProjectBinding, get_binding, register_binding
from src.sra.methodology_defaults_reader import MethodologyDefaultsError
from src.sra.strawman_reader import read_run_config

REPO = Path(__file__).resolve().parents[1]
WIND_DIR = REPO / "tests" / "fixtures" / "synthetic_wind"
WIND_REGISTER = WIND_DIR / "wind_register.xlsx"
WIND_XER = WIND_DIR / "wind_schedule.xer"

ITERATIONS_ROW = "Iterations (n)"
P_LEVEL_ROW = "Reporting P-Level (pct)"
DECLARES = "configtruth01_declares"
DECLARES_NONE = "configtruth01_declares_none"


@pytest.fixture
def bindings():
    """Two test bindings: one declaring both defaults, one declaring neither."""
    register_binding(ProjectBinding(name=DECLARES, project_name="Declares", float_threshold_hours=480.0,
                                    default_iterations=7000, default_reporting_p_level=85),
                     overwrite=True)
    register_binding(ProjectBinding(name=DECLARES_NONE, project_name="Declares none",
                                    float_threshold_hours=480.0), overwrite=True)
    yield
    for name in (DECLARES, DECLARES_NONE):
        registry._REGISTRY.pop(name, None)


def _register_without(tmp_path: Path, *labels: str) -> Path:
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    wb = openpyxl.load_workbook(register)
    for label in labels:
        for row in wb["MC_Params"].iter_rows():
            if row[1].value == label:
                row[1].value = None
                row[2].value = None
                break
        else:
            raise AssertionError(f"the wind fixture no longer carries a {label!r} row - update this test")
    wb.save(register)
    wb.close()
    return register


def _register_blank(tmp_path: Path, label: str) -> Path:
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    wb = openpyxl.load_workbook(register)
    for row in wb["MC_Params"].iter_rows():
        if row[1].value == label:
            row[2].value = None
            break
    wb.save(register)
    wb.close()
    return register


def _strawman_without(tmp_path: Path, missing: str) -> Path:
    """A separate strawman: Run_Config_Preview, Parameter/Value in columns A/B."""
    rows = {"Iterations": 50000, "P-Level_reporting": 80,
            "Day_rate_default ($/wd)": 35000, "Schedule duration (weeks)": 90}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Run_Config_Preview"
    ws.append(["Parameter", "Value"])
    for label, value in rows.items():
        if label != missing:
            ws.append([label, value])
    path = tmp_path / "strawman.xlsx"
    wb.save(path)
    wb.close()
    return path


def _build(register: Path, strawman: Path, project: str, **kw):
    return build_stage3_config(register_path=register, xer_path=WIND_XER, strawman_path=strawman,
                               project=project, caller_module="test_configtruth01", seed=42, **kw)


# --- the reader --------------------------------------------------------------------------------

@pytest.mark.parametrize("key, label", [("iterations", ITERATIONS_ROW), ("p_level_reporting", P_LEVEL_ROW)])
def test_the_reader_fills_no_default_for_a_missing_row(key, label, tmp_path):
    cfg = read_run_config(_register_without(tmp_path, label))
    assert key not in cfg, cfg.get(key)


def test_the_reader_says_which_sheet_it_read(tmp_path):
    assert read_run_config(WIND_REGISTER)["run_config_sheet"] == "MC_Params"
    assert read_run_config(_strawman_without(tmp_path, "nothing"))["run_config_sheet"] == "Run_Config_Preview"


# --- the binding -------------------------------------------------------------------------------

@pytest.mark.parametrize("project", ["synthetic_wind", "synthetic_road"])
def test_each_shipped_binding_declares_what_its_register_carries(project):
    binding = get_binding(project)
    cfg = read_run_config(REPO / binding.data_dirs["register"] / binding.default_filenames["register"])
    assert (binding.default_iterations, binding.default_reporting_p_level) == (50000, 80)
    assert (cfg["iterations"], cfg["p_level_reporting"]) == (50000, 80)


@pytest.mark.parametrize("field", ["default_iterations", "default_reporting_p_level"])
@pytest.mark.parametrize("value", [0, -1, True, 80.0, "80"])
def test_a_binding_cannot_declare_a_default_that_is_not_a_positive_whole_number(field, value):
    # Without this, "unexpected keyword argument" also matches the field's name, and the test
    # passes on a binding that has no such field at all.
    assert field in ProjectBinding.__dataclass_fields__
    with pytest.raises((TypeError, ValueError), match=field):
        ProjectBinding(name="configtruth01_bad", project_name="Bad", **{field: value})


# --- resolution, and the source the manifest records -------------------------------------------

@pytest.mark.parametrize("label, key, expected", [
    (ITERATIONS_ROW, "iterations", 7000), (P_LEVEL_ROW, "reporting_p_level", 85)])
def test_a_missing_row_resolves_from_the_binding_and_the_manifest_says_so(
        label, key, expected, bindings, tmp_path):
    register = _register_without(tmp_path, label)
    summary = stage3_config_summary(_build(register, register, DECLARES))
    assert summary[key] == expected
    assert summary[f"{key}_source"] == "binding default"


def test_a_present_row_is_the_registers_and_an_override_says_override(bindings):
    summary = stage3_config_summary(_build(WIND_REGISTER, WIND_REGISTER, DECLARES))
    assert (summary["iterations"], summary["iterations_source"]) == (50000, "register")
    assert (summary["reporting_p_level"], summary["reporting_p_level_source"]) == (80, "register")
    summary = stage3_config_summary(_build(WIND_REGISTER, WIND_REGISTER, DECLARES, iterations=1000))
    assert (summary["iterations"], summary["iterations_source"]) == (1000, "override")


@pytest.mark.parametrize("label", [ITERATIONS_ROW, P_LEVEL_ROW])
def test_a_missing_row_with_no_binding_default_is_refused_naming_row_sheet_and_binding(
        label, bindings, tmp_path):
    register = _register_without(tmp_path, label)
    with pytest.raises(MethodologyDefaultsError) as exc:
        _build(register, register, DECLARES_NONE)
    message = str(exc.value)
    assert label in message and "MC_Params" in message and "declares none" in message, message


@pytest.mark.parametrize("label", ["Iterations", "P-Level_reporting"])
def test_a_missing_run_config_preview_row_is_refused_naming_that_sheet(label, bindings, tmp_path):
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    strawman = _strawman_without(tmp_path, label)
    with pytest.raises(MethodologyDefaultsError) as exc:
        _build(register, strawman, DECLARES_NONE)
    assert "Run_Config_Preview" in str(exc.value) and "declares none" in str(exc.value), str(exc.value)


@pytest.mark.parametrize("label", [ITERATIONS_ROW, P_LEVEL_ROW])
def test_a_blank_row_is_still_refused_even_where_the_binding_declares_a_default(label, bindings, tmp_path):
    register = _register_blank(tmp_path, label)
    with pytest.raises(MethodologyDefaultsError, match="blank"):
        _build(register, register, DECLARES)


# --- the deliverable's pre-flight --------------------------------------------------------------

def _preflight(register: Path, project: str) -> dict:
    env = {"QRA_PROJECT": project, "REGISTER_FILE": str(register), "P6_FILE": str(WIND_XER),
           "STRAWMAN_FILE": str(register)}
    return {r.label: r for r in run_preflight(env)}


@pytest.mark.parametrize("label, check", [(ITERATIONS_ROW, "iterations"), (P_LEVEL_ROW, "reporting P-level")])
def test_the_preflight_passes_a_missing_row_the_binding_supplies_and_says_where_from(
        label, check, bindings, tmp_path):
    result = _preflight(_register_without(tmp_path, label), DECLARES)[check]
    assert result.ok and "binding default" in result.detail, result


@pytest.mark.parametrize("label, check", [(ITERATIONS_ROW, "iterations"), (P_LEVEL_ROW, "reporting P-level")])
def test_the_preflight_refuses_a_missing_row_the_binding_does_not_supply(label, check, bindings, tmp_path):
    result = _preflight(_register_without(tmp_path, label), DECLARES_NONE)[check]
    assert not result.ok and "absent" in result.detail and "declares none" in result.detail, result
