"""3c Part D-1 / SPEC §3 limb 4, first half: two surfaces, one register, one number.

SPEC 1.2 §3 says the product is one engine behind several surfaces, and limb 4 is what makes
the other three checkable: *"the same register through Tier 1 and Tier 2 produces the same
numbers, gated the way the wind bit-identity anchor is."*

Until now that was an assertion. The xlsx surface reaches the cost engine through
``src/tier1_headless`` - the same code the xlwings bundle runs in the task pane - and the
notebook surface reaches it through ``src/final/deliverable``, rebuilt at Part A. Nothing
compared them.

**Which two arms, measured rather than assumed.** The headless path is bit-identical to the
deliverable's *integrated DP_Review* cost arm, and to neither of the other two candidates: the
standalone cost MC differs (they are the two sides of cross-validation check 1, which carries a
tolerance precisely because they are independent runs), and the Stage 3 cost arm differs again
because it consumes severity uniforms shared with the schedule arm. Pairing the wrong two would
have produced a test that could only pass with a tolerance - which GOLDEN01 rule 2 bans, and
rightly: a tolerance here would hide the drift this test exists to catch.

Equality is exact, on ``float.hex()``, as the goldens are.
"""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from src.register_contract import CONTRACT_VERSION, REQUIRED_NAMED_RANGES, check_contract
from tests.input_skips import private_documents_skip

REPO = Path(__file__).resolve().parents[1]
WIND_DIR = REPO / "tests" / "fixtures" / "synthetic_wind"
WIND_REGISTER = WIND_DIR / "wind_register.xlsx"
WIND_XER = WIND_DIR / "wind_schedule.xer"

#: Small enough to run in the stage lane, large enough that a real divergence shows.
ITERATIONS = 3000
SEED = 42


@pytest.fixture
def wind_project():
    import src.sra.cpm.bindings.synthetic_wind  # noqa: F401  self-registers

    saved = os.environ.get("QRA_PROJECT")
    os.environ["QRA_PROJECT"] = "synthetic_wind"
    try:
        yield
    finally:
        if saved is None:
            os.environ.pop("QRA_PROJECT", None)
        else:
            os.environ["QRA_PROJECT"] = saved


@pytest.fixture
def wind_env() -> dict[str, str]:
    return {
        "QRA_PROJECT": "synthetic_wind",
        "REGISTER_FILE": str(WIND_REGISTER),
        "P6_FILE": str(WIND_XER),
        "STRAWMAN_FILE": str(WIND_REGISTER),
        "SR9_ITERATIONS": str(ITERATIONS),
        # CONFIGTRUTH-01: unset is now a drawn seed; the Tier-1 side is pinned at SEED.
        "SR9_SEED": str(SEED),
    }


# --- limb 4, the number --------------------------------------------------------------------

@pytest.mark.slow
def test_the_two_surfaces_produce_the_same_cost_headline(wind_env, wind_project):
    """The whole point of limb 4, on one register, with no tolerance anywhere."""
    from src.final.deliverable import resolve_inputs, run_engines
    from src.tier1_headless import run_headless_from_workbook

    notebook = np.asarray(
        run_engines(resolve_inputs(wind_env, output_dir=Path(tempfile.mkdtemp())))
        .integrated_results["cost_dollars_array"], dtype=float)

    xlsx = run_headless_from_workbook(WIND_REGISTER, "WT", ITERATIONS, SEED, "shared-pool")

    assert xlsx["p_table"], "the Tier-1 surface produced no P-table"
    mismatches = [
        (p, value.hex(), float(np.percentile(notebook, p)).hex())
        for p, value in xlsx["p_table"]
        if value.hex() != float(np.percentile(notebook, p)).hex()
    ]
    assert mismatches == [], (
        "the xlsx surface and the notebook surface disagree on the cost headline from the "
        f"same register: {mismatches[:3]}"
    )


@pytest.mark.slow
def test_the_two_surfaces_agree_on_the_mean_not_only_the_percentiles(wind_env, wind_project):
    """A P-table can match while the underlying distribution differs - percentiles are a
    lossy view. The mean is computed from every sample, so it catches what they miss."""
    from src.final.deliverable import resolve_inputs, run_engines
    from src.tier1_headless import run_headless_from_workbook

    notebook = np.asarray(
        run_engines(resolve_inputs(wind_env, output_dir=Path(tempfile.mkdtemp())))
        .integrated_results["cost_dollars_array"], dtype=float)
    xlsx = run_headless_from_workbook(WIND_REGISTER, "WT", ITERATIONS, SEED, "shared-pool")

    assert float(xlsx["cost_mean"]).hex() == float(notebook.mean()).hex(), (
        f"means differ: xlsx {xlsx['cost_mean']!r} vs notebook {notebook.mean()!r}")


@pytest.mark.slow
def test_the_agreement_is_not_an_accident_of_the_iteration_count(wind_env, wind_project):
    """Run both at a different count. If they agreed only at one n, the equality above would
    be a coincidence of the seed rather than one engine behind two surfaces."""
    from src.final.deliverable import resolve_inputs, run_engines
    from src.tier1_headless import run_headless_from_workbook

    n = 1500
    env = dict(wind_env, SR9_ITERATIONS=str(n))
    notebook = np.asarray(
        run_engines(resolve_inputs(env, output_dir=Path(tempfile.mkdtemp())))
        .integrated_results["cost_dollars_array"], dtype=float)
    xlsx = run_headless_from_workbook(WIND_REGISTER, "WT", n, SEED, "shared-pool")

    assert float(xlsx["cost_mean"]).hex() == float(notebook.mean()).hex()


# --- limb 4, the contract ------------------------------------------------------------------

def test_both_surfaces_see_the_same_contract_version():
    """§3 limb 2: 'the xlsx and notebook surfaces read the same named ranges at the same
    CONTRACT_VERSION'. One register, one version, and the engine's own constant."""
    wb = openpyxl.load_workbook(WIND_REGISTER, data_only=True)
    try:
        version = wb["Settings"]["B6"].value
        present = set(wb.defined_names)
        grid = [list(r) for r in wb["Methodology"].iter_rows(values_only=True)]
    finally:
        wb.close()

    assert str(version).strip() == CONTRACT_VERSION, (
        f"the register declares contract {version!r}, the engine expects {CONTRACT_VERSION!r}")

    ok, message = check_contract(present, grid, version)
    assert ok, f"the shared register fails the contract both surfaces read it under: {message}"


def test_the_register_carries_every_named_range_the_contract_requires():
    wb = openpyxl.load_workbook(WIND_REGISTER)
    try:
        missing = sorted(REQUIRED_NAMED_RANGES - set(wb.defined_names))
    finally:
        wb.close()
    assert missing == [], f"the register is missing contract names: {missing}"


def test_the_shipped_template_satisfies_the_same_contract():
    """The register a user is handed must pass the check both surfaces apply to it.

    The wind register is a filled fixture; the template is what actually ships. A contract
    that only the fixture satisfies would be green here and useless to a user.
    """
    template = REPO / "templates" / "register_template.xlsx"
    wb = openpyxl.load_workbook(template, data_only=True)
    try:
        ok, message = check_contract(
            set(wb.defined_names),
            [list(r) for r in wb["Methodology"].iter_rows(values_only=True)],
            wb["Settings"]["B6"].value,
        )
    finally:
        wb.close()
    assert ok, f"the shipped template fails the register contract: {message}"


@pytest.mark.slow
def test_the_notebook_surface_reads_the_register_the_environment_named(wind_env, wind_project):
    """Limb 4 is only meaningful if both surfaces read the SAME file. The xlsx side is handed
    the path; the notebook side resolves it from the environment, and this pins that they
    land on one register rather than two that happen to agree."""
    from src.final.deliverable import resolve_inputs

    inputs = resolve_inputs(wind_env, output_dir=Path(tempfile.mkdtemp()))
    assert inputs.register_path == WIND_REGISTER
    assert inputs.register_path.samefile(WIND_REGISTER)


# --- D-2 / SPEC §9.3: the manuals describe the surfaces as they are ------------------------
#
# The spec's lock action §9.3 was a doc pass, owed since 02/09: the agent manual, README.md and
# ARCHITECTURE.md described notebooks that were not in the tree. After Part A there is exactly
# one, so the correction is not "delete the lines" but "say one where they said four".
#
# A-4's guard already holds that no manual names a retired .py/.bat ENTRY POINT as present.
# This is its counterpart for the retired NOTEBOOKS, which that set does not cover.

#: Retired at 3c: their content is cells of `final_runner.ipynb`.
RETIRED_NOTEBOOKS = (
    "sra_runner.ipynb",
    "integrated_runner.ipynb",
    "client_runner.ipynb",
    "client_appendices.ipynb",
    "report_template.ipynb",
)

# The superseded ARCHITECTURE.md copy left out since PUBLIC-01 reduced it to a
# pointer: it names no notebook at all, so checking it would pass vacuously.
MANUALS = ("CLAUDE.md", "README.md", "ARCHITECTURE.md")  # link-rule: the agent manual is checked for notebook names when present, skipped when absent


@pytest.mark.parametrize("manual", private_documents_skip(MANUALS))
def test_no_manual_lists_a_retired_notebook_as_present(manual):
    """Naming one as retired is fine; listing it as something a reader can run is not.

    Checked per PARAGRAPH, not per line. Prose wraps, so "sra_runner, integrated_runner and
    client_runner are / retired" puts the mention and its qualifier on different lines, and the
    line-based first version called that an offender - a false positive from reading the wrong
    unit. The paragraph is the unit a reader takes meaning from, so it is the unit to check.
    """
    text = (REPO / manual).read_text(encoding="utf-8")
    blocks = [b for b in text.split("\n\n") if b.strip()]
    excuses = ("retire", "superseded", "removed", "no longer", "do not look for")
    offenders = [
        f"{manual}: {b.strip()[:110]}"
        for b in blocks
        if any(n.replace(".ipynb", "") in b for n in RETIRED_NOTEBOOKS)
        and not any(word in b.lower() for word in excuses)
    ]
    assert offenders == [], (
        "these passages present a retired notebook as available:\n" + "\n".join(offenders))


@pytest.mark.parametrize("manual", private_documents_skip(MANUALS))
def test_every_manual_names_the_one_notebook_that_does_exist(manual):
    """The other half. A doc pass that only deletes leaves a reader knowing what NOT to run."""
    assert "final_runner.ipynb" in (REPO / manual).read_text(encoding="utf-8"), (
        f"{manual} does not mention the deliverable notebook at all")


def test_the_notebook_the_manuals_name_is_in_the_tree():
    """§9.3's original defect was manuals describing an ABSENT notebook as present. The
    inverse of this test is the one that was missing, so it belongs with them."""
    assert (REPO / "final_runner.ipynb").is_file()
