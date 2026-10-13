"""REG-02 locks: theming + EMV/Mode analysis are presentation-only.

The load-bearing guarantee (Director §11): theming must NOT mask the REG-01
builder-side Mode assertion. A themed sheet with a Mode dropdown that quietly
disabled the blank/unrecognised-Mode build failure would be the worst kind of
regression - it looks fine. So this proves the guard still FIRES, and fires
BEFORE any themed output is written (fail-fast, upstream of theming/save).
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
BUILD_SCRIPT = REPO / "scripts" / "build_synthetic_wind_fixture.py"
WIND_REGISTER = REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"


def _load_builder():
    spec = importlib.util.spec_from_file_location("_reg02_builder", BUILD_SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


BUILDER = _load_builder()


def test_mode_guard_still_fires_after_theming_and_writes_no_file(tmp_path, monkeypatch):
    """A bad Mode fails the themed build - and fails BEFORE any file is written
    (the guard is upstream of theming + save; theming cannot mask it)."""
    good = BUILDER.WIND_RISKS
    corrupted, hit = [], False
    for r in good:
        r = list(r)
        if r[13] == "Yes" and not hit:  # r[13]=Time flag, r[14]=Mode
            r[14] = ""
            hit = True
        corrupted.append(tuple(r))
    assert hit
    monkeypatch.setattr(BUILDER, "WIND_RISKS", corrupted)

    out = tmp_path / "themed_bad_mode.xlsx"
    with pytest.raises(ValueError, match="no-silent-zero"):
        BUILDER.build_register(out)
    assert not out.exists(), "guard must fail-fast before any themed file is saved"


def test_register_is_themed():
    """Themed Risk Register: house accent tab, freeze panes, Mode dropdown."""
    wb = load_workbook(WIND_REGISTER)
    ws = wb["Risk Register"]
    assert ws.freeze_panes == "D5"
    assert ws.sheet_properties.tabColor.rgb.endswith("0A7DFF")  # accent2
    # header fill is the house accent1
    assert ws.cell(row=4, column=1).fill.fgColor.rgb.endswith("231EDC")
    # a Mode data-validation dropdown exists over the Mode column (col 26 = Z)
    dvs = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    assert any("Z5" in str(dv.sqref) for dv in dvs), "Mode dropdown missing"
    wb.close()


def test_emv_mode_analysis_sheet_present_with_native_charts():
    """The analytics sheet exists with native openpyxl charts (not images)."""
    import zipfile
    wb = load_workbook(WIND_REGISTER)
    assert "EMV & Mode Analysis" in wb.sheetnames
    wb.close()
    with zipfile.ZipFile(WIND_REGISTER) as z:
        charts = [n for n in z.namelist()
                  if n.startswith("xl/charts/chart") and n.endswith(".xml")]
    assert len(charts) >= 3, f"expected >=3 native charts, got {charts}"


def test_theming_value_neutral_reads(monkeypatch):
    """Theming touches no value the engine reads: 35 risks, 10 Time=Yes.

    Uses monkeypatch.setenv so QRA_PROJECT is RESTORED after the test -
    a bare os.environ set leaks the WT binding into later RefProject tests.
    """
    spec = importlib.util.spec_from_file_location(
        "_reg02_wind_binding",
        REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_binding.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # self-registers 'synthetic_wind'
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")

    from src.sra.strawman_reader import read_risk_scope
    from src.v5_register_reader import read_risks_v5
    risks = read_risks_v5(str(WIND_REGISTER), risk_id_prefix="WT")
    assert len(risks) == 35
    scope = read_risk_scope(str(WIND_REGISTER))
    assert (scope["time_impact"] == "Yes").sum() == 10
