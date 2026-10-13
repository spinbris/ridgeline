"""REG03 S3 - dynamic settings cascade (S3a) + Layer1_Correlation auto-resize (S3b).

S3a: the six S1 inline dropdowns become named-range lists on a Lists sheet, plus the
S1-deferred RBS L1/L2 and Package/Factor lists. Each List_* is an OFFSET dynamic range
with a MAX(...,1) guard (extends as the user adds entries; no #REF! on empty). Value-
neutral (input constraints only).

S3b: `resize_layer1_correlation` grows/shrinks the factor matrix to the factor list,
preserving existing coefficients and defaulting new cells (diag 1, off-diag 0). Value-
neutral for an unchanged factor set.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.register_contract import LIST_NAMED_RANGES, REQUIRED_NAMED_RANGES
from tools.register_scaffold import (
    LIST_NAMES, LISTS_SHEET, build_template, resize_layer1_correlation,
)


# --- S3a: the cascade -----------------------------------------------------

@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    return build_template(tmp_path_factory.mktemp("s3") / "t.xlsx")


def test_lists_sheet_and_all_list_names_present(template):
    wb = openpyxl.load_workbook(template)
    try:
        assert LISTS_SHEET in wb.sheetnames
        got = {n for n in wb.defined_names if n.startswith("List_")}
    finally:
        wb.close()
    assert got == set(LIST_NAMES) == set(LIST_NAMED_RANGES)
    assert len(got) == 10


def test_list_names_are_offset_max_guarded(template):
    """Dynamic (extends on add) AND #REF!-safe on an empty list (MAX(...,1))."""
    wb = openpyxl.load_workbook(template)
    try:
        for name in LIST_NAMES:
            ref = wb.defined_names[name].attr_text
            assert ref.startswith("OFFSET("), (name, ref)
            assert "MAX(COUNTA(" in ref, (name, ref)
    finally:
        wb.close()


def test_dropdowns_are_list_sourced_including_rbs(template):
    """The Risk Register dropdowns reference the List_* names; RBS L1/L2 (cols
    E/F = 5/6, S1-deferred) now have dropdowns."""
    wb = openpyxl.load_workbook(template)
    try:
        rr = wb["Risk Register"]
        dvs = rr.data_validations.dataValidation
        formulas = {str(dv.formula1) for dv in dvs}
        sqrefs = " ".join(str(dv.sqref) for dv in dvs)
    finally:
        wb.close()
    assert "List_Distribution" in formulas and "List_CGCode" in formulas
    assert "List_RBS_L1" in formulas and "List_RBS_L2" in formulas
    assert not any('"' in f for f in formulas)  # no inline literal lists remain
    assert "E5" in sqrefs and "F5" in sqrefs   # RBS L1/L2 dropdowns present


def test_list_names_in_required_contract():
    assert LIST_NAMED_RANGES <= REQUIRED_NAMED_RANGES


# --- S3b: the auto-resize -------------------------------------------------

def _corr_sheet(names, coeffs=None, mode="elicited"):
    """A minimal Layer1_Correlation sheet: Corr_Mode / PSD_Status / blank / Factor
    header / lower-triangle. coeffs is {frozenset(a,b): v} for off-diagonal."""
    coeffs = coeffs or {}
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Layer1_Correlation"
    s.append(["Corr_Mode", mode, "(zero | elicited | one)"])
    s.append(["PSD_Status", "valid", ""])
    s.append([])
    s.append(["Factor", *names])
    for i, a in enumerate(names):
        row = [a]
        for j in range(len(names)):
            if j > i:
                row.append(None)
            elif j == i:
                row.append(1.0)
            else:
                row.append(coeffs.get(frozenset((a, names[j])), 0.0))
        s.append(row)
    return wb, s


def _grid(s, ncol=10):
    return [[s.cell(row=r, column=c).value for c in range(1, ncol + 1)]
            for r in range(1, s.max_row + 1)]


def test_resize_same_set_is_byte_unchanged():
    names = ["A", "B", "C"]
    coeffs = {frozenset(("A", "B")): 0.3, frozenset(("A", "C")): 0.2,
              frozenset(("B", "C")): 0.4}
    wb, s = _corr_sheet(names, coeffs)
    before = _grid(s)
    resize_layer1_correlation(s, names)
    assert _grid(s) == before  # value-neutral


def test_resize_grow_preserves_and_defaults_new():
    names = ["A", "B"]
    wb, s = _corr_sheet(names, {frozenset(("A", "B")): 0.5})
    resize_layer1_correlation(s, ["A", "B", "C"])
    # header row = 4; C is the 3rd factor -> row 7
    assert s.cell(row=7, column=1).value == "C"
    assert s.cell(row=7, column=2).value == 0.0   # C-A off-diagonal default
    assert s.cell(row=7, column=3).value == 0.0   # C-B off-diagonal default
    assert s.cell(row=7, column=4).value == 1.0   # C-C diagonal
    assert s.cell(row=6, column=2).value == 0.5   # A-B preserved


def test_resize_shrink_drops_factor():
    names = ["A", "B", "C"]
    coeffs = {frozenset(("A", "B")): 0.3, frozenset(("B", "C")): 0.4}
    wb, s = _corr_sheet(names, coeffs)
    resize_layer1_correlation(s, ["A", "B"])
    hdr = [s.cell(row=4, column=c).value for c in range(2, 6)]
    assert hdr == ["A", "B", None, None]           # C dropped from header
    assert s.cell(row=6, column=2).value == 0.3     # A-B preserved
    assert s.cell(row=7, column=1).value is None    # C's row gone


def test_resize_updates_mode_only_when_asked():
    wb, s = _corr_sheet(["A", "B"], mode="elicited")
    resize_layer1_correlation(s, ["A", "B"])
    assert s.cell(row=1, column=2).value == "elicited"  # untouched by default
    resize_layer1_correlation(s, ["A", "B"], corr_mode="zero")
    assert s.cell(row=1, column=2).value == "zero"
