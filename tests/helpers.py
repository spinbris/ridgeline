"""Shared test helpers for SR9 v3.2 content-identity and grep checks."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path


def extract_full_docx_text(path: str | Path) -> str:
    """Return all paragraph text + table cell text from a docx as one
    newline-joined string.

    Uses python-docx for paragraphs and tables. Headers/footers are not
    included by default; callers that need footers should use
    extract_docx_with_footers.
    """
    from docx import Document
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_docx_with_footers(path: str | Path) -> str:
    """Return body text + footer text from a docx.

    Used for tests that assert on visible footer content.
    """
    from docx import Document
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in doc.sections:
        for p in section.footer.paragraphs:
            parts.append(p.text)
    return "\n".join(parts)


def extract_html_text(path: str | Path) -> str:
    """Return concatenated visible text from an HTML file (strip tags)."""
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    # Strip <script>...</script> and <style>...</style> blocks first
    raw = re.sub(r"<script.*?</script>", " ", raw, flags=re.DOTALL | re.IGNORECASE)
    raw = re.sub(r"<style.*?</style>", " ", raw, flags=re.DOTALL | re.IGNORECASE)
    # Strip remaining tags
    return re.sub(r"<[^>]+>", " ", raw)


def diff_docx_text(path_a: str | Path, path_b: str | Path) -> list[str]:
    """Return a list of diff lines between two docx files, comparing
    canonical text content only (not bytes).

    Empty list = identical text. Differences include both A-only and
    B-only lines so caller can see what shifted.
    """
    import difflib
    a = extract_full_docx_text(path_a).splitlines()
    b = extract_full_docx_text(path_b).splitlines()
    return [
        line for line in difflib.unified_diff(a, b, lineterm="")
        if not line.startswith(("---", "+++", "@@"))
    ]


def docx_byte_size(path: str | Path) -> int:
    """File size in bytes."""
    return Path(path).stat().st_size
