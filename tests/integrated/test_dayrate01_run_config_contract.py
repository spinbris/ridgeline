"""DAYRATE-01: the run-config contract shared by both integrated loaders.

Before DAYRATE-01, ``load_integrated_config`` and
``load_integrated_config_stage3`` each carried their own copy of

    try:
        cfg = read_run_config(strawman_path)
    except Exception:
        cfg = {}

so a strawman that existed but would not parse was swallowed and the run
proceeded on four hardcoded values with nothing raised - including the
client's $/wd day rate, which would then be a plausible wrong number from
a different project. The same block also tested
``if cfg.get("day_rate_default")``, which is falsy at zero, so a deliberate
zero was silently replaced by the hardcoded rate.

These are the positive controls for that mechanism. Each must be witnessed
red before its pass is evidence (PC-1 .. PC-3 in the brief). They use only
synthetic workbooks, so they run on a fresh clone with no confidential input
present.
"""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pytest

from src.integrated.run_config_contract import (
    DAY_RATE_KEY,
    DAY_RATE_SOURCE,
    RunConfigError,
    day_rate_audit_fields,
    read_run_config_strict,
    resolve_day_rate,
)


from tests._fixture_values import FIXTURE_DAY_RATE
def _write_strawman(path: Path, day_rate: float | None) -> Path:
    """Minimal Run_Config_Preview strawman. ``None`` omits the day-rate row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Run_Config_Preview"
    rows: list[tuple[str, object]] = [
        ("iterations", 5000),
        ("p-level_reporting", 80),
    ]
    if day_rate is not None:
        rows.append(("day_rate_default ($/wd)", day_rate))
    for i, (label, value) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# PC-2 - the load-bearing control: the swallowing path must be able to fail
# ---------------------------------------------------------------------------

class TestMalformedStrawmanStopsTheRun:
    """PC-2. The swallow at the old ``except Exception: cfg = {}`` had never
    been shown capable of failing; until it is, its pass is not evidence."""

    def test_read_run_config_strict_raises_on_malformed_workbook(self, tmp_path):
        bad = tmp_path / "corrupt.xlsx"
        bad.write_bytes(b"this is not a workbook")

        from src.sra.strawman_reader import read_run_config

        with pytest.raises(RunConfigError) as exc:
            read_run_config_strict(bad, read_run_config)

        # The refusal must name the file, so the operator knows what to fix.
        assert str(bad) in str(exc.value)

    def test_read_run_config_strict_chains_the_underlying_cause(self, tmp_path):
        """The swallowed exception is preserved, not discarded."""
        bad = tmp_path / "corrupt.xlsx"
        bad.write_bytes(b"this is not a workbook")

        from src.sra.strawman_reader import read_run_config

        with pytest.raises(RunConfigError) as exc:
            read_run_config_strict(bad, read_run_config)

        assert exc.value.__cause__ is not None

    def test_absent_strawman_file_stops_the_run(self, tmp_path):
        from src.sra.strawman_reader import read_run_config

        missing = tmp_path / "no_such_strawman.xlsx"
        with pytest.raises(RunConfigError) as exc:
            read_run_config_strict(missing, read_run_config)

        assert str(missing) in str(exc.value)

    def test_loaders_no_longer_swallow(self, tmp_path):
        """Both Class C sites, through their public entry points.

        The raise happens while resolving the strawman, which is before either
        loader reads the register, so this needs no confidential input.
        """
        bad = tmp_path / "corrupt.xlsx"
        bad.write_bytes(b"this is not a workbook")
        reg = tmp_path / "reg.xlsx"
        p6 = tmp_path / "p6.xlsx"
        reg.write_bytes(b"")
        p6.write_bytes(b"")

        from src.integrated.integrated_simulation import load_integrated_config
        from src.integrated.integrated_simulation_stage3 import (
            load_integrated_config_stage3,
        )

        with pytest.raises(RunConfigError):
            load_integrated_config(str(reg), str(p6), str(bad))

        with pytest.raises(RunConfigError):
            load_integrated_config_stage3(str(reg), str(p6), str(bad))


# ---------------------------------------------------------------------------
# PC-1 - an absent day rate stops the run, naming the input and its source
# ---------------------------------------------------------------------------

class TestAbsentDayRateStopsTheRun:

    def test_absent_day_rate_raises_naming_input_and_source(self, tmp_path):
        from src.sra.strawman_reader import read_run_config

        straw = _write_strawman(tmp_path / "no_rate.xlsx", day_rate=None)
        cfg = read_run_config_strict(straw, read_run_config)

        with pytest.raises(RunConfigError) as exc:
            resolve_day_rate(cfg, straw)

        msg = str(exc.value)
        assert DAY_RATE_KEY in msg, "the refusal must name the missing input"
        assert DAY_RATE_SOURCE in msg, "the refusal must name where it comes from"
        assert str(straw) in msg, "the refusal must name the workbook"

    def test_refusal_does_not_quote_a_substitute_rate(self, tmp_path):
        """The message must not carry the client's commercial rate.

        REG-01's equivalent refusal embeds ``FIXTURE_DAY_RATE`` in its text.
        This contract names the input instead, so the refusal adds no further
        instance of a confidential figure (GOLDEN01 item 8).
        """
        from src.sra.strawman_reader import read_run_config

        straw = _write_strawman(tmp_path / "no_rate.xlsx", day_rate=None)
        cfg = read_run_config_strict(straw, read_run_config)

        with pytest.raises(RunConfigError) as exc:
            resolve_day_rate(cfg, straw)

        msg = str(exc.value)
        # Both spellings are DERIVED from the constant, never typed: a test
        # that guards a confidential figure must not itself put that figure
        # into tracked source (DAYRATE-01 review finding 2). The underscored
        # spelling is not covered because it cannot arise here - any leak
        # would come from interpolating the float, which renders plain.
        from tests._fixture_values import FIXTURE_DAY_RATE

        for form in (f"{int(FIXTURE_DAY_RATE)}", f"{int(FIXTURE_DAY_RATE):,}"):
            assert form not in msg, (
                f"the refusal leaked the client day rate ({form!r} spelling)"
            )

    @pytest.mark.parametrize("bad", ["", "   ", "N/A", None])
    def test_blank_or_non_numeric_day_rate_raises(self, tmp_path, bad):
        straw = tmp_path / "s.xlsx"
        with pytest.raises(RunConfigError):
            resolve_day_rate({DAY_RATE_KEY: bad}, straw)

    def test_non_finite_day_rate_raises(self, tmp_path):
        straw = tmp_path / "s.xlsx"
        with pytest.raises(RunConfigError):
            resolve_day_rate({DAY_RATE_KEY: float("nan")}, straw)
        with pytest.raises(RunConfigError):
            resolve_day_rate({DAY_RATE_KEY: float("inf")}, straw)


class TestConfiguredDayRateFlowsThrough:
    """The counter-control: the contract is not simply always-raise."""

    def test_workbook_day_rate_is_returned(self, tmp_path):
        from src.sra.strawman_reader import read_run_config

        straw = _write_strawman(tmp_path / "s.xlsx", day_rate=35_000.0)
        cfg = read_run_config_strict(straw, read_run_config)

        assert resolve_day_rate(cfg, straw) == 35_000.0


# ---------------------------------------------------------------------------
# PC-3 - an explicit zero is permitted, warned, and stamped
# ---------------------------------------------------------------------------

class TestExplicitZeroIsPermittedWarnedAndStamped:
    """A zero day rate is a legitimate cost-only run. Refusing it would push
    users to fake it with a near-zero value, which is worse. But it silently
    removes all time-related cost and is indistinguishable in the output from
    a correct run, so it warns and it is stamped."""

    def test_explicit_zero_is_returned_not_replaced(self, tmp_path):
        from src.sra.strawman_reader import read_run_config

        straw = _write_strawman(tmp_path / "zero.xlsx", day_rate=0)
        cfg = read_run_config_strict(straw, read_run_config)

        # The pre-DAYRATE-01 falsy test replaced this with the hardcoded rate.
        assert resolve_day_rate(cfg, straw) == 0.0

    def test_explicit_zero_warns(self, tmp_path, caplog):
        straw = tmp_path / "s.xlsx"
        with caplog.at_level(
            logging.WARNING, logger="src.integrated.run_config_contract",
        ):
            resolve_day_rate({DAY_RATE_KEY: 0.0}, straw)

        assert any(r.levelno >= logging.WARNING for r in caplog.records), (
            "a zero day rate must warn"
        )
        assert any("zero" in r.getMessage().lower() for r in caplog.records)

    def test_zero_is_stamped_and_a_real_rate_is_not(self):
        """The stamp is what survives the run; the warning scrolls past."""
        assert day_rate_audit_fields(0.0) == {"day_rate_explicit_zero": True}
        assert day_rate_audit_fields(35_000.0) == {}

    def test_stamp_is_absent_for_every_non_zero_rate(self):
        """Number-neutrality: no existing run artefact gains a key."""
        for rate in (1.0, 35_000.0, 1e6):
            assert day_rate_audit_fields(rate) == {}
