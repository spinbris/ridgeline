"""REG05-A1 - the client-ready scaffolder finish (value-neutral).

One regenerable finish (number formats, project-identity metadata, dynamic TOC list,
single-source cell classes + machine-read marker, tab colours, print setup, correlation
heatmap, autofilter, blank-input CF), applied to the empty template + the wind fixture
via shared helpers. Every check confirms formatting/metadata only - no read value moves.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.register_contract import (
    CONTRACT_VERSION, PROJECT_META_NAMES, REQUIRED_NAMED_RANGES,
)
from tools.register_scaffold import (
    EMITTED_ONLY_NAMES, RISK_HEADERS, _TOC_LEGEND_ROW, _TOC_TABLIST_FIRST_ROW,
    apply_finish, build_template,
)

REPO = Path(__file__).resolve().parents[1]
FIXTURE = REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    return build_template(tmp_path_factory.mktemp("reg05a1") / "t.xlsx")


def test_contract_bumped_with_project_meta(template):
    # REG05-A2 moved the contract to 1.4 (Layer-1 promoted); the A1 project-identity
    # metadata + exact-name contract still hold on the template.
    assert CONTRACT_VERSION == "1.4"
    assert PROJECT_META_NAMES <= REQUIRED_NAMED_RANGES
    wb = openpyxl.load_workbook(template)
    try:
        # SRA-02 S-B: Grid_Mapping is emitted, not required (D-3 - requiring it
        # would be a contract-version bump on the client-number-bearing lane).
        assert set(wb.defined_names) == set(REQUIRED_NAMED_RANGES) | set(EMITTED_ONLY_NAMES)
        assert wb["Settings"]["B6"].value == "1.4"
    finally:
        wb.close()


def test_project_block_and_core_properties():
    """Metadata block on TOC + core properties driven from it (fixture seeds them)."""
    wb = openpyxl.load_workbook(FIXTURE)
    try:
        names = set(wb.defined_names)
        assert PROJECT_META_NAMES <= names
        sh, coord = list(wb.defined_names["Set_ProjectName"].destinations)[0]
        assert sh == "TOC & Instructions"
        assert wb[sh][coord].value == "Tropic Ridge Wind Farm"
        assert wb.properties.title == "Tropic Ridge Wind Farm"
        # analysis-sheet title driven from the project name (hardcoded literal retired)
        assert wb["EMV & Mode Analysis"]["A1"].value == "EMV & Mode Analysis - Tropic Ridge Wind Farm"
    finally:
        wb.close()


def test_number_formats_by_column_class(template):
    wb = openpyxl.load_workbook(template)
    try:
        rr = wb["Risk Register"]
        idx = {h: i + 1 for i, h in enumerate(RISK_HEADERS)}
        # '$' is escaped (\$) so Excel's file-level validation accepts the stored
        # format code and does not repair the workbook on open.
        assert rr.cell(row=5, column=idx["Optimistic ($)"]).number_format == r"\$#,##0"
        assert rr.cell(row=5, column=idx["EMV_cost ($)"]).number_format == r"\$#,##0"
        assert rr.cell(row=5, column=idx["Probability"]).number_format == "0%"
        assert rr.cell(row=5, column=idx["Optimistic (days)"]).number_format == "0.0"
        assert wb["Methodology"]["B2"].number_format == "0.00"
    finally:
        wb.close()


def test_single_class_map_covers_the_escapee_columns(template):
    """The three cols S1 skipped (Risk ID/Status/Description) now carry the input class."""
    wb = openpyxl.load_workbook(template)
    try:
        rr = wb["Risk Register"]
        for col in (1, 2, 3):
            assert rr.cell(row=5, column=col).font.color.rgb.endswith("0000CC"), col
    finally:
        wb.close()


def test_machine_read_marker_on_the_machine_set(template):
    """Thin border on the model-read cells incl. MC_day_rate + Set_* + Grid_* inputs."""
    wb = openpyxl.load_workbook(template)
    try:
        st = wb["Settings"]
        assert st["B1"].border.left.style == "thin"          # Set_iterations
        mp_sh, mp_coord = list(wb.defined_names["MC_day_rate"].destinations)[0]
        assert wb[mp_sh][mp_coord].border.left.style == "thin"
    finally:
        wb.close()


def test_toc_tablist_generated_from_inventory(template):
    wb = openpyxl.load_workbook(template)
    try:
        toc = wb["TOC & Instructions"]
        # Read the reserved block from the module constants, not a literal window:
        # S-B grew the inventory past the old hardcoded 5-20 and this test was the
        # thing that noticed. Tying it to the constants stops the next growth
        # silently truncating what the test can see.
        listed = {toc.cell(row=r, column=1).value
                  for r in range(_TOC_TABLIST_FIRST_ROW, _TOC_LEGEND_ROW)
                  if toc.cell(row=r, column=1).value}
        actual = {s.title for s in wb.worksheets if s.title != "TOC & Instructions"}
        assert actual <= listed  # every real tab appears in the generated list
    finally:
        wb.close()


def test_tab_colours_and_gridlines(template):
    wb = openpyxl.load_workbook(template)
    try:
        assert all(s.sheet_properties.tabColor is not None for s in wb.worksheets)  # no uncoloured tab
        assert wb["Lists"].sheet_view.showGridLines is False
        assert wb["Reference (xlwings)"].sheet_view.showGridLines is False
    finally:
        wb.close()


def test_print_setup_and_autofilter(template):
    wb = openpyxl.load_workbook(template)
    try:
        rr = wb["Risk Register"]
        assert rr.page_setup.orientation == "landscape"
        assert (rr.print_title_rows or "").replace("$", "") == "3:4"
        assert rr.auto_filter.ref is not None
    finally:
        wb.close()


def test_correlation_matrix_styling_and_heatmap_in_fixture():
    """Layer-1 opt-in surface (present in the fixture): header styling + colour-scale CF."""
    wb = openpyxl.load_workbook(FIXTURE)
    try:
        cs = wb["Layer1_Correlation"]
        rules = [r for cf in cs.conditional_formatting for r in cs.conditional_formatting[cf]]
        assert any(r.type == "colorScale" for r in rules), "heatmap colour-scale missing"
    finally:
        wb.close()


def test_blank_required_input_conditional_format(template):
    wb = openpyxl.load_workbook(template)
    try:
        rr = wb["Risk Register"]
        rules = [r for cf in rr.conditional_formatting for r in rr.conditional_formatting[cf]]
        assert any(r.type == "expression" for r in rules), "blank-input formula CF missing"
    finally:
        wb.close()


def test_apply_finish_is_value_neutral_and_idempotent(template):
    """Re-running the finish moves no cell value (formatting/metadata only)."""
    wb = openpyxl.load_workbook(template)
    try:
        rr = wb["Risk Register"]
        before = [[c.value for c in row] for row in rr.iter_rows()]
        apply_finish(wb, project=None)   # idempotent second pass
        after = [[c.value for c in row] for row in rr.iter_rows()]
        assert before == after
    finally:
        wb.close()
