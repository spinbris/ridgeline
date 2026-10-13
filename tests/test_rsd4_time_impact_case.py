"""RS-D-4 (D-16, TIMEYES-CASE-01): ``time_impact`` case must not decide which rows simulate.

Excel counts ``<>"Yes"`` case-insensitively; before RS-D-4 the Python cross-check compared
exact-case and the reader kept whatever case the cell carried. Property: a scope frame whose
``time_impact`` cells differ from a canonical twin only in case and surrounding whitespace
simulates the same rows, on every consumer, and the reader emits the canonical form.

Written red-first: on the pre-RS-D-4 tree the cross-check test and the reader test fail.
"""
from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from src.sra import strawman_reader
from src.sra.cpm.bindings import active_risk_id_prefix
from src.sra.schedule_cross_check import compute_schedule_comparators
from src.sra.sra_simulation import simulating_risk_ids

GRID = {"Series": 1.0, "Hybrid": 0.5, "Parallel": 0.0}
COLS = ["risk_id", "time_impact", "mode_aggregated",
        "probability", "opt_days", "ml_days", "pess_days"]
VARIANTS = ["Yes", "yes", "YES", " YeS ", "yes "]


def _frame(flags: list[str]) -> pd.DataFrame:
    rows = [
        {"risk_id": f"R{i:03d}", "time_impact": flag, "mode_aggregated": "Series",
         "probability": 0.5, "opt_days": 1.0, "ml_days": 2.0, "pess_days": 3.0}
        for i, flag in enumerate(flags)
    ]
    rows.append({"risk_id": "R999", "time_impact": "No", "mode_aggregated": "Series",
                 "probability": 0.5, "opt_days": 1.0, "ml_days": 2.0, "pess_days": 3.0})
    return pd.DataFrame(rows, columns=COLS)


def test_simulating_ids_ignore_case():
    assert simulating_risk_ids(_frame(VARIANTS)) == simulating_risk_ids(_frame(["Yes"] * len(VARIANTS)))


def test_cross_check_counts_every_case_variant():
    mixed = compute_schedule_comparators(_frame(VARIANTS), GRID, day_rate=1000.0)
    canon = compute_schedule_comparators(_frame(["Yes"] * len(VARIANTS)), GRID, day_rate=1000.0)
    assert mixed == canon
    assert canon["comparator_1_emv_days_pert"] > 0.0  # the check could have failed


@pytest.mark.parametrize("raw", VARIANTS)
def test_reader_canonicalises_time_impact(raw):
    assert strawman_reader.canonical_time_impact(raw) == "Yes"


@pytest.mark.parametrize("raw, expected", [("No", "No"), (" no ", "No"), ("", ""), (None, ""), ("n/a", "n/a")])
def test_reader_canonical_leaves_non_yes_alone(raw, expected):
    assert strawman_reader.canonical_time_impact(raw) == expected


def _strawman(path, flags: list[str]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk_Scope"
    headers = ["Risk_ID", "Time_Impact", "Mode_Aggregated", "Probability", "Probability_Time",
               "Opt_Days", "ML_Days", "Pess_Days", "Pattern_ID", "Prefix_List", "Activity_Count"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c, value=h)
    prefix = active_risk_id_prefix()  # the reader drops rows without the binding's prefix
    for i, flag in enumerate(flags + ["No"]):
        vals = [f"{prefix}-{i:03d}", flag, "Series", 0.5, 0.5, 1, 2, 3, "", "", 1]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=6 + i, column=c, value=v)
    wb.save(str(path))


def test_read_risk_scope_emits_canonical_time_impact(tmp_path):
    mixed, canon = tmp_path / "mixed.xlsx", tmp_path / "canon.xlsx"
    _strawman(mixed, VARIANTS)
    _strawman(canon, ["Yes"] * len(VARIANTS))
    df_m = strawman_reader.read_risk_scope(str(mixed))
    df_c = strawman_reader.read_risk_scope(str(canon))
    assert list(df_m["time_impact"]) == list(df_c["time_impact"]) == ["Yes"] * len(VARIANTS) + ["No"]
    prefix = active_risk_id_prefix()
    assert simulating_risk_ids(df_m) == simulating_risk_ids(df_c) == [f"{prefix}-{i:03d}" for i in range(len(VARIANTS))]
