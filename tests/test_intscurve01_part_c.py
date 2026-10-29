"""INTSCURVE-01 Part C: the joint confidence level, in the report and the app.

P11  the JCL on cost (base + event-risk, delay cost excluded) against finish, at the D-3 target
     pair, with the 70% frontier; the basis states that joint sampling raises it.
P12  the period JSON carries the integrated curve points and the JCL (schema 1.5).
"""
from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

from src.integrated.jcl import JCL_FRONTIER_LEVEL, jcl_data, jcl_frontier, joint_confidence
from src.integrated.integrated_dollars import integrated_dollars_array
from src.final.sections.integrated_total import jcl_sentence, stage3_jcl

DAY_RATE = 35_000.0


def _paired(n: int = 4_000, seed: int = 11):
    """Cost and finish driven by one shared fire draw, as the Stage 3 run pairs them."""
    rng = np.random.default_rng(seed)
    fired = rng.random((n, 6)) < 0.35
    sev = rng.random((n, 6))
    cost = (fired * sev * 2.0e6).sum(axis=1)
    days = (fired * sev * 30.0).sum(axis=1)
    return cost, days, 400.0 + days


# ---------------------------------------------------------------------------
# The statistic
# ---------------------------------------------------------------------------

def test_joint_confidence_counts_iterations_meeting_both():
    cost = np.array([1.0, 2.0, 3.0, 4.0])
    finish = np.array([4.0, 3.0, 2.0, 1.0])
    assert joint_confidence(cost, finish, 2.0, 3.0) == 0.25
    assert joint_confidence(cost, finish, 4.0, 4.0) == 1.0
    with pytest.raises(ValueError):
        joint_confidence(cost, finish[:3], 1.0, 1.0)


def test_every_frontier_pair_holds_the_level_and_none_is_cheaper():
    cost, _, finish = _paired()
    frontier = jcl_frontier(cost, finish)
    assert frontier, "a frontier is drawn"
    need = np.ceil(JCL_FRONTIER_LEVEL * cost.size) / cost.size
    for f_t, c_t in frontier:
        assert joint_confidence(cost, finish, c_t, f_t) >= need
        below = np.nextafter(c_t, -np.inf)
        assert joint_confidence(cost, finish, below, f_t) < need
    costs = [c for _, c in frontier]
    assert costs == sorted(costs, reverse=True), "later finish, cheaper budget"


def test_shared_fire_decisions_raise_the_jcl_over_independence():
    """The basis sentence's claim, shown: shuffling one arm breaks the pairing and lowers it."""
    cost, _, finish = _paired()
    c_t, f_t = np.percentile(cost, 80), np.percentile(finish, 80)
    paired = joint_confidence(cost, finish, c_t, f_t)
    shuffled = joint_confidence(np.random.default_rng(3).permutation(cost), finish, c_t, f_t)
    assert paired > shuffled + 0.05


# ---------------------------------------------------------------------------
# The deliverable's JCL: cost without delay cost, D-3 targets
# ---------------------------------------------------------------------------

def _stage3(cost, days, finish):
    return {"cost_result": {"cost_dollars_array": cost}, "schedule_days_array": days,
            "schedule_result": SimpleNamespace(project_finish_days=finish)}


def test_the_cost_axis_excludes_delay_cost_and_the_targets_are_d3():
    cost, days, finish = _paired()
    base = np.full(cost.size, 50e6) + np.random.default_rng(1).normal(0, 2e6, cost.size)
    total = {"base_dollars": base, "base_total_dollars": 50e6}
    jcl = stage3_jcl(_stage3(cost, days, finish), total, DAY_RATE)
    assert np.array_equal(jcl["cost_axis"], cost + base), "no delay cost on the cost axis"
    headline = float(np.percentile(integrated_dollars_array(cost, days, DAY_RATE), 80))
    assert jcl["cost_target"] == 50e6 + headline
    assert jcl["finish_target"] == float(np.percentile(finish, 80))
    assert jcl["jcl"] == joint_confidence(cost + base, finish, jcl["cost_target"],
                                          jcl["finish_target"])
    assert jcl["level"] == 0.70 and jcl["frontier"]


def test_without_a_base_the_cost_axis_is_event_cost_alone():
    cost, days, finish = _paired()
    jcl = stage3_jcl(_stage3(cost, days, finish), None, DAY_RATE)
    assert np.array_equal(jcl["cost_axis"], cost) and jcl["base_drawn"] is False
    assert "base" not in jcl_sentence(jcl)
    assert stage3_jcl(_stage3(cost, days, finish[:-1]), None, DAY_RATE) is None
    assert stage3_jcl(None, None, DAY_RATE) is None


def test_the_jcl_section_renders_in_html_and_word(tmp_path, monkeypatch):
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")
    from docx import Document

    from src.final import commentary_text, composer
    from src.final.charts import make_jcl_scatter
    from src.final.html_blocks import blocks
    from src.final.word_generator import generate_word_summary
    from tests.test_sr9_v3_word_generator import _generate_v3_doc

    _, composed = _generate_v3_doc(tmp_path)
    assert composed["page2_distribution"]["jcl_path"] == "", "no Stage 3 arrays, no JCL"
    cost, days, finish = _paired()
    jcl = stage3_jcl(_stage3(cost, days, finish), None, DAY_RATE)
    png = make_jcl_scatter(jcl, title="JCL", output_path=tmp_path / "jcl.png")
    assert png.exists()
    composer._apply_jcl(composed, jcl, {"jcl_path": str(png)})

    p2 = composed["page2_distribution"]
    assert p2["jcl_basis"] == commentary_text.JCL_BASIS
    assert "joint" in p2["jcl_basis"].lower() and "higher" in p2["jcl_basis"]
    assert f"{jcl['jcl']:.0%}" in p2["jcl_sentence"]
    page = blocks(p2, section="page2_distribution")
    assert str(Path(png)) in {b[1] for b in page if b[0] == "image"}
    assert ("text", "#### " + p2["jcl_heading"]) in page

    out = tmp_path / "with_jcl.docx"
    generate_word_summary(
        summary_data=composed["executive_summary"],
        audit_trail_data=composed["appendix_audit_trail"],
        disclaimer_paragraphs=composed["disclaimer_paragraphs"],
        cross_validation=composed["cross_validation"], output_path=out, composed=composed)
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert p2["jcl_heading"] in headings
    text = "\n".join(p.text for p in doc.paragraphs)
    assert p2["jcl_sentence"] in text and p2["jcl_basis"] in text


# ---------------------------------------------------------------------------
# P12 - the period JSON follows the report (contract 1.5)
# ---------------------------------------------------------------------------

REPO = Path(__file__).resolve().parents[1]
WIND_01 = REPO / "status_periods" / "wind_period_01.xlsx"


@pytest.fixture(scope="module")
def period_01(tmp_path_factory: pytest.TempPathFactory) -> dict:
    import datetime as _dt
    import json

    from src.am.run_period import run_period
    from tests.input_skips import skip_missing

    if not WIND_01.is_file():
        skip_missing(WIND_01, detail="run scripts/build_wind_am_periods.py")
    path = run_period(
        workbook_path=WIND_01, project="synthetic_wind",
        output_dir=tmp_path_factory.mktemp("c_p01"), main_iterations=2000,
        days_at_risk_iterations=500, counterfactual_iterations=500, seed=42,
        criticality_measurement_mode="both", project_start=_dt.date(2024, 1, 1),
        allow_dirty=True)
    return json.loads(path.read_bytes().decode("utf-8"))


@pytest.mark.integration
class TestThePeriodJsonCarriesTheIntegratedCurveAndJcl:
    def test_the_curve_has_the_finish_curve_s_shape(self, period_01: dict):
        pts = period_01["integrated"]["curve_points"]
        finish = period_01["distribution"]["s_curve_points"]
        assert [p["cumulative_prob"] for p in pts] == [p["cumulative_prob"] for p in finish]
        dollars = [p["dollars"] for p in pts]
        assert dollars == sorted(dollars) and len(dollars) == 21

    def test_the_curve_agrees_with_the_block_s_own_percentiles(self, period_01: dict):
        b = period_01["integrated"]
        at = {p["cumulative_prob"]: p["dollars"] for p in b["curve_points"]}
        assert at[0.5] == b["p50_dollars"] and at[0.8] == b["p80_dollars"]
        assert at[0.9] == b["p90_dollars"] and at[0.95] == b["p95_dollars"]

    def test_the_jcl_is_at_the_d3_pair_and_below_either_confidence(self, period_01: dict):
        b = period_01["integrated"]
        j = b["jcl"]
        assert j["cost_target_dollars"] == b["p80_dollars"]
        finish_p80 = {p["cumulative_prob"]: p["date"]
                      for p in period_01["distribution"]["s_curve_points"]}[0.8]
        assert j["finish_target_date"] == finish_p80
        assert 0 < j["jcl"] <= min(j["cost_confidence"], j["finish_confidence"])
        assert j["frontier_level"] == 0.70
        assert "delay cost excluded" in j["basis"]


def test_the_period_jcl_leaves_delay_cost_off_its_cost_axis():
    """All the money is delay: with delay off the cost axis every iteration meets the budget."""
    import datetime as _dt

    from src.am.json_writer import _build_jcl

    days = np.linspace(0.0, 50.0, 1_000)
    result = SimpleNamespace(
        cost_dollars_array=np.zeros(days.size), schedule_days_array=days,
        integrated_dollars_array=integrated_dollars_array(np.zeros(days.size), days, DAY_RATE),
        deterministic_baseline_days=300.0)
    j = _build_jcl(result, _dt.date(2024, 1, 1))
    assert j["cost_confidence"] == 1.0
    assert j["jcl"] == j["finish_confidence"]
