"""CONFIGTRUTH-01 P5 (register item 5): the register's seed row does not promise what no route does.

The MC_Params row was labelled "Random seed (0=random)", and only the legacy cost-MC reader
(``read_params_v5``) reads it; the deliverable's seed is ``SR9_SEED``. The row is relabelled, not
dropped (task -2, ruling B): dropping it would shift every row beneath it. The reader still accepts
the old label, so an existing register keeps its seed. The Settings ``Set_seed`` cell is not touched.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parents[1]
FIXTURES = REPO / "tests" / "fixtures"
#: Literal on purpose: the red run on the tree before P5 must fail on the workbook's label, not on
#: an import of a constant that does not exist yet.
NEW = "Random seed - legacy cost MC only (0=random; the deliverable's seed is SR9_SEED)"
OLD = "Random seed (0=random)"
REGISTERS = [FIXTURES / "synthetic_wind" / "wind_register.xlsx",
             FIXTURES / "synthetic_road" / "road_register.xlsx"]


def _seed_row_labels(wb) -> list[str]:
    return [str(row[1].value) for row in wb["MC_Params"].iter_rows()
            if row[1].value is not None and "seed" in str(row[1].value).lower()]


def test_the_reader_names_the_same_labels():
    from src.v5_register_reader import SEED_LABEL, SEED_LABEL_BEFORE_CONFIGTRUTH01

    assert (SEED_LABEL, SEED_LABEL_BEFORE_CONFIGTRUTH01) == (NEW, OLD)


@pytest.mark.parametrize("register", REGISTERS, ids=lambda p: p.name)
def test_each_fixture_register_carries_the_new_label_and_not_the_old(register):
    wb = openpyxl.load_workbook(register, read_only=True)
    try:
        assert _seed_row_labels(wb) == [NEW]
    finally:
        wb.close()


def test_the_scaffold_writes_the_new_label():
    from tools.register_scaffold import build_mc_params_sheet

    wb = openpyxl.Workbook()
    build_mc_params_sheet(wb)
    assert _seed_row_labels(wb) == [NEW]


@pytest.mark.parametrize("label", ["new", "old"])
def test_the_legacy_reader_reads_the_seed_under_either_label(label, tmp_path, monkeypatch):
    import src.sra.cpm.bindings.synthetic_wind  # noqa: F401  self-registers
    from src.v5_register_reader import read_params_v5

    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(REGISTERS[0], register)
    wb = openpyxl.load_workbook(register)
    for row in wb["MC_Params"].iter_rows():
        if row[1].value is not None and "seed" in str(row[1].value).lower():
            row[1].value = NEW if label == "new" else OLD
            row[2].value = 7
            break
    wb.save(register)
    wb.close()
    config = read_params_v5(register)[0]
    assert config.seed == 7
