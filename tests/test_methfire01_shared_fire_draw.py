"""METHFIRE-01: one risk, one fire decision, on every Stage 3 route (SPEC R-10).

Brief: METHFIRE-01. The properties each
test holds are named in its docstring. The fast tests use synthetic graphs, drivers and risks;
the tests that run a shipped project through an engine are `slow`.
"""
from __future__ import annotations

import os
import shutil
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import openpyxl
import pytest

from src.qra_core import draw_fire_matrix
from src.sra.cpm.cpm_simulation import ScheduleMCConfig, run_schedule_mc
from src.sra.cpm.risk_driver_schema import RiskDriver
from tests.conftest import TEST_MODE_MULTIPLIERS
from tests.sra.cpm.test_cpm_simulation import _driver, _five_activity_chain

REPO = Path(__file__).resolve().parents[1]
ROAD_REGISTER = REPO / "tests" / "fixtures" / "synthetic_road" / "road_register.xlsx"
ROAD_XER = REPO / "tests" / "fixtures" / "synthetic_road" / "road_schedule.xer"
WIND_REGISTER = REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"
WIND_XER = REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_schedule.xer"


def _risk(rid: str, probability: float, group: str) -> SimpleNamespace:
    return SimpleNamespace(risk_id=rid, probability=probability, correlation_group=group)


def _three_drivers() -> dict[str, RiskDriver]:
    return {
        "R001": _driver("R001", affected=("a1",), probability=0.4),
        "R002": _driver("R002", affected=("a3",), probability=0.6),
        "R003": _driver("R003", affected=("a5",), probability=0.2),
    }


# ---------------------------------------------------------------------------
# The draw itself
# ---------------------------------------------------------------------------

class TestDrawFireMatrix:

    def test_it_fires_on_probability_and_is_seeded(self) -> None:
        risks = [_risk("R001", 0.3, "CG01"), _risk("R002", 0.7, "CG01"), _risk("R003", 0.5, "S1")]
        a = draw_fire_matrix(20_000, risks, {"CG01": 0.6}, np.random.default_rng(3))
        b = draw_fire_matrix(20_000, risks, {"CG01": 0.6}, np.random.default_rng(3))
        assert a.dtype == np.bool_ and a.shape == (20_000, 3)
        assert np.array_equal(a, b)
        assert np.allclose(a.mean(axis=0), [0.3, 0.7, 0.5], atol=0.02)

    def test_group_members_fire_together_more_than_independence_implies(self) -> None:
        """The correlation D-1 carries to the schedule arm is real in the draw."""
        risks = [_risk("R001", 0.5, "CG01"), _risk("R002", 0.5, "CG01")]
        fired = draw_fire_matrix(50_000, risks, {"CG01": 0.8}, np.random.default_rng(11))
        both = float((fired[:, 0] & fired[:, 1]).mean())
        assert both > 0.25 + 0.1


# ---------------------------------------------------------------------------
# The kernel consumes a supplied mask
# ---------------------------------------------------------------------------

class TestKernelFireMask:

    def _cfg(self, batch_size: int = 100) -> ScheduleMCConfig:
        return ScheduleMCConfig(n_iter=400, batch_size=batch_size, seed=5,
                                mode_multipliers=TEST_MODE_MULTIPLIERS)

    def test_the_kernel_fires_exactly_on_the_supplied_mask(self) -> None:
        """P-1 at the kernel: the firing record is the mask, not a draw of its own."""
        mask = np.random.default_rng(1).random((400, 3)) < 0.5
        result = run_schedule_mc(graph=_five_activity_chain(), risk_drivers=_three_drivers(),
                                 config=self._cfg(), fire_mask=mask)
        assert np.array_equal(result.risk_fired_array, mask)

    @pytest.mark.parametrize("batch_size", [100, 400, 137])
    @pytest.mark.parametrize("mode", ["mode_conditioned", "both"])
    def test_given_its_own_draw_the_kernel_reproduces_itself(self, batch_size: int, mode: str) -> None:
        """P-3: a supplied mask equal to the kernel's own draw changes nothing, at any batch size
        and on both passes of the "both" measurement mode. So a moved number is the shared draw."""
        graph, drivers = _five_activity_chain(), _three_drivers()
        sev = np.random.default_rng(7).random((400, 3)).astype(np.float32)
        own = run_schedule_mc(graph=graph, risk_drivers=drivers, config=self._cfg(batch_size),
                              severity_uniforms=sev, rng=np.random.default_rng(9),
                              criticality_measurement_mode=mode)
        supplied = run_schedule_mc(graph=graph, risk_drivers=drivers, config=self._cfg(batch_size),
                                   severity_uniforms=sev, rng=np.random.default_rng(9),
                                   criticality_measurement_mode=mode,
                                   fire_mask=np.asarray(own.risk_fired_array))
        assert np.array_equal(own.project_finish_hours, supplied.project_finish_hours)
        assert own.criticality_index == supplied.criticality_index
        assert own.mode_independent_criticality_index == supplied.mode_independent_criticality_index

    @pytest.mark.parametrize("bad", [np.zeros((400, 2), dtype=bool), np.zeros((400, 3), dtype=float)])
    def test_a_mask_of_the_wrong_shape_or_type_is_refused(self, bad: np.ndarray) -> None:
        with pytest.raises(ValueError, match="fire_mask"):
            run_schedule_mc(graph=_five_activity_chain(), risk_drivers=_three_drivers(),
                            config=self._cfg(), fire_mask=bad)


# ---------------------------------------------------------------------------
# The Stage 3 orchestrator projects the cost arm's record
# ---------------------------------------------------------------------------

class TestCostArmFireMask:

    def test_it_projects_the_cost_record_onto_driver_order(self) -> None:
        from src.integrated.integrated_simulation_stage3 import cost_arm_fire_mask
        fired = np.random.default_rng(2).random((50, 4)) < 0.5
        mask = cost_arm_fire_mask({"fired_array": fired, "risk_ids_all": ["A", "B", "C", "D"]},
                                  ("C", "A"))
        assert np.array_equal(mask, fired[:, [2, 0]])

    def test_a_driver_the_cost_arm_does_not_price_is_refused(self) -> None:
        from src.integrated.integrated_simulation_stage3 import cost_arm_fire_mask
        with pytest.raises(ValueError, match="no risk in the cost arm"):
            cost_arm_fire_mask({"fired_array": np.zeros((5, 1), dtype=bool), "risk_ids_all": ["A"]},
                               ("A", "Z"))


# ---------------------------------------------------------------------------
# The period route refuses what it cannot share exactly
# ---------------------------------------------------------------------------

def test_the_period_route_refuses_a_chain_target_among_its_drivers(monkeypatch) -> None:
    """P-1 on the period route: its reproduced draw is pre-chain, so a chained driver is refused
    rather than fired on a draw that is only nearly the cost arm's."""
    import src.defaults as defaults
    from src.am.conditional_runner import period_fire_mask
    monkeypatch.setattr(defaults, "get_chain_risk_ids",
                        lambda: {"c1_trigger": "X", "c1_target": "R002", "c2_trigger": None, "c2_targets": []})
    with pytest.raises(ValueError, match="event-chain targets"):
        period_fire_mask(risks=[], pinned_states=[], schedule_risk_ids=("R001", "R002"),
                         register_path=ROAD_REGISTER, n_iter=10, seed=1)


# ---------------------------------------------------------------------------
# D-2: the reader refuses a register that gives one risk two probabilities
# ---------------------------------------------------------------------------

def _road_copy(tmp_path: Path, mutate) -> Path:
    dst = tmp_path / "road_register.xlsx"
    wb = openpyxl.load_workbook(ROAD_REGISTER)
    ws = wb["Risk Register"]
    header = next(r for r in ws.iter_rows() if any(c.value == "Probability (Time)" for c in r))
    col = {c.value: c.column for c in header}
    mutate(ws, col, header[0].row + 1)
    wb.save(dst)
    return dst


class TestUnequalProbabilitiesRefused:

    def test_the_shipped_road_register_reads(self, monkeypatch) -> None:
        from src.sra.strawman_reader import read_risk_scope
        monkeypatch.setenv("QRA_PROJECT", "synthetic_road")
        scope = read_risk_scope(ROAD_REGISTER)
        yes = scope[scope["time_impact"] == "Yes"]
        assert len(yes) == 10 and (yes["probability"] > 0).all()

    def test_a_time_yes_row_with_a_different_probability_time_is_refused_by_name(
            self, tmp_path, monkeypatch) -> None:
        """P-6, red first: KH-002's Probability (Time) moved away from its Probability."""
        from src.sra.strawman_reader import read_risk_scope
        monkeypatch.setenv("QRA_PROJECT", "synthetic_road")

        def mutate(ws, col, start):
            for r in range(start, ws.max_row + 1):
                if ws.cell(r, col["Risk ID"]).value == "KH-002":
                    ws.cell(r, col["Probability (Time)"]).value = 0.99
        with pytest.raises(ValueError, match=r"KH-002 \(Probability=0\.45, Probability \(Time\)=0\.99\)"):
            read_risk_scope(_road_copy(tmp_path, mutate))

    def test_a_cost_only_row_may_carry_a_different_probability_time(self, tmp_path, monkeypatch) -> None:
        """The refusal is about one risk given two probabilities on the schedule side, nothing wider."""
        from src.sra.strawman_reader import read_risk_scope
        monkeypatch.setenv("QRA_PROJECT", "synthetic_road")

        def mutate(ws, col, start):
            for r in range(start, ws.max_row + 1):
                if str(ws.cell(r, col["Time Impact"]).value).strip() != "Yes" and ws.cell(r, col["Risk ID"]).value:
                    ws.cell(r, col["Probability (Time)"]).value = 0.77
        read_risk_scope(_road_copy(tmp_path, mutate))

    def test_probability_is_matched_exactly_never_by_prefix(self, tmp_path, monkeypatch) -> None:
        """A strawman carrying only Probability_Time must not have it read as Probability - that
        would make the refusal compare the column with itself."""
        from src.sra.strawman_reader import read_risk_scope
        monkeypatch.setenv("QRA_PROJECT", "synthetic_road")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Risk_Scope"
        ws.append(["Risk_ID", "Time_Impact", "Mode_Aggregated", "Probability_Time",
                   "Opt_Days", "ML_Days", "Pess_Days"])
        ws.append(["KH-001", "Yes", "Series", 0.4, 1, 2, 3])
        path = tmp_path / "strawman.xlsx"
        wb.save(path)
        with pytest.raises(ValueError, match="Required column 'probability' missing"):
            read_risk_scope(path)


# ---------------------------------------------------------------------------
# Shipped projects, through the engines
# ---------------------------------------------------------------------------

def _stage3(project: str, register: Path, xer: Path, n_iter: int, seed: int):
    from src.integrated.integrated_simulation_stage3 import run_integrated_simulation_stage3
    from src.integrated.stage3_config import build_stage3_config, build_stage3_rules_from_scope
    from src.sra.cpm.auto_mapper import expand_auto_rules
    from src.integrated.stage3_inputs import Stage3BindingConfig
    from src.sra.cpm.graph_builder import build_schedule_graph
    from src.sra.cpm.xer_reader import parse_xer
    from src.sra.strawman_reader import read_risk_scope
    from src.v5_register_reader import read_risks_v5

    os.environ["QRA_PROJECT"] = project
    xer_tables = parse_xer(str(xer))
    graph = build_schedule_graph(xer_tables)
    scope = read_risk_scope(str(register))
    drivers = expand_auto_rules(build_stage3_rules_from_scope(
        scope, graph, Stage3BindingConfig(register_path=register, xer_path=xer,
                                          strawman_path=register,
                                          float_threshold_hours=480.0)), graph)
    cfg = build_stage3_config(register_path=register, xer_path=xer, strawman_path=register,
                              iterations=n_iter, seed=seed, caller_module="tests/methfire01")
    return run_integrated_simulation_stage3(cfg, risks=read_risks_v5(str(register)),
                                            risk_scope_df=scope, schedule_graph=graph,
                                            risk_drivers=drivers, xer_tables=xer_tables)


def _columns(cost_ids: list[str], ids) -> list[int]:
    return [list(cost_ids).index(r) for r in ids]


@pytest.mark.slow
@pytest.mark.parametrize("project,register,xer", [
    ("synthetic_road", ROAD_REGISTER, ROAD_XER), ("synthetic_wind", WIND_REGISTER, WIND_XER)])
class TestShippedProjects:

    def test_the_headline_fires_each_risk_once_for_both_arms(self, project, register, xer, monkeypatch) -> None:
        """P-1 on the Stage 3 headline, and a dual-coded risk that fires contributes on both arms."""
        monkeypatch.setenv("QRA_PROJECT", project)
        res = _stage3(project, register, xer, n_iter=2_000, seed=42)
        cost = res.cost_result
        sched_ids = list(res.schedule_result.risk_ids)
        cost_fired = np.asarray(cost["fired_array"])[:, _columns(cost["risk_ids_all"], sched_ids)]
        assert np.array_equal(np.asarray(res.schedule_result.risk_fired_array), cost_fired)
        per_risk_cost = np.asarray(cost["per_risk_cost_array"])[:, _columns(cost["risk_ids_all"], sched_ids)]
        assert (per_risk_cost[~cost_fired] == 0).all()
        assert per_risk_cost[cost_fired].mean() > 0

    def test_same_seed_same_headline(self, project, register, xer, monkeypatch) -> None:
        """P-10."""
        monkeypatch.setenv("QRA_PROJECT", project)
        a = _stage3(project, register, xer, n_iter=1_000, seed=7)
        b = _stage3(project, register, xer, n_iter=1_000, seed=7)
        assert np.array_equal(a.schedule_days_array, b.schedule_days_array)
        assert np.array_equal(a.cost_result["cost_dollars_array"], b.cost_result["cost_dollars_array"])

    def test_the_period_route_fires_as_its_own_cost_arm_would(self, project, register, xer, monkeypatch) -> None:
        """P-4: the period schedule run fires on exactly the draw the period cost arm makes, and a
        Closed risk fires on neither."""
        from src.am.conditional_runner import ConditionalRunInputs, run_conditional_mc
        from src.am.integrated_period import apply_pinning_to_risks
        from src.am.workbook_reader import WorkbookData
        from src.integrated.integrated_simulation import run_integrated_simulation
        from src.integrated.stage3_config import build_stage3_config
        from src.sra.strawman_reader import read_risk_scope
        from src.v5_register_reader import read_risks_v5

        monkeypatch.setenv("QRA_PROJECT", project)
        n_iter, seed = 1_000, 42
        scope = read_risk_scope(str(register))
        closed = str(scope[scope["time_impact"] == "Yes"]["risk_id"].iloc[0])
        result = run_conditional_mc(ConditionalRunInputs(
            workbook_data=WorkbookData(period_info={}, milestones=[], mitigations=[], am_risk_ids=[],
                                       risk_status=[{"risk_id": closed, "status_this_period": "Closed"}]),
            register_path=register, xer_path=xer, strawman_path=register,
            n_iter=n_iter, seed=seed, criticality_measurement_mode="mode_conditioned", project=project))
        cfg = build_stage3_config(register_path=register, xer_path=xer, strawman_path=register,
                                  iterations=n_iter, seed=seed, caller_module="tests/methfire01")
        priced, _ = apply_pinning_to_risks(read_risks_v5(str(register)), result.pinned_states)
        cost = run_integrated_simulation(config=cfg.integrated_config, risks=priced, risk_scope_df=scope,
                                         precomputed_severity_uniforms=result.severity_uniforms_by_risk_id)
        sched = result.schedule_result
        cost_fired = np.asarray(cost["fired_array"])[:, _columns(cost["risk_ids_all"], sched.risk_ids)]
        assert np.array_equal(np.asarray(sched.risk_fired_array), cost_fired)
        closed_col = list(sched.risk_ids).index(closed)
        assert not np.asarray(sched.risk_fired_array)[:, closed_col].any()


def test_a_schedule_only_risk_without_probability_is_told_how_to_fill_it(tmp_path, monkeypatch) -> None:
    """Director 17/09/2026, ruling (a): a risk that only delays still has its one probability in
    `Probability`; the refusal says so rather than only that the cell is blank."""
    from src.sra.strawman_reader import read_risk_scope
    monkeypatch.setenv("QRA_PROJECT", "synthetic_road")

    def mutate(ws, col, start):
        for r in range(start, ws.max_row + 1):
            if ws.cell(r, col["Risk ID"]).value == "KH-002":
                ws.cell(r, col["Probability"]).value = None
    with pytest.raises(ValueError, match="including a risk that only delays"):
        read_risk_scope(_road_copy(tmp_path, mutate))
