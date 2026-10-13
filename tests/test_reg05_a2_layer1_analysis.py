"""REG05-A2 - Layer-1 + the analysis sheet promoted into the per-project scaffolder.

Gates (build brief REG05-A2 §5-6, §8.4):
- the empty template emits the four Layer1_* sheets + Grid_Layer1_* + the
  Set_Layer1_Enabled off-signal + the EMV & Mode Analysis sheet;
- Set_Layer1_Enabled ships FALSE, so an empty, present Layer-1 reads to NIL impact
  (the flag skips the parser - the empty-present crash of brief §8.1 never fires),
  and toggling the flag off on a POPULATED register makes the cost read byte-
  identical to the Layer-1-off (shared-pool) path;
- flag TRUE on the wind fixture keeps its filled Layer-1 reading ON;
- the analysis sheet is GENERIC: title driven from Set_ProjectName, totals/helper
  tables/charts reference the register EMV columns over the data block (no wind row
  count or risk IDs baked in);
- contract is 1.4 and the headless Validate passes on the fresh template.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.register_contract import (
    CONTRACT_VERSION, LAYER1_NAMED_RANGES, REQUIRED_NAMED_RANGES,
    check_contract, layer1_enabled_flag,
)
from src.tier1_headless import (
    _layer1_from_grids, extract_grids_openpyxl, run_headless,
)
from tools.register_scaffold import apply_finish, build_template

REPO = Path(__file__).resolve().parents[1]
FIXTURE = REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"
_LAYER1_SHEETS = ("Layer1_Base", "Layer1_Factors", "Layer1_Allocation",
                  "Layer1_Correlation")


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    return build_template(tmp_path_factory.mktemp("reg05a2") / "t.xlsx", prefix="RSK")


# --- 1. the template now emits Layer-1 + analysis (present-but-off) -----------

def test_template_emits_layer1_sheets_grids_flag_and_analysis(template):
    wb = openpyxl.load_workbook(template)
    try:
        for sh in (*_LAYER1_SHEETS, "EMV & Mode Analysis"):
            assert sh in wb.sheetnames, sh
        names = set(wb.defined_names)
        assert LAYER1_NAMED_RANGES <= names          # 4 grids + the flag
        assert wb["Settings"]["B8"].value is False    # ships OFF
    finally:
        wb.close()


def test_contract_is_1_4_and_validate_passes(template):
    assert CONTRACT_VERSION == "1.4"
    assert LAYER1_NAMED_RANGES <= REQUIRED_NAMED_RANGES
    wb = openpyxl.load_workbook(template)
    try:
        present = list(wb.defined_names)
        version = wb["Settings"]["B6"].value
    finally:
        wb.close()
    grids = extract_grids_openpyxl(str(template))
    ok, msg = check_contract(present, grids["methodology"], version)
    assert ok, msg
    assert "v1.4" in msg


# --- 2. the flag governs the read (nil impact when off) -----------------------

def test_flag_false_reads_nil_and_never_crashes(template):
    """Empty, present Layer-1 + flag FALSE -> layer1 skipped, nil P-table, no crash
    (the brief §8.1 empty-present parser crash is gated out by the flag)."""
    grids = extract_grids_openpyxl(str(template))
    assert grids["layer1_enabled"] is False
    assert len(grids["layer1"]) == 4                 # present-but-empty
    out = run_headless(grids, risk_id_prefix="RSK", n_iter=2000, seed=42, view="owner")
    assert out["layer1_enabled"] is False
    assert all(v == 0.0 for _p, v in out["p_table"])


def test_flag_toggle_is_nil_impact_on_a_populated_register():
    """On the wind register (Layer-1 populated), forcing the flag OFF makes the
    owner-view cost read byte-identical to the Layer-1-off (shared-pool) path, and
    ON genuinely changes it - so the flag is the on/off with nil impact when off."""
    grids = extract_grids_openpyxl(str(FIXTURE))
    assert grids["layer1_enabled"] is True           # fixture ships ON

    on = run_headless(grids, "WT", n_iter=2000, seed=42, view="owner")
    off = run_headless({**grids, "layer1_enabled": False}, "WT", 2000, 42, "owner")
    shared = run_headless(grids, "WT", n_iter=2000, seed=42, view="shared-pool")

    assert on["layer1_enabled"] is True
    assert off["layer1_enabled"] is False
    # flag OFF (owner) == Layer-1-off (shared-pool): identical cost, bit-for-bit.
    assert off["p_table"] == shared["p_table"]
    # and ON actually moved the owner cost (the flag is load-bearing).
    assert on["p_table"] != off["p_table"]


def test_flag_true_with_no_layer1_sheets_returns_none():
    """flag TRUE but no Layer-1 grids (a client register) -> None, no crash."""
    assert _layer1_from_grids({"layer1_enabled": True, "layer1": {}}) is None
    assert _layer1_from_grids({"layer1_enabled": False, "layer1": {}}) is None


def test_layer1_enabled_flag_coercion():
    for truthy in (True, 1, "TRUE", "true", "Yes", "1", "on"):
        assert layer1_enabled_flag(truthy) is True, truthy
    for falsy in (False, 0, "FALSE", "no", "", None, "maybe"):
        assert layer1_enabled_flag(falsy) is False, falsy


# --- 3. the fixture keeps Layer-1 ON (flag TRUE, byte-identical) --------------

def test_fixture_flag_true_layer1_reads_on():
    grids = extract_grids_openpyxl(str(FIXTURE))
    assert grids["layer1_enabled"] is True
    layer1 = _layer1_from_grids(grids)
    assert layer1 is not None
    assert layer1.declared_total == 50_000_000.0     # the wind $50M base
    assert len(layer1.factors) == 14


# --- 4. the analysis sheet is generic (title + register-referenced) -----------

def test_analysis_title_driven_from_project_name(template, tmp_path):
    # template (project=None) -> generic placeholder title.
    wb = openpyxl.load_workbook(template)
    try:
        assert wb["EMV & Mode Analysis"]["A1"].value == "EMV & Mode Analysis - this project"
    finally:
        wb.close()
    # a named project drives the title (retiring the hardcoded wind literal).
    wb = openpyxl.load_workbook(template)
    try:
        apply_finish(wb, project={"Set_ProjectName": "Acme Solar BOP"})
        assert wb["EMV & Mode Analysis"]["A1"].value == "EMV & Mode Analysis - Acme Solar BOP"
    finally:
        wb.close()


def test_analysis_references_register_ranges_not_hardcoded():
    """The analysis formulas reference the Risk Register EMV columns over the data
    block - no wind row count / risk IDs baked in. Two templates with different
    block sizes get correspondingly different reference extents."""
    from tools.register_scaffold import DATA_START_ROW
    import tempfile

    def analysis(n_rows: int):
        p = Path(tempfile.mkdtemp()) / "t.xlsx"
        build_template(p, n_blank_input_rows=n_rows)
        wb = openpyxl.load_workbook(p)
        try:
            ws = wb["EMV & Mode Analysis"]
            return ws["B5"].value, ws["A12"].value, ws["B12"].value
        finally:
            wb.close()

    b5_40, a12_40, b12_40 = analysis(40)
    b5_20, _a12_20, _b12_20 = analysis(20)

    rr = "'Risk Register'"
    assert b5_40.startswith(f"=SUM({rr}!AB")           # cost total over EMV_cost col
    assert a12_40 == f"={rr}!A5"                        # helper Risk ID from register
    assert b12_40 == f"={rr}!AB5"                       # helper EMV_cost from register
    # the reference extent tracks the block size (generic, not a fixed 35/wind).
    assert b5_40.endswith("AB44)")                      # 5 + 40 - 1
    assert b5_20.endswith("AB24)")                      # 5 + 20 - 1
