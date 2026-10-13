"""REGUX-01 A-0.4: the Layer-1 grids ship at their real shape, and a Layer 1 switched on before it
is filled is refused with a message that names the problem.

As shipped before this, `Layer1_Allocation` was one cell and the TOC gave a three-step procedure to
fill grids that were not there. Characterised on 08/09 before building: seeding the grids with zeros
parses SILENTLY to a Layer 1 that is on and contributes nothing; seeding blanks was refused with a
raw `TypeError` from `float()`. So the shape and the refusal are one change, not two - blank grids
that the parser refuses by name.

What this suite can establish is the shape and the parser. Whether Excel lets a person type into
the blue cells and refuses the grey ones is the Excel pass, REGUX-01 §6a, run through
`scripts/ask_excel_register.py`.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pytest

from src.qra_register_parse import RegisterReadError
from src.tier1_headless import _layer1_from_grids, extract_grids_openpyxl
from tools.register_scaffold import (
    LAYER1_GENERIC_COLS, _N_LAYER1_FACTORS, _N_LAYER1_WBS_ROWS, build_template,
)

INPUT_BLUE = "000000CC"
LAYER1 = ("Layer1_Base", "Layer1_Factors", "Layer1_Allocation", "Layer1_Correlation")
WBS = tuple(f"WBS-{i:02d}" for i in range(1, _N_LAYER1_WBS_ROWS + 1))
CORR_HEADER_ROW = 4


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    return build_template(tmp_path_factory.mktemp("regux01") / "t.xlsx")


@pytest.fixture(scope="module")
def wb(template):
    book = openpyxl.load_workbook(template)
    yield book
    book.close()


def _blue(cell) -> bool:
    colour = cell.font.color.rgb if (cell.font and cell.font.color and cell.font.color.rgb) else None
    return str(colour) == INPUT_BLUE


def _locked(cell) -> bool:
    return True if cell.protection is None else bool(cell.protection.locked)


class TestTheGridsHaveTheirRealShape:
    def test_base_is_wbs_rows_by_resource_columns_and_blank(self, wb):
        """WBS rows x the eight generic resource columns both real projects adopted unchanged, in
        $M, every value blank - a new project cannot inherit a number, and blank is what the parser
        refuses (a zero would parse)."""
        ws = wb["Layer1_Base"]
        assert [c.value for c in ws[1]] == ["WBS Row", *LAYER1_GENERIC_COLS]
        assert tuple(ws.cell(row=r, column=1).value for r in range(2, 2 + len(WBS))) == WBS
        for r in range(2, 2 + len(WBS)):
            for c in range(2, 2 + len(LAYER1_GENERIC_COLS)):
                cell = ws.cell(row=r, column=c)
                assert cell.value is None, cell.coordinate
                assert _blue(cell) and not _locked(cell), cell.coordinate
        labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert "Declared_Total ($M)" in labels and "Units" in labels
        declared = ws.cell(row=labels.index("Declared_Total ($M)") + 1, column=2)
        assert declared.value is None and _blue(declared) and not _locked(declared)

    def test_allocation_columns_are_the_base_rows_by_the_same_names(self, wb):
        """The parser matches Allocation columns to Base rows BY NAME, so the two must ship
        identical, and a rename must be made in both places. That coupling is the user's; the
        template arrives consistent and the validator catches a drift loudly."""
        ws = wb["Layer1_Allocation"]
        assert [c.value for c in ws[1]] == ["Factor", *WBS]
        assert ws.max_row == 1 + _N_LAYER1_FACTORS
        for r in range(2, ws.max_row + 1):
            for c in range(1, 2 + len(WBS)):
                cell = ws.cell(row=r, column=c)
                assert cell.value is None and _blue(cell) and not _locked(cell), cell.coordinate
        # the header labels are inputs too - a project renames its WBS rows
        for c in range(2, 2 + len(WBS)):
            assert _blue(ws.cell(row=1, column=c)) and not _locked(ws.cell(row=1, column=c))
        assert _locked(ws["A1"]) and not _blue(ws["A1"])

    def test_factors_is_the_fifth_table(self, wb):
        """§5.4 ruled five Tables and part 2b built four; this is the fifth, and it was left for
        here because its shape is this part's to change."""
        ws = wb["Layer1_Factors"]
        assert list(ws.tables) == ["tblLayer1_Factors"]
        assert ws.tables["tblLayer1_Factors"].ref == f"A1:F{1 + _N_LAYER1_FACTORS}"
        assert ws.protection.sheet is True
        assert ws.protection.sort is False and ws.protection.autoFilter is False  # permitted
        assert ws.freeze_panes == "A2"
        for r in range(2, 2 + _N_LAYER1_FACTORS):
            for c in range(1, 7):
                cell = ws.cell(row=r, column=c)
                assert cell.value is None and _blue(cell) and not _locked(cell), cell.coordinate

    def test_correlation_is_a_lower_triangle_with_the_rest_visibly_not_read(self, wb):
        """Names across the header and down column A, coefficients strictly below the diagonal.
        The diagonal and the upper triangle are locked and grey: the parser never reads them, and
        a cell the parser never reads must not accept typing."""
        ws = wb["Layer1_Correlation"]
        k = _N_LAYER1_FACTORS
        assert ws.cell(row=CORR_HEADER_ROW, column=1).value == "Factor"
        for j in range(k):
            hdr = ws.cell(row=CORR_HEADER_ROW, column=2 + j)
            assert hdr.value is None and _blue(hdr) and not _locked(hdr), hdr.coordinate
        for i in range(k):
            name = ws.cell(row=CORR_HEADER_ROW + 1 + i, column=1)
            assert name.value is None and _blue(name) and not _locked(name), name.coordinate
            for j in range(k):
                cell = ws.cell(row=CORR_HEADER_ROW + 1 + i, column=2 + j)
                assert cell.value is None, cell.coordinate
                if j < i:
                    assert _blue(cell) and not _locked(cell), cell.coordinate
                else:
                    assert _locked(cell) and not _blue(cell), cell.coordinate
        assert not _locked(ws["B1"])  # Corr_Mode toggle, unchanged

    @pytest.mark.parametrize("sheet", LAYER1)
    def test_unlocked_if_and_only_if_blue(self, sheet, wb):
        """A-0.1's invariant, held as a biconditional on the four Layer-1 sheets: these are unlocked
        BY STYLE at finalise time, not by a range table, so a lower triangle can be unlocked without
        unlocking the rectangle around it."""
        ws = wb[sheet]
        wrong = [c.coordinate for c in ws._cells.values() if _blue(c) == _locked(c)]
        assert not wrong, f"{sheet}: unlocked xor blue at {wrong[:12]}"

    def test_grid_names_cover_the_grids(self, wb):
        n_c, n_r, k = len(LAYER1_GENERIC_COLS), len(WBS), _N_LAYER1_FACTORS
        want = {
            "Grid_Layer1_Base": f"$A$1:${openpyxl.utils.get_column_letter(1 + n_c)}${1 + n_r + 3}",
            "Grid_Layer1_Factors": f"$A$1:$F${1 + k}",
            "Grid_Layer1_Allocation": f"$A$1:${openpyxl.utils.get_column_letter(1 + n_r)}${1 + k}",
            "Grid_Layer1_Correlation": f"$A$1:${openpyxl.utils.get_column_letter(1 + k)}${CORR_HEADER_ROW + k}",
        }
        for name, ref in want.items():
            (sheet, coord), = wb.defined_names[name].destinations
            assert coord == ref, f"{name} = {sheet}!{coord}, expected {ref}"


# --- the parser --------------------------------------------------------------------------

def _grids_after(template: Path, tmp_path: Path, mutate=None) -> dict:
    """Copy the template, let ``mutate(wb)`` type into it as a user would, switch Layer 1 ON, and
    return the headless grids - the exact input the read gate sees."""
    wb = openpyxl.load_workbook(template)
    if mutate is not None:
        mutate(wb)
    wb["Settings"]["B8"] = True  # Set_Layer1_Enabled
    p = tmp_path / "filled.xlsx"
    wb.save(p)
    wb.close()
    grids = extract_grids_openpyxl(str(p))
    assert grids["layer1_enabled"] is True
    return grids


def _minimal_valid_fill(wb) -> None:
    """Two WBS rows renamed IN BOTH PLACES, numbers typed, two factors, ticks, one coefficient -
    and every other placeholder left exactly as shipped."""
    b = wb["Layer1_Base"]
    b["A2"], b["A3"] = "Civils", "Electrical"
    b["B2"], b["C2"] = 10.0, 5.0          # Staff, Labour
    b["B3"], b["C3"] = 8.0, 2.0
    labels = [b.cell(row=r, column=1).value for r in range(1, b.max_row + 1)]
    b.cell(row=labels.index("Declared_Total ($M)") + 1, column=2, value=25.0)

    f = wb["Layer1_Factors"]
    f["A2"], f["B2"], f["C2"], f["D2"], f["E2"], f["F2"] = \
        "Labour rate", "labour_rate", -0.05, 0.0, 0.1, "Labour"
    f["A3"], f["B3"], f["C3"], f["D3"], f["E3"], f["F3"] = \
        "Plant rate", "plant_rate", -0.05, 0.0, 0.1, "ALL"

    a = wb["Layer1_Allocation"]
    a["B1"], a["C1"] = "Civils", "Electrical"
    a["A2"], a["B2"], a["C2"] = "Labour rate", "x", "x"
    a["A3"], a["B3"] = "Plant rate", "x"

    c = wb["Layer1_Correlation"]
    c["B4"], c["C4"] = "Labour rate", "Plant rate"
    c["A5"], c["A6"] = "Labour rate", "Plant rate"
    c["B6"] = 0.2


class TestAnEnabledButUnfilledLayer1IsRefusedByName:
    def test_the_untouched_template_with_the_flag_on_is_refused_naming_the_flag(self, template, tmp_path):
        """The trap the characterisation found: with zeros this parsed silently to a Layer 1 that
        contributes nothing. Blank must be refused, and the message must say what to do."""
        grids = _grids_after(template, tmp_path)
        with pytest.raises(RegisterReadError, match="Set_Layer1_Enabled") as exc:
            _layer1_from_grids(grids)
        assert "Layer1_Base" in str(exc.value) and "Layer1_Factors" in str(exc.value)

    def test_a_user_fill_over_the_placeholders_round_trips(self, template, tmp_path):
        """Unused placeholders are not the user's problem: an unfilled WBS row and an unticked
        Allocation column are ignored, and the eight columns stay (blank = 0 within a filled row)."""
        cfg = _layer1_from_grids(_grids_after(template, tmp_path, _minimal_valid_fill))
        assert cfg is not None
        assert cfg.row_labels == ("Civils", "Electrical")
        assert cfg.col_labels == LAYER1_GENERIC_COLS
        assert cfg.base.shape == (2, 8)
        assert np.array_equal(cfg.base[:, :2], np.array([[10e6, 5e6], [8e6, 2e6]]))
        assert cfg.base[:, 2:].sum() == 0.0
        assert cfg.declared_total == 25e6
        assert [f.name for f in cfg.factors] == ["Labour rate", "Plant rate"]
        assert cfg.factors[0].rows == ("ALL",) and cfg.factors[1].rows == ("Civils",)
        assert cfg.corr_names == ("Labour rate", "Plant rate")
        assert cfg.corr_matrix[1, 0] == 0.2

    def test_a_tick_under_an_unfilled_wbs_placeholder_is_refused(self, template, tmp_path):
        def mutate(wb):
            _minimal_valid_fill(wb)
            wb["Layer1_Allocation"]["D2"] = "x"   # WBS-03: no numbers on Layer1_Base
        with pytest.raises(RegisterReadError, match="WBS-03"):
            _layer1_from_grids(_grids_after(template, tmp_path, mutate))

    def test_a_renamed_row_not_renamed_in_both_places_is_refused(self, template, tmp_path):
        def mutate(wb):
            _minimal_valid_fill(wb)
            wb["Layer1_Allocation"]["C1"] = "Electricals"
        with pytest.raises(RegisterReadError, match="Electricals"):
            _layer1_from_grids(_grids_after(template, tmp_path, mutate))

    def test_a_factor_with_a_blank_bound_is_named(self, template, tmp_path):
        """Was a raw TypeError from float(); now names the factor and the field."""
        def mutate(wb):
            _minimal_valid_fill(wb)
            wb["Layer1_Factors"]["E3"] = None
        with pytest.raises(RegisterReadError, match="'Plant rate'.*P90.*blank"):
            _layer1_from_grids(_grids_after(template, tmp_path, mutate))

    def test_a_value_above_the_diagonal_is_refused(self, template, tmp_path):
        """A coefficient typed where the parser never looks would otherwise vanish silently."""
        def mutate(wb):
            _minimal_valid_fill(wb)
            wb["Layer1_Correlation"]["C5"] = 0.3
        with pytest.raises(RegisterReadError, match="above the diagonal"):
            _layer1_from_grids(_grids_after(template, tmp_path, mutate))

    def test_correlation_names_must_agree_across_the_header_and_down_the_column(self, template, tmp_path):
        """Rows are read in header order; a column-A name that disagrees was silently mis-assigned."""
        def mutate(wb):
            _minimal_valid_fill(wb)
            wb["Layer1_Correlation"]["A6"] = "Labour rate"
        with pytest.raises(RegisterReadError, match="same order"):
            _layer1_from_grids(_grids_after(template, tmp_path, mutate))

    def test_the_flag_off_still_reads_to_nil(self, template):
        """Unchanged: the shipped state is OFF, and the shaped-but-blank grids are a read no-op."""
        grids = extract_grids_openpyxl(str(template))
        assert grids["layer1_enabled"] is False
        assert _layer1_from_grids(grids) is None
