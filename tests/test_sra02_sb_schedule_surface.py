"""SRA-02 S-B - the scaffolder emits the schedule surface.

Gates for the four new tabs (Risk_Activity_Patterns, Risk_Pattern_Assignments,
Summary, Treatments) plus the Grid_Mapping named range, per
the SRA02 S-A pinned schedule contract, sections 1.1-1.6
and 6.

The central discipline here: **no header literal in this file stands alone.** Every
header the scaffolder emits is cross-checked against the reader's own declared
variant list in ``src.sra.strawman_reader``, so a drift on either side reds. The
scaffolder keeps its own literals (it must not import the reader at runtime - it is
the standalone/Pyodide-adjacent surface); this test is where the two are tied
together.

The nil-read gates are the S-B analogue of the S1 cost nil gate and of SRA-01:71:
a register that ships the schedule tabs present-and-empty must read to a well-formed
nil, not a crash and not a silent default.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.register_contract import REQUIRED_NAMED_RANGES
from src.sra.strawman_reader import (
    _ASSIGNMENT_HEADERS,
    _COMPARATOR_LABELS,
    _PATTERN_HEADERS,
    _TREATMENTS_HEADERS,
    _normalise,
    get_strawman_calibration_target,
    read_risk_activity_patterns,
    read_risk_pattern_assignments,
    read_treatments,
)
from tools.register_scaffold import (
    ASSIGNMENTS_SHEET,
    CALIBRATION_LABELS,
    CALIBRATION_SHEET,
    COL_EMV_SCHED_DAYS,
    EMITTED_ONLY_NAMES,
    calibration_formula_strings,
    GRID_NAMES,
    MAPPING_SHEET,
    PATTERNS_SHEET,
    RISK_HEADERS,
    TREATMENTS_SHEET,
    _TOC_LEGEND_ROW,
    _TOC_TABLIST_FIRST_ROW,
    _generate_toc_tablist,
    ASSIGNMENT_HEADERS,
    PATTERN_HEADERS,
    TREATMENT_HEADERS,
    build_template,
)


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("sra02_sb") / "empty_template.xlsx"
    return build_template(out, prefix="RSK")


def _emv_block(register_ws) -> tuple[int, int]:
    """The register rows ``write_emv_formulas`` filled, read off the sheet itself.

    Derived rather than restated so the Summary range is compared against the block
    that actually exists, not against a repeat of the arithmetic that built it.
    """
    rows = [r for r in range(1, register_ws.max_row + 1)
            if str(register_ws.cell(row=r, column=COL_EMV_SCHED_DAYS).value or "").startswith("=")]
    assert rows, "no EMV_sched_days formulas found - the block this pins is gone"
    return rows[0], rows[-1]


# ---------------------------------------------------------------------------
# 1. The four tabs exist, on the exact sheet names the readers hard-code
# ---------------------------------------------------------------------------

def test_four_schedule_tabs_present(template):
    """The readers address these sheets by literal name, so the names are the
    contract - not a naming preference."""
    wb = openpyxl.load_workbook(template)
    try:
        for tab in (PATTERNS_SHEET, ASSIGNMENTS_SHEET, CALIBRATION_SHEET,
                    TREATMENTS_SHEET):
            assert tab in wb.sheetnames, tab
        # and the pre-existing surfaces are untouched
        assert MAPPING_SHEET in wb.sheetnames
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# 2. Every emitted header is the reader's FIRST declared variant
# ---------------------------------------------------------------------------

def _assert_emits_first_variant(emitted: list[str], reader_headers: dict[str, list[str]],
                                label: str) -> None:
    """The scaffolder emits one header per reader field, in reader field order, and
    each is the first (underscored) variant the reader declares."""
    expected = [variants[0] for variants in reader_headers.values()]
    assert emitted == expected, (
        f"{label}: scaffolder headers drifted from the reader's declared variants.\n"
        f"  scaffolder emits: {emitted}\n"
        f"  reader wants    : {expected}"
    )


def test_pattern_headers_are_reader_first_variants():
    _assert_emits_first_variant(PATTERN_HEADERS, _PATTERN_HEADERS,
                                "Risk_Activity_Patterns")


def test_assignment_headers_are_reader_first_variants():
    _assert_emits_first_variant(ASSIGNMENT_HEADERS, _ASSIGNMENT_HEADERS,
                                "Risk_Pattern_Assignments")


def test_treatment_headers_are_reader_first_variants():
    _assert_emits_first_variant(TREATMENT_HEADERS, _TREATMENTS_HEADERS,
                                "Treatments")


def test_emitted_headers_land_where_the_reader_finds_them_in_column_order(template):
    """Column order as declared, on the row the readers find by scanning column A
    (``_find_header_row``, 20 rows). Was "row 1": REGUX-01 part 3 writes a note block above
    the header, which the readers were measured to tolerate before it was promised (brief §2),
    so the property is "the reader's own scan lands on it", not a row number."""
    from src.sra.strawman_reader import _find_header_row

    wb = openpyxl.load_workbook(template)
    try:
        for sheet, headers in ((PATTERNS_SHEET, PATTERN_HEADERS),
                               (ASSIGNMENTS_SHEET, ASSIGNMENT_HEADERS),
                               (TREATMENTS_SHEET, TREATMENT_HEADERS)):
            ws = wb[sheet]
            hdr = _find_header_row(ws, headers[0])
            got = [ws.cell(row=hdr, column=c).value for c in range(1, len(headers) + 1)]
            assert got == headers, sheet
            # the rows above carry notes in column A only - nothing a reader could mistake
            above = [c.value for r in ws.iter_rows(min_row=1, max_row=hdr - 1, min_col=2)
                     for c in r if c.value is not None]
            assert not above, f"{sheet}: content beside the note block: {above[:3]}"
            # REGUX-01: was `ws.max_row == 1`. The property this guards is that the template
            # ships with NO SEEDED DATA, and max_row was proxying for it - but max_row counts
            # materialised cells, so it also forbade an EMPTY styled body. That body is what
            # makes these sheets fillable at all: `_unlock` clamps to the sheet extent, so a
            # header-only sheet clamps to nothing and shipped entirely locked. Asserted on
            # values now, which is the intent, and verified: the readers return 0 rows.
            seeded = [r for r in ws.iter_rows(min_row=hdr + 1, values_only=True)
                      if any(v is not None and str(v).strip() for v in r)]
            assert not seeded, f"{sheet} ships seeded data rows: {seeded[:3]}"
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# 3. Calibration is NOT header-driven - it is a label walk down column A
# ---------------------------------------------------------------------------

def test_calibration_labels_normalise_to_reader_literals():
    """``get_strawman_calibration_target`` matches ``_normalise``d column A against
    four exact literals. The scaffolder emits human-readable labels; what must hold
    is that they normalise onto the reader's literals. This is the one reader a
    header-based scaffold silently misses (contract 1.5)."""
    normalised = [_normalise(lbl) for lbl in CALIBRATION_LABELS]
    assert normalised == list(_COMPARATOR_LABELS.values()), (
        "Calibration labels do not normalise onto the reader's literals.\n"
        f"  scaffolder -> {normalised}\n"
        f"  reader     -> {list(_COMPARATOR_LABELS.values())}"
    )


def test_calibration_labels_in_column_a_comparator_formulas_in_column_b(template):
    """Labels in column A, the S-E comparator formulas in column B.

    Positional by design: the reader takes ``row[1]`` regardless of any header.
    Supersedes the S-B property that these cells ship blank - S-E computes the
    cross-check instead of asking the workshop to type it. The labels are the half
    the reader contract depends on and they are unchanged; the test that pins them
    is ``test_calibration_labels_normalise_to_reader_literals``.

    The expected strings come from ``calibration_formula_strings`` rather than being
    written out here, so this asserts *placement and row range*, not the formula
    text - the arithmetic is P-E2's and P-E3's to prove.
    """
    wb = openpyxl.load_workbook(template)
    try:
        first_row, last_row = _emv_block(wb["Risk Register"])
        want = calibration_formula_strings(first_row, last_row)
        ws = wb[CALIBRATION_SHEET]
        col_a = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        for label, expected in zip(CALIBRATION_LABELS, want):
            assert label in col_a, label
            r = col_a.index(label) + 1
            assert ws.cell(row=r, column=2).value == expected, (
                f"{label}: column B must carry the comparator formula"
            )
    finally:
        wb.close()


def test_calibration_formulas_span_exactly_the_emv_input_block(template):
    """The Summary sums must cover the same rows ``write_emv_formulas`` filled.

    The two surfaces are written from one ``last_input`` in ``build_template``; this
    ties them together so a change to the input-block size on one side and not the
    other reds, instead of producing a Summary that silently sums a shorter block
    than the register carries.
    """
    wb = openpyxl.load_workbook(template)
    try:
        first_row, last_row = _emv_block(wb["Risk Register"])
        b3 = wb[CALIBRATION_SHEET].cell(row=3, column=2).value
        assert b3.endswith(f"{last_row})"), b3
        assert f"{first_row}:" in b3, b3
    finally:
        wb.close()


def test_calibration_labels_appear_exactly_once(template):
    """The reader takes the FIRST occurrence only (contract 1.5). A scaffolded tab
    that repeated a label would silently pin the wrong one."""
    wb = openpyxl.load_workbook(template)
    try:
        ws = wb[CALIBRATION_SHEET]
        col_a = [_normalise(ws.cell(row=r, column=1).value)
                 for r in range(1, ws.max_row + 1)]
        for literal in _COMPARATOR_LABELS.values():
            assert col_a.count(literal) == 1, f"{literal} appears {col_a.count(literal)}x"
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# 4. The nil-read gates - SRA-01:71 / G4 analogue, per reader
# ---------------------------------------------------------------------------

def test_patterns_read_to_nil(template):
    df = read_risk_activity_patterns(template)
    assert list(df.columns) == list(_PATTERN_HEADERS)
    assert len(df) == 0


def test_assignments_read_to_nil(template):
    df = read_risk_pattern_assignments(template)
    assert list(df.columns) == list(_ASSIGNMENT_HEADERS)
    assert len(df) == 0


def test_treatments_read_to_nil(template):
    """Note the extra sheet-presence limb. ``read_treatments`` returns the same
    empty four-column frame whether the sheet is present-and-empty or absent
    entirely (contract 1.6, ``:734-735``), so the frame assertions alone pass
    against a scaffolder that never emitted the tab. Red-first caught exactly that:
    this gate passed before the tab existed. The presence limb is what makes it
    capable of failing on the property it is cited for."""
    wb = openpyxl.load_workbook(template)
    try:
        assert TREATMENTS_SHEET in wb.sheetnames, "tab absent - the frame check alone cannot see this"
    finally:
        wb.close()
    df = read_treatments(template)
    assert list(df.columns) == list(_TREATMENTS_HEADERS)
    assert len(df) == 0


def test_calibration_reads_to_nil(template):
    """All four comparator keys resolve to None on an unfilled tab - no raise, and
    no silent default standing in for a workshop number."""
    out = get_strawman_calibration_target(template)
    for key in _COMPARATOR_LABELS:
        assert key in out, key
        assert out[key] is None, f"{key} should be None on an unfilled Calibration tab"
    assert out["source_summary_cells"] == {}


# ---------------------------------------------------------------------------
# 5. Grid_Mapping - emitted, but deliberately NOT contract-required
# ---------------------------------------------------------------------------

def test_grid_mapping_emitted_over_the_mapping_sheet(template):
    wb = openpyxl.load_workbook(template)
    try:
        assert "Grid_Mapping" in wb.defined_names
        sh, _coord = list(wb.defined_names["Grid_Mapping"].destinations)[0]
        assert sh == MAPPING_SHEET
    finally:
        wb.close()
    assert GRID_NAMES[MAPPING_SHEET] == "Grid_Mapping"


def test_grid_mapping_is_emitted_only_not_required():
    """Adding Grid_Mapping to REQUIRED_NAMED_RANGES would fail Validate on every
    register scaffolded before it, which is a contract-version bump on the
    client-number-bearing lane - forbidden by D-3. It is emitted, not required."""
    assert "Grid_Mapping" in EMITTED_ONLY_NAMES
    assert "Grid_Mapping" not in REQUIRED_NAMED_RANGES


def test_template_names_are_required_plus_emitted_only(template):
    wb = openpyxl.load_workbook(template)
    try:
        got = set(wb.defined_names)
    finally:
        wb.close()
    assert got == set(REQUIRED_NAMED_RANGES) | set(EMITTED_ONLY_NAMES)


# ---------------------------------------------------------------------------
# 6. Fail-closed guards
# ---------------------------------------------------------------------------

def test_toc_tablist_cannot_silently_overwrite_the_legend(template):
    """The generated tab-list fills from _TOC_TABLIST_FIRST_ROW down. S-B takes the
    sheet inventory to 17, which reaches row 21; the legend moved to
    _TOC_LEGEND_ROW so the block holds 21 rows with four spare. The guard turns an
    overrun into a loud stop instead of a silently overwritten legend."""
    wb = openpyxl.load_workbook(template)
    try:
        toc = wb["TOC & Instructions"]
        n_sheets = len(wb.sheetnames) - 1  # TOC excludes itself
        last_used = _TOC_TABLIST_FIRST_ROW + n_sheets - 1
        assert last_used < _TOC_LEGEND_ROW, (
            f"tab-list reaches row {last_used}, legend sits at {_TOC_LEGEND_ROW}"
        )
        assert toc.cell(row=_TOC_LEGEND_ROW, column=1).value == "Cell-type legend"
    finally:
        wb.close()


def test_toc_tablist_overflow_raises_rather_than_overwriting(template):
    """The positive control for the guard above.

    The passing assertion in ``test_toc_tablist_cannot_silently_overwrite_the_legend``
    only says today's inventory fits - it would pass just as happily if the guard did
    not exist. This one drives the inventory past capacity and requires the loud
    stop, which is what makes the guard's green mean something."""
    wb = openpyxl.load_workbook(template)
    try:
        capacity = _TOC_LEGEND_ROW - _TOC_TABLIST_FIRST_ROW
        while len(wb.sheetnames) - 1 <= capacity:
            wb.create_sheet(f"_overflow_{len(wb.sheetnames)}")
        with pytest.raises(ValueError, match="TOC tab-list overflow"):
            _generate_toc_tablist(wb)
        # and the legend it protects is still intact
        assert wb["TOC & Instructions"].cell(
            row=_TOC_LEGEND_ROW, column=1).value == "Cell-type legend"
    finally:
        wb.close()


def test_frozen_cost_contract_untouched():
    """Cols 1-30 are the frozen contract-v2 set and col 31 stays the inert binding
    column. S-B adds sheets; it must not have moved a register column."""
    assert len(RISK_HEADERS) == 31
    assert RISK_HEADERS[30] == "Schedule Activity ID (P6)"
    assert RISK_HEADERS[19] == "Time Impact"
    assert RISK_HEADERS[25] == "Mode"


def test_schedule_tabs_do_not_touch_the_register_sheet(template):
    """Additivity, asserted rather than assumed: the Risk Register header row is
    byte-identical to RISK_HEADERS after the schedule surface lands."""
    wb = openpyxl.load_workbook(template)
    try:
        ws = wb["Risk Register"]
        got = [ws.cell(row=4, column=c).value for c in range(1, len(RISK_HEADERS) + 1)]
        assert got == RISK_HEADERS
    finally:
        wb.close()
