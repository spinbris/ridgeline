"""Stage 4 F-28 (D-2, first-rank): an empty pattern catalogue is no catalogue.

Found on 08/09/2026 running the hand-filled Kestrel Highway register through RUN_FINAL.bat: the
schedule arm chose its rule source by whether the register HAD a `Risk_Activity_Patterns` sheet.
Every scaffolded register has one - the scaffolder ships it, header only, on every register - so
every hand-made register was forced onto the pattern route, which its empty pattern sheets cannot
satisfy (`Risk 'KH-002' is Time=Yes but has no Pattern_ID in Risk_Pattern_Assignments`), and the
structure lock stops the filler deleting the sheet. Neither fixture carries the sheet, so both take
the prefix route from `Mapping` and never met this.

The property: the route follows what the register CONTAINS. A catalogue with pattern rows below the
header the reader finds is a catalogue; a header-only sheet, with or without a note block above it,
is not, and the register routes by its `Mapping` prefixes exactly as the fixtures do.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.integrated.stage3_config import register_carries_pattern_catalogue
from tools.register_scaffold import build_template


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    return build_template(tmp_path_factory.mktemp("f28") / "t.xlsx")


def _header_row(ws) -> int:
    for r in range(1, 21):
        if str(ws.cell(row=r, column=1).value or "").strip().lower() == "pattern_id":
            return r
    raise AssertionError("no Pattern_ID header in the first 20 rows")


def test_a_scaffolded_register_with_no_patterns_carries_no_catalogue(template):
    """Red before the fix: the sheet exists, so this returned True and the schedule arm demanded
    assignments the workshop never produced."""
    assert register_carries_pattern_catalogue(template) is False


def test_the_note_block_above_the_header_is_not_a_pattern(template):
    """Part 3 writes notes in column A above the header; they must not read as catalogue rows."""
    wb = openpyxl.load_workbook(template)
    ws = wb["Risk_Activity_Patterns"]
    assert _header_row(ws) > 1 and ws.cell(row=1, column=1).value  # the notes are there
    wb.close()
    assert register_carries_pattern_catalogue(template) is False


def test_one_pattern_row_below_the_header_makes_it_a_catalogue(template, tmp_path):
    wb = openpyxl.load_workbook(template)
    ws = wb["Risk_Activity_Patterns"]
    hdr = _header_row(ws)
    ws.cell(row=hdr + 1, column=1, value="P1")
    ws.cell(row=hdr + 1, column=2, value="Piling")
    ws.cell(row=hdr + 1, column=3, value="30.")
    p = tmp_path / "with_pattern.xlsx"
    wb.save(p)
    wb.close()
    assert register_carries_pattern_catalogue(p) is True


def test_a_register_without_the_sheet_carries_no_catalogue(tmp_path):
    """The fixtures' shape - unchanged by the fix."""
    wb = openpyxl.Workbook()
    wb.active.title = "Risk Register"
    p = tmp_path / "no_sheet.xlsx"
    wb.save(p)
    wb.close()
    assert register_carries_pattern_catalogue(p) is False


# ---------------------------------------------------------------------------
# REFRETIRE-01 (D-3): the route the detector used to select has retired, and the
# deliverable says so rather than running a catalogue register down a path that no
# longer joins its pattern sheets.
# ---------------------------------------------------------------------------

def test_the_deliverable_refuses_a_catalogue_register_by_name(template, tmp_path):
    """The refusal names the catalogue and what to do instead. Built on a real workbook, so it
    exercises the detector rather than a stubbed flag."""
    from src.final.deliverable import refuse_a_catalogue_register

    wb = openpyxl.load_workbook(template)
    ws = wb["Risk_Activity_Patterns"]
    hdr = _header_row(ws)
    ws.cell(row=hdr + 1, column=1, value="P1")
    ws.cell(row=hdr + 1, column=2, value="Piling")
    ws.cell(row=hdr + 1, column=3, value="30.")
    catalogue = tmp_path / "with_catalogue.xlsx"
    wb.save(catalogue)
    wb.close()

    with pytest.raises(ValueError, match="pattern catalogue"):
        refuse_a_catalogue_register(catalogue)

    # and a register without one passes straight through
    plain = openpyxl.Workbook()
    plain.active.title = "Risk Register"
    plain_path = tmp_path / "no_catalogue.xlsx"
    plain.save(plain_path)
    plain.close()
    assert refuse_a_catalogue_register(plain_path) is None
