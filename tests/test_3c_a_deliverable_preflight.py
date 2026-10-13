"""3c Part A A-1: one pre-flight, every missing setting in one message.

Director 04/09/2026 ("add the single pre-flight"): the same refusal posture as the individual
readers, one round trip instead of four. The property under test is not "it refuses" - each
reader already did that - but that **one** refusal names **everything** that is missing, and
names where to fix each one.

The green path runs against the committed wind fixture, so it is demonstrable on a fresh clone.
The red paths are built by removing one setting from a copy of that register: a pre-flight
whose failure has never been seen is not evidence that it can fail.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest

from src.final.deliverable_preflight import (
    DeliverablePreflightError,
    assert_ready,
    check_register_settings,
    run_preflight,
)

REPO = Path(__file__).resolve().parents[1]
WIND_DIR = REPO / "tests" / "fixtures" / "synthetic_wind"
WIND_REGISTER = WIND_DIR / "wind_register.xlsx"
WIND_XER = WIND_DIR / "wind_schedule.xer"


def _env(**overrides: str) -> dict[str, str]:
    env = {
        "QRA_PROJECT": "synthetic_wind",
        "REGISTER_FILE": str(WIND_REGISTER),
        "P6_FILE": str(WIND_XER),
        "STRAWMAN_FILE": str(WIND_REGISTER),
    }
    env.update(overrides)
    return {k: v for k, v in env.items() if v != ""}


def test_the_green_path_resolves_every_setting_on_a_committed_fixture():
    results = assert_ready(env=_env())
    assert all(r.ok for r in results)
    labels = {r.label for r in results}
    for expected in ("project binding", "register path", "p6 path", "strawman path",
                     "Mode multiplier grid", "CG correlation coefficients",
                     "five calibration tolerances", "default injection method",
                     "day rate", "programme days", "iterations", "reporting P-level"):
        assert expected in labels, f"{expected} is not checked by the pre-flight"


def test_nothing_set_reports_all_four_at_once_not_the_first():
    with pytest.raises(DeliverablePreflightError) as excinfo:
        assert_ready(env={})
    message = str(excinfo.value)
    assert "4 unmet precondition(s)" in message
    for var in ("QRA_PROJECT", "REGISTER_FILE", "P6_FILE", "STRAWMAN_FILE"):
        assert var in message, f"{var} is missing from the single message"
    # honest about what it could not reach
    assert "were NOT checked" in message


def test_an_unset_path_names_the_variable_and_does_not_invent_a_default():
    with pytest.raises(DeliverablePreflightError) as excinfo:
        assert_ready(env=_env(REGISTER_FILE=""))
    message = str(excinfo.value)
    assert "REGISTER_FILE" in message
    assert "1 unmet precondition(s)" in message


def test_a_path_set_to_a_missing_file_is_distinguished_from_an_unset_one(tmp_path):
    with pytest.raises(DeliverablePreflightError) as excinfo:
        assert_ready(env=_env(P6_FILE=str(tmp_path / "absent.xer")))
    assert "no such file" in str(excinfo.value)


def test_a_workbook_open_in_excel_is_refused_by_its_lock_file(tmp_path):
    """A stale cached value reading as a real one is the failure this prevents."""
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    (tmp_path / "~$wind_register.xlsx").write_bytes(b"")
    with pytest.raises(DeliverablePreflightError) as excinfo:
        assert_ready(env=_env(REGISTER_FILE=str(register), STRAWMAN_FILE=str(register)))
    assert "open in Excel" in str(excinfo.value)


@pytest.mark.parametrize(
    "row_label, expected_label",
    [
        ("Series", "Mode multiplier grid"),
        ("Cost_Calibration_Tolerance_Dollars", "five calibration tolerances"),
        ("Default_Injection_Method", "default injection method"),
    ],
)
def test_removing_one_register_setting_is_reported_by_name(tmp_path, row_label, expected_label):
    """Red-first, one setting at a time: the pre-flight must name the one that went."""
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    wb = openpyxl.load_workbook(register)
    sheet = wb["Methodology"]
    removed = False
    for row in sheet.iter_rows():
        if row[0].value == row_label:
            for cell in row:
                cell.value = None
            removed = True
            break
    assert removed, f"fixture no longer carries a {row_label!r} row - update this test"
    wb.save(register)
    wb.close()

    results = check_register_settings(register, register)
    failed = {r.label for r in results if not r.ok}
    assert expected_label in failed, (
        f"removing {row_label!r} did not surface as {expected_label!r}; "
        f"failures were {failed or 'none - the check cannot fail'}"
    )


def test_the_register_checks_run_even_when_an_earlier_one_fails(tmp_path):
    """The point of the design: a failure is a row, not a halt, so one message covers all."""
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    wb = openpyxl.load_workbook(register)
    sheet = wb["Methodology"]
    for row in sheet.iter_rows():
        if row[0].value in {"Series", "Default_Injection_Method"}:
            for cell in row:
                cell.value = None
    wb.save(register)
    wb.close()

    results = check_register_settings(register, register)
    failed = {r.label for r in results if not r.ok}
    ok = {r.label for r in results if r.ok}

    assert {"Mode multiplier grid", "default injection method"} <= failed, failed
    # It did not stop at the first failure: the settings that come from the OTHER surface
    # (MC_Params, not the Methodology grid) still resolved and are still reported.
    assert {"day rate", "programme days", "iterations", "reporting P-level"} <= ok, ok
    # Every check is present in the result either way - that is what makes it one round trip.
    assert len(results) == 8, [r.label for r in results]


def test_a_blanked_grid_row_fails_the_whole_methodology_grid_not_just_its_own_setting(tmp_path):
    """Recorded because the pre-flight's message is only as precise as the parser under it.

    ``resolve_methodology_from_grid`` parses the ``Methodology`` sheet as a unit, so blanking
    ONE row makes all four grid-sourced settings unreadable and the single message names four
    when one was removed. That over-reports; it does not under-report, so it cannot let a
    missing setting through - which is the property that matters here. Narrowing it to the
    offending row is a grid-parser change on the Excel side (GRIDPARSE-01), not a pre-flight
    change, and this test states the behaviour so the next reader meets it as a decision rather
    than as a surprise. If GRIDPARSE-01 lands, this is the test that should change.
    """
    register = tmp_path / "wind_register.xlsx"
    shutil.copy(WIND_REGISTER, register)
    wb = openpyxl.load_workbook(register)
    for row in wb["Methodology"].iter_rows():
        if row[0].value == "Series":
            for cell in row:
                cell.value = None
            break
    wb.save(register)
    wb.close()

    grid_sourced = {"Mode multiplier grid", "CG correlation coefficients",
                    "five calibration tolerances", "default injection method"}
    failed = {r.label for r in check_register_settings(register, register) if not r.ok}
    assert grid_sourced <= failed, (
        "one blanked Mode row no longer takes the whole grid down - if GRIDPARSE-01 has "
        f"landed, narrow this expectation to the Mode row alone. Failures: {failed}"
    )


def test_every_failure_row_says_where_to_fix_it():
    results = run_preflight(env={})
    for r in results:
        if not r.ok:
            assert r.where.strip(), f"{r.label} reports no remedy location"
