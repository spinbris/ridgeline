"""REG-01 §8 lock: the register's EMV columns reconcile to the Monte-Carlo means
the engine actually produces.

This is the whole point of the deterministic Layer-1 EMV columns: an analyst
gets an instant cross-check (no MC) that should land on the mean the MC
converges to. EMV_cost uses each risk's OWN distribution mean (§12.5), so it
reconciles to the cost MC mean; EMV_sched_days is the mode-weighted PERT mean,
so it reconciles to the Stage 2.5 schedule MC mean.

REG03 S2: the EMV columns are now LIVE Excel formulas (not builder-computed
literals), so the totals are obtained by evaluating the formulas headless via the
`formulas` lib rather than reading cached openpyxl values (which are None until
Excel opens the file). Skips cleanly when `formulas` is unavailable.
"""
from __future__ import annotations
from tests.conftest import TEST_MODE_MULTIPLIERS

import importlib.util
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
FIXTURE_DIR = REPO / "tests" / "fixtures" / "synthetic_wind"
WIND_REGISTER = FIXTURE_DIR / "wind_register.xlsx"
WIND_XER = FIXTURE_DIR / "wind_schedule.xer"


@pytest.fixture(autouse=True)
def _wind_project(monkeypatch):
    """Activate the synthetic-wind binding so readers use the WT prefix."""
    spec = importlib.util.spec_from_file_location(
        "wind_binding_emv_fixture", FIXTURE_DIR / "wind_binding.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # self-registers 'synthetic_wind'
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")


def _emv_column_totals() -> dict[str, float]:
    """Sum the EMV columns from the themed Risk Register by EVALUATING the live
    Excel formulas (REG03 S2) headless via the `formulas` lib."""
    formulas = pytest.importorskip("formulas")
    import warnings
    from openpyxl.utils import get_column_letter
    warnings.filterwarnings("ignore")

    wb = load_workbook(str(WIND_REGISTER))  # formulas, not data_only
    ws = wb["Risk Register"]
    header = [c.value for c in ws[4]]
    idx = {h: i + 1 for i, h in enumerate(header)}  # 1-based column
    cols = {
        "emv_cost": idx["EMV_cost ($)"],
        "emv_sched_days": idx["EMV_sched_days"],
        "emv_sched_dollars": idx["EMV_sched_dollars ($)"],
    }
    wt_rows = [
        r for r in range(5, ws.max_row + 1)
        if str(ws.cell(row=r, column=1).value or "").startswith("WT")
    ]
    wb.close()

    sol = formulas.ExcelModel().loads(str(WIND_REGISTER)).finish().calculate()

    def _val(col: int, row: int) -> float:
        addr = f"{get_column_letter(col)}{row}"
        for k, v in sol.items():
            if k.upper().endswith("RISK REGISTER'!" + addr):
                try:
                    cell = v.value[0, 0]
                except Exception:
                    cell = v.value
                return float(cell) if isinstance(cell, (int, float)) else 0.0
        return 0.0

    totals = {k: 0.0 for k in cols}
    for r in wt_rows:
        for k, ci in cols.items():
            totals[k] += _val(ci, r)
    return totals


def test_emv_cost_total_matches_cost_mc_mean():
    """EMV_cost column total ~ cost MC mean (per-distribution means, §12.5)."""
    from src.config import ChainConfig, SimConfig
    from src.simulation import MonteCarloEngine
    from src.v5_register_reader import read_risks_v5

    totals = _emv_column_totals()
    risks = read_risks_v5(str(WIND_REGISTER))
    cfg = SimConfig(n_iter=40000, seed=42)
    cg_coeffs = {f"CG{n:02d}": 0.5 for n in range(1, 14)}
    engine = MonteCarloEngine(
        cfg, buckets=[], chains=ChainConfig(), risks=risks,
        corr_coefficients=cg_coeffs, seed=42,
    )
    mc_mean = float(np.asarray(engine.run().ev_costs).sum(axis=1).mean())
    emv = totals["emv_cost"]
    rel = abs(emv - mc_mean) / mc_mean
    assert rel < 0.02, (
        f"EMV_cost total ${emv:,.0f} vs cost MC mean ${mc_mean:,.0f} "
        f"(rel {rel:.3%}) - deterministic column does not reconcile to the MC"
    )


def test_emv_sched_days_total_matches_stage25_mc_mean():
    """EMV_sched_days column total ~ Stage 2.5 MC schedule mean within 0.1d."""
    from src.sra.sra_config import SRAConfig
    from src.sra.sra_simulation import run_schedule_simulation
    from src.sra.strawman_reader import read_risk_scope

    totals = _emv_column_totals()
    risk_scope = read_risk_scope(str(WIND_REGISTER))
    cfg = SRAConfig(
        register_path=WIND_REGISTER, p6_path=WIND_XER, strawman_path=WIND_REGISTER,
        iterations=80000, random_seed=42, day_rate=35_000.0,
    mode_multipliers=TEST_MODE_MULTIPLIERS)
    result = run_schedule_simulation(cfg, risk_scope)
    mc_mean = float(np.asarray(result["finish_delay_days_array"]).mean())
    emv = totals["emv_sched_days"]
    assert abs(emv - mc_mean) < 0.1, (
        f"EMV_sched_days total {emv:.4f}d vs Stage 2.5 MC mean {mc_mean:.4f}d "
        f"(gap {abs(emv - mc_mean):.4f}d) exceeds 0.1d"
    )
