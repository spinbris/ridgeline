"""REGUX-01 part 3: the workbook earns its "standalone" claim.

Stage 4 F-7: the workbook said "fill it in, it is standalone" and carried column headers and
nothing else. F-10: the README that answers F-7 had been written once, on the Tropic Ridge factor
model, and lost when the template was generalised. This part puts the answers where the eye is -
a units/basis block on the front sheet (A-1), a note block above every list sheet's headers (A-2),
every specified column's format on its header cell (A-3), and a one-line banner on every sheet
nobody fills (§5.3) - and every word of it is rendered from `src.register_field_spec`, which the
drift test holds against the parsers (B-1). That is what keeps this from rotting the way
`SOFTWARE_USER_GUIDE.md` §3.1 did.

What this suite cannot establish: that a person can still type where they should and not where
they should not once rows have moved. That is `scripts/ask_excel_register.py` (§6a).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.register_field_spec import FIELD_SPECS, FILLABLE_SHEETS, note_lines, specs_for
from src.sra.strawman_reader import (
    _find_header_row, _read_mapping_sheet, read_risk_activity_patterns,
    read_risk_pattern_assignments, read_treatments,
)
from tools.register_scaffold import (
    LIST_SPECS, _BANNERS, _LISTS_HEADROOM_LAST_ROW, _TOC_FILL_ORDER, build_template,
)

INPUT_BLUE = "000000CC"
LIST_SHEETS = ("Mapping", "Risk_Activity_Patterns", "Risk_Pattern_Assignments", "Treatments")


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    return build_template(tmp_path_factory.mktemp("regux01p3") / "t.xlsx")


@pytest.fixture(scope="module")
def wb(template):
    book = openpyxl.load_workbook(template)
    yield book
    book.close()


def _blue(cell) -> bool:
    colour = cell.font.color.rgb if (cell.font and cell.font.color and cell.font.color.rgb) else None
    return str(colour) == INPUT_BLUE


def _locked(cell) -> bool:
    return True if cell.protection is None else bool(cell.protection.locked)


def _header_row(ws) -> int:
    (table,) = ws.tables.values()
    return int("".join(ch for ch in table.ref.split(":")[0] if ch.isdigit()))


class TestTheNoteBlocksSitAboveTheHeaders:
    @pytest.mark.parametrize("sheet", LIST_SHEETS)
    def test_the_block_is_the_field_spec_rendered_and_nothing_else(self, sheet, wb):
        """Every column line above the header is `note_lines(sheet)` verbatim, so the words a
        filler reads are the words the drift test holds against the parser."""
        ws = wb[sheet]
        hdr = _header_row(ws)
        assert hdr > 1, f"{sheet}: no note block above the header"
        above = [ws.cell(row=r, column=1).value for r in range(1, hdr)]
        assert above[-1] is None, "a blank spacer row separates the notes from the header"
        assert above[1:-1] == note_lines(sheet)
        assert "Blue cells are yours" in str(above[0])

    @pytest.mark.parametrize("sheet", LIST_SHEETS)
    def test_the_reader_still_finds_the_header_and_the_table_follows_it(self, sheet, wb):
        ws = wb[sheet]
        hdr = _header_row(ws)
        first_header = ws.cell(row=hdr, column=1).value
        assert _find_header_row(ws, first_header) == hdr
        assert ws.freeze_panes == f"A{hdr + 1}"
        (table,) = ws.tables.values()
        assert table.ref.startswith(f"A{hdr}:")

    @pytest.mark.parametrize("sheet", LIST_SHEETS)
    def test_notes_and_header_are_locked_and_the_body_is_not(self, sheet, wb):
        """Unlock-by-style, held as the biconditional: notes black and locked, header locked, body
        blue and open. A range table could not express this; the style can."""
        ws = wb[sheet]
        hdr = _header_row(ws)
        for r in range(1, hdr + 1):
            cell = ws.cell(row=r, column=1)
            assert _locked(cell) and not _blue(cell), cell.coordinate
        body = ws.cell(row=hdr + 1, column=1)
        assert _blue(body) and not _locked(body)
        wrong = [c.coordinate for c in ws._cells.values() if _blue(c) == _locked(c)]
        assert not wrong, f"{sheet}: unlocked xor blue at {wrong[:8]}"

    def test_a_filled_copy_reads_through_every_reader(self, template, tmp_path):
        """The property the brief measured before promising A-2 (§2): a note block above the
        headers parses cleanly. Asserted through the four readers, not by inspection."""
        wb = openpyxl.load_workbook(template)
        rows = {
            "Mapping": ["RSK-001", "P1", "20.", 3, "A1000", "additive_pert_days"],
            "Risk_Activity_Patterns": ["P1", "Piling", "20.", "No", None, None, None, None],
            "Risk_Pattern_Assignments": ["RSK-001", "P1", "high", "n"],
            "Treatments": ["RSK-001", "Extra rig", 30, "firm"],
        }
        for sheet, values in rows.items():
            ws = wb[sheet]
            hdr = _header_row(ws)
            for c, v in enumerate(values, start=1):
                ws.cell(row=hdr + 1, column=c, value=v)
        p = tmp_path / "filled.xlsx"
        wb.save(p)
        wb.close()
        assert list(read_risk_activity_patterns(p)["pattern_id"]) == ["P1"]
        assert list(read_risk_pattern_assignments(p)["risk_id"]) == ["RSK-001"]
        assert list(read_treatments(p)["risk_id"]) == ["RSK-001"]
        book = openpyxl.load_workbook(p, data_only=True)
        try:
            mapped = _read_mapping_sheet(book)
        finally:
            book.close()
        assert "RSK-001" in mapped


class TestEveryColumnsFormatIsOnItsHeader:
    @pytest.mark.parametrize("sheet", FILLABLE_SHEETS)
    def test_every_specified_column_carries_its_guidance_as_a_header_note(self, sheet, wb):
        ws = wb[sheet]
        hdr = 4 if sheet == "Risk Register" else _header_row(ws)
        by_text = {str(c.value).strip(): c for c in ws[hdr] if c.value is not None}
        for spec in specs_for(sheet):
            cell = by_text[spec.column]
            assert cell.comment is not None, f"{sheet}!{cell.coordinate} has no note"
            assert spec.guidance in cell.comment.text
            if spec.legal_values:
                for v in spec.legal_values:
                    assert v in cell.comment.text

    def test_the_risk_register_says_its_piece_above_the_section_band(self, wb):
        """Its header is row 4 by contract, so the two rows above the band carry the note."""
        ws = wb["Risk Register"]
        assert "FRACTION" in str(ws["A2"].value)
        assert "TOC & Instructions" in str(ws["A1"].value)
        assert _locked(ws["A1"]) and _locked(ws["A2"])
        # the contract row is untouched
        assert ws["A4"].value == "Risk ID"


class TestTheSheetsNobodyFillsSaySo:
    @pytest.mark.parametrize("sheet,cell", [(s, c) for s, (c, _t) in _BANNERS.items()])
    def test_banner_present_and_locked(self, sheet, cell, wb):
        c = wb[sheet][cell]
        assert c.value and ("do not type" in c.value.lower() or "nothing to fill" in c.value.lower())
        assert _locked(c)

    def test_summary_is_no_longer_called_a_workshop_output(self, wb):
        """Stage 4 F-9: a filler following the label destroyed the comparators."""
        toc = wb["TOC & Instructions"]
        row = next(r for r in range(5, 26) if toc.cell(row=r, column=1).value == "Summary")
        desc = str(toc.cell(row=row, column=2).value)
        assert "workshop output" not in desc.lower() and "calculated" in desc.lower()
        assert toc.cell(row=row, column=3).value == "calculated"


class TestTheFrontSheetIsAReadme:
    def test_units_basis_and_formats_are_stated(self, wb):
        toc = wb["TOC & Instructions"]
        text = " ".join(str(c.value) for row in toc.iter_rows(min_row=40, max_row=59)
                        for c in row if c.value)
        for must in ("working days", "$M", "before escalation", "FRACTION", "P10 and P90",
                     "Binary/PERT", "one severity draw", "(1 + factor)", "MC Include is the switch",
                     "never used by the engine"):
            assert must in text, f"the units/basis block does not say: {must!r}"

    def test_the_tab_list_is_in_fill_order_with_who_writes_each(self, wb):
        toc = wb["TOC & Instructions"]
        listed = [toc.cell(row=r, column=1).value for r in range(5, 26)
                  if toc.cell(row=r, column=1).value]
        assert listed[:len(_TOC_FILL_ORDER)] == list(_TOC_FILL_ORDER)
        kinds = {toc.cell(row=r, column=1).value: toc.cell(row=r, column=3).value
                 for r in range(5, 26) if toc.cell(row=r, column=1).value}
        assert kinds["Risk Register"] == "you fill"
        assert kinds["Simulation"] == "model writes"
        assert kinds["Summary"] == "calculated"

    def test_the_units_block_does_not_reach_the_project_identity_block(self, wb):
        toc = wb["TOC & Instructions"]
        assert toc["A60"].value == "Project identity"
        assert toc.cell(row=59, column=1).value is None and toc.cell(row=59, column=2).value is None


class TestTheRemainingFormatting:
    def test_binary_pert_is_offered_and_the_dropdown_matches_the_spec(self):
        """The EMV formula and the sampler both handled Binary/PERT; the dropdown did not offer
        it, so a filler typing it tripped validation. Held equal to the spec in both directions."""
        seeds = next(seeds for _c, name, _d, seeds in LIST_SPECS if name == "List_Distribution")
        spec = next(f for f in FIELD_SPECS if f.column == "Distribution")
        assert set(seeds) == set(spec.legal_values)

    def test_the_lists_have_headroom_and_it_is_open(self, wb):
        """The unlock used to stop at the foot of each seeded list, so a project could not add a
        package or a factor without unprotecting the sheet."""
        ls = wb["Lists"]
        for col, _n, _d, _s in LIST_SPECS:
            cell = ls[f"{col}{_LISTS_HEADROOM_LAST_ROW}"]
            assert _blue(cell) and not _locked(cell), cell.coordinate
        assert _locked(ls["A2"]) and _locked(ls["A1"])
        assert ls.freeze_panes == "A3"

    def test_the_prose_columns_on_the_register_are_wide_and_wrap(self, wb):
        rr = wb["Risk Register"]
        for letter in ("K", "L", "M"):
            assert rr.column_dimensions[letter].width >= 30
            assert rr[f"{letter}5"].alignment.wrap_text

    def test_matrix_sheets_freeze_their_labels(self, wb):
        assert wb["Layer1_Base"].freeze_panes == "B2"
        assert wb["Layer1_Allocation"].freeze_panes == "B2"
        assert wb["Layer1_Correlation"].freeze_panes == "B5"
