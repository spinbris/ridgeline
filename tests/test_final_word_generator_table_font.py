"""Beta-19b item 9 — table body font reduction unit tests.

Locks the contract:
- `_TABLE_BODY_FONT_PT` exposed as a module constant equal to 9.
- `_apply_data_table_font` walks every Light Grid / Light List styled
  table and applies Calibri 9pt to body cells / 9pt bold to headers.
- Banner tables (no style attribute) are left untouched.
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt

from src.final.word_generator import (
    _TABLE_BODY_FONT_PT,
    _TABLE_FONT_NAME,
    _apply_data_table_font,
)


def test_table_body_font_constant_is_9pt() -> None:
    """The Beta-19b table-body constant is 9pt (Calibri)."""
    assert _TABLE_BODY_FONT_PT == 9
    assert _TABLE_FONT_NAME == "Calibri"


def test_apply_data_table_font_sets_calibri_9pt_on_data_tables() -> None:
    """Every cell in a Light-Grid styled table picks up 9pt Calibri."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Body A"
    table.rows[1].cells[1].text = "Body B"

    _apply_data_table_font(doc)

    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    assert run.font.name == "Calibri"
                    assert run.font.size == Pt(9)
                    # Header row should be bolded; body row stays
                    # regular.
                    if row_idx == 0:
                        assert run.font.bold is True


def test_apply_data_table_font_skips_tables_without_light_style() -> None:
    """Banner tables (no style) are left untouched.

    Adds a styled cell to the unstyled table with an explicit 14pt
    pre-existing run; helper must NOT overwrite that run's size.
    """
    doc = Document()
    banner_table = doc.add_table(rows=1, cols=1)
    # NOTE: no banner_table.style assignment — banner pattern.
    cell = banner_table.cell(0, 0)
    cell.text = "BANNER"
    for run in cell.paragraphs[0].runs:
        run.font.size = Pt(14)

    _apply_data_table_font(doc)

    for run in banner_table.cell(0, 0).paragraphs[0].runs:
        assert run.font.size == Pt(14)  # unchanged
