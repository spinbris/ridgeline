"""Regression: cp:revision must be numeric-or-absent so Excel does not run
"file level validation and repair" on open.

Root cause (22/07/2026): build_project_block set the workbook cp:revision core
property to the engineering revision letter ("A"); Excel treats cp:revision as a
numeric save-counter and repairs any file where it is not one. The revision letter
belongs in the Set_Revision cell (display/contract surface), not the core property.
See the REG05 fix brief.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from tools.register_scaffold import build_project_block, build_template

REPO = Path(__file__).resolve().parents[1]
FIXTURE = REPO / "tests" / "fixtures" / "synthetic_wind" / "wind_register.xlsx"


def _cp_revision(xlsx_path) -> str | None:
    with zipfile.ZipFile(xlsx_path) as z:
        if "docProps/core.xml" not in z.namelist():
            return None
        core = z.read("docProps/core.xml").decode()
    m = re.search(r"<cp:revision>(.*?)</cp:revision>", core)
    return m.group(1) if m else None


def test_scaffolder_core_revision_numeric_or_absent(tmp_path):
    out = build_template(tmp_path / "reg.xlsx")
    rev = _cp_revision(out)
    assert rev is None or rev.isdigit(), f"cp:revision {rev!r} must be numeric or absent"


def test_scaffolder_keeps_revision_letter_in_cell(tmp_path):
    out = build_template(tmp_path / "reg.xlsx")
    wb = openpyxl.load_workbook(out)
    sheet, coord = next(iter(wb.defined_names["Set_Revision"].destinations))
    assert wb[sheet][coord].value == "A"


def test_committed_fixture_core_revision_numeric_or_absent():
    rev = _cp_revision(FIXTURE)
    assert rev is None or rev.isdigit(), f"fixture cp:revision {rev!r} must be numeric or absent"


def test_build_project_block_keeps_letter_out_of_core_revision():
    # Even when a caller passes an engineering revision letter, it must not reach
    # cp:revision - that is exactly what triggered the Excel file-level repair.
    wb = Workbook()
    wb.create_sheet("TOC & Instructions")
    build_project_block(wb, {"Set_Revision": "B"})
    rev = wb.properties.revision
    assert rev is None or str(rev).isdigit(), f"props.revision {rev!r} must be numeric or unset"
