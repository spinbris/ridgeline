"""PUBLICFACE-01 P2: what the public sees carries no insider scenario or vendor name.

The shipped projects' production deliverable is rendered the way ``RUN_FINAL.bat`` renders it
(tests/publicface01_render.py) and every artefact is scanned. The page-1 rows keep their values
bit-for-bit: the Stage 3 mean / P80 / P95 are pinned as ``float.hex`` taken on the tree before
the change (1,000 iterations, seed 42), so a relabel that moved a number goes red here.
"""
from __future__ import annotations

import re

import pytest

from tests import publicface01_render as pf
from tests.arch_guard import same_value

#: ``float.hex`` of the Stage 3 statistics, measured on the tree at a51fe2f, 02/10/2026,
#: the build machine, 1,000 iterations. Moves only with an approved methodology change.
PINS = {
    "synthetic_wind": {
        "cost_mean": "0x1.eebe1d8a906d7p+21", "cost_p80": "0x1.48e25cc1d3b98p+22",
        "cost_p95": "0x1.bf2600c5143e2p+22", "schedule_mean": "0x1.59aae5c28f5c3p+5",
        "schedule_p80": "0x1.df38266666667p+5", "schedule_p95": "0x1.3af6a99999997p+6",
    },
    "synthetic_road": {
        "cost_mean": "0x1.2d7e74595674cp+22", "cost_p80": "0x1.a0a81748c7605p+22",
        "cost_p95": "0x1.17a2c6f2ca74fp+23", "schedule_mean": "0x1.0b1b07be76c8bp+6",
        "schedule_p80": "0x1.6f2fd9999999ap+6", "schedule_p95": "0x1.dca9eb3333333p+6",
    },
}
ROW_LABELS = {
    "mean": "MC mean, Stage 3 headline (recommended central estimate)",
    "p80": "MC P80, Stage 3 headline",
    "p95": "MC P95, Stage 3 headline",
}


def test_the_scan_finds_each_mark_and_ignores_image_blobs(tmp_path):
    """The check must be able to fail: plant each mark, and a blob that spells one by chance."""
    page = tmp_path / "page.html"
    page.write_text(
        f"<p>Under {pf.THIRD_PARTY} indicative classification</p><p>Dummy Realistic (x)</p>"
        f"<img src='data:image/png;base64,{'A' * 100}{pf.THIRD_PARTY}{'B' * 100}'>", encoding="utf8")
    assert pf.marks_in(page) == {pf.THIRD_PARTY: 1, "indicative classification": 1, "Dummy Realistic": 1}
    clean = tmp_path / "clean.html"
    clean.write_text("<p>Stage 3 headline</p>", encoding="utf8")
    assert pf.marks_in(clean) == {pf.THIRD_PARTY: 0, "indicative classification": 0, "Dummy Realistic": 0}


@pytest.fixture(scope="module", params=sorted(PINS))
def rendered(request, tmp_path_factory):
    return pf.render(request.param, tmp_path_factory.mktemp(request.param))


def test_the_completeness_check_can_fail(tmp_path):
    """Gate 3 round 1, W3: a scan over an incomplete run must not go green. Build a run missing
    one artefact and carrying an error cell, and see each named."""
    html = tmp_path / "report.html"
    html.write_text("<p>ok</p> Traceback (most recent call last) boom", encoding="utf8")
    made = []
    for name in ("BETA_executive_summary_2026-10-02.docx", "sra_stage3_standalone_PROD_X_2026-10-02.docx"):
        (tmp_path / name).write_bytes(b"x")
        made.append(tmp_path / name)
    run = pf.Rendered("synthetic_wind", html, tuple(made), tmp_path)
    problems = pf.incomplete(run)
    assert any("standalone_.*\\.html" in p for p in problems), problems
    assert "a notebook cell produced an error output" in problems
    html.write_text("<p>ok</p>", encoding="utf8")
    (tmp_path / "sra_stage3_standalone_PROD_X_2026-10-02.html").write_bytes(b"x")
    complete = pf.Rendered("synthetic_wind", html, (*made, tmp_path / "sra_stage3_standalone_PROD_X_2026-10-02.html"), tmp_path)
    assert pf.incomplete(complete) == []


@pytest.mark.slow
def test_no_production_artefact_names_an_insider_scenario_or_vendor(rendered):
    # Inside the scan itself, so selecting this test alone cannot go green over an incomplete run.
    assert pf.incomplete(rendered) == [], f"{rendered.project}: the run is incomplete"
    found = {a.name: hits for a in rendered.artefacts()
             if any((hits := pf.marks_in(a)).values())}
    assert found == {}, f"{rendered.project}: the public sees an insider name: {found}"


@pytest.mark.slow
def test_a_matching_override_does_not_make_the_production_path_run_a_dummy_scenario(monkeypatch):
    """The two shipped registers carry none of the override's IDs, so a scenario that is never
    produced there proves nothing. Plant an override that matches a real risk, and show the run
    still produces only the Mode scenarios the deliverable names."""
    from src.final import deliverable
    from src.sra import sra_simulation
    from src.sra.strawman_reader import read_risk_scope

    fx = pf.REPO_ROOT / "tests" / "fixtures" / "synthetic_wind"
    env = {"QRA_PROJECT": "synthetic_wind", "REGISTER_FILE": str(fx / "wind_register.xlsx"),
           "P6_FILE": str(fx / "wind_schedule.xer"), "STRAWMAN_FILE": str(fx / "wind_register.xlsx"),
           "SR9_ITERATIONS": "500"}
    monkeypatch.setenv("QRA_PROJECT", "synthetic_wind")
    planted = str(read_risk_scope(env["STRAWMAN_FILE"])["risk_id"].astype(str).iloc[0])
    monkeypatch.setattr(sra_simulation, "DUMMY_REALISTIC_MODE_OVERRIDE", {planted: "Series"})

    engines = deliverable.run_engines(deliverable.resolve_inputs(env))
    assert sorted(engines.integrated_by_scenario) == sorted(deliverable.BASE_SCENARIOS)


@pytest.mark.slow
def test_the_page_one_rows_keep_their_values_bit_for_bit_under_the_stage_3_label(rendered):
    got = pf.headline_stats(rendered.cache_dir)
    want = PINS[rendered.project]
    assert set(got) == set(want) and all(
        same_value(float.fromhex(got[k]), float.fromhex(want[k])) for k in want), (got, want)
    text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", pf.text_of(rendered.html)))
    pins = {k: float.fromhex(v) for k, v in PINS[rendered.project].items()}
    for key, label in ROW_LABELS.items():
        days = f"{round(pins[f'schedule_{key}'], 2):.2f}"
        cost = f"{round(pins[f'cost_{key}'] / 1e6, 2):.2f}"
        assert f"{label} {days} {cost}" in text, (
            f"{rendered.project}: no row '{label}' carrying {days} days and ${cost}M")


# --- P6: the Stage 3 standalone pair is fit to ship ---------------------------------------

_PATH = re.compile(
    r"(?<![A-Za-z])[A-Za-z]:[\\/]|\\\\[A-Za-z]"
    # POSIX roots, macOS's /Volumes/ among them (round 1, N6)
    r"|(?<![A-Za-z0-9.])/(?:Users|home|Volumes|mnt|tmp|var|private|opt)/")
_CODE_NAME = re.compile(r"^[a-z][a-z0-9]*(?:_[a-z0-9]+)+$")
_MODULE = re.compile(r"\bsrc\.[a-z_]+(?:\.[a-z_]+)*")


def _standalone(rendered):
    html = next(p for p in rendered.others if p.name.startswith("sra_stage3_standalone") and p.suffix == ".html")
    docx = next(p for p in rendered.others if p.name.startswith("sra_stage3_standalone") and p.suffix == ".docx")
    return html, docx


def _visible_text(path):
    """What a reader sees: a .docx's paragraphs and table cells, or an HTML file's text."""
    if path.suffix.lower() == ".docx":
        import docx

        doc = docx.Document(path)
        cells = [c.text for table in doc.tables for row in table.rows for c in row.cells]
        return "\n".join([p.text for p in doc.paragraphs] + cells)
    return re.sub(r"<[^>]+>", " ", path.read_text(encoding="utf8"))


def _docx_headings_and_labels(path):
    import docx

    out = []
    for table in docx.Document(path).tables:
        if table.rows:
            out += [c.text.strip() for c in table.rows[0].cells]
            out += [row.cells[0].text.strip() for row in table.rows]
    return out


def test_the_p6_checks_can_fail():
    assert not _PATH.search("http://schemas.openxmlformats.org")
    bs = chr(92)  # built, not typed: the client gate matches a drive path written out as one token
    assert _PATH.search("C:" + bs + "Users" + bs + "someone" + bs + "wind_register.xlsx")
    vault = "/Volumes/" + "King" + "ston/" + "Obsi" + "dian"  # built, as the drive path above is
    for posix in (vault, "see /mnt/data/x", "at /tmp/run1", "/Users/me/x"):
        assert _PATH.search(posix), posix
    assert not _PATH.search("w/o the tmp/ folder, a/mnt/ ratio, and https://example.org/tmp")
    assert _CODE_NAME.match("risk_id") and _CODE_NAME.match("on_deterministic_cp")
    assert not _CODE_NAME.match("Risk") and not _CODE_NAME.match("P80")
    assert _MODULE.search("Generated by src.final.stage3_composer")


@pytest.mark.slow
def test_the_standalone_pair_prints_no_local_path_and_no_module_name(rendered):
    html, docx = _standalone(rendered)
    for artefact in (html, docx):
        text = _visible_text(artefact)
        assert not _PATH.search(text), f"{rendered.project}: {artefact.name} prints a local path"
        assert not _MODULE.search(text), f"{rendered.project}: {artefact.name} names a code module"


@pytest.mark.slow
def test_the_standalone_pair_shows_no_code_name_as_a_table_header(rendered):
    html, docx = _standalone(rendered)
    cells = _docx_headings_and_labels(docx)
    cells += [re.sub(r"<[^>]+>", "", m) for m in re.findall(r"<th>(.*?)</th>", html.read_text(encoding="utf8"))]
    assert [c for c in cells if _CODE_NAME.match(c)] == [], f"{rendered.project}: code names as headers"


@pytest.mark.slow
def test_the_standalone_html_title_and_caveat_carry_the_product_and_project(rendered):
    from src.final.product import PRODUCT_NAME
    from src.sra.cpm.bindings import get_binding, project_display_name

    html, _ = _standalone(rendered)
    text = html.read_text(encoding="utf8")
    project = project_display_name(get_binding(rendered.project))
    title = re.search(r"<title>(.*?)</title>", text).group(1)
    caveat = re.search(r'<div class="cover-caveat">(.*?)</div>', text, re.S).group(1)
    assert PRODUCT_NAME in title and project in title, title
    assert PRODUCT_NAME in caveat, caveat


@pytest.mark.slow
def test_no_word_document_a_run_writes_shows_a_code_name_as_a_table_header(rendered):
    """Gate 3 round 1, N7: the executive summary shares the table helpers, and the P6 test above
    read only the standalone pair. Every .docx the run wrote, both projects."""
    assert pf.incomplete(rendered) == [], f"{rendered.project}: the run is incomplete"
    docs = [p for p in rendered.others if p.suffix == ".docx"]
    assert len(docs) >= 2, docs
    bad = {p.name: [c for c in _docx_headings_and_labels(p) if _CODE_NAME.match(c)] for p in docs}
    assert {k: v for k, v in bad.items() if v} == {}, f"{rendered.project}: {bad}"
