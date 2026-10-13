"""3c Part A A-2/A-3/A-5: the §6a sequence, D-11's halt behaviour, and the wind reproduction.

These run engines, so they carry ``slow`` and land in the serial ``outputs`` group with the
other co-tenants (``conftest.pytest_collection_modifyitems``). Iteration counts are small; the
properties held here are structural, and the *numbers* are held by the wind golden.

The load-bearing one is :func:`test_the_deliverable_reproduces_the_showcase_headline_exactly`.
A-5 says the deliverable on the wind binding reproduces the showcase headline exactly, and the
Director's option-A decision was taken so that "exactly" would be structural. This asserts it
element by element on ``float.hex()``, the same equality the goldens use - not a tolerance,
which would be a silent methodology change (GOLDEN01 rule 2).
"""
from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pytest

from src.final.deliverable import (
    DeliverableRun,
    resolve_inputs,
    run_deliverable,
    run_engines,
)
from src.final.deliverable_preflight import DeliverablePreflightError

REPO = Path(__file__).resolve().parents[1]
WIND_DIR = REPO / "tests" / "fixtures" / "synthetic_wind"
ITERATIONS = "2000"


@pytest.fixture
def wind_env() -> dict[str, str]:
    return {
        "QRA_PROJECT": "synthetic_wind",
        "REGISTER_FILE": str(WIND_DIR / "wind_register.xlsx"),
        "P6_FILE": str(WIND_DIR / "wind_schedule.xer"),
        "STRAWMAN_FILE": str(WIND_DIR / "wind_register.xlsx"),
        "SR9_ITERATIONS": ITERATIONS,
        # CONFIGTRUTH-01: unset is now a drawn seed; the drawn-seed property is
        # tests/test_configtruth01_seed.py. Here the seed is the override, and says so.
        "SR9_SEED": "42",
    }


@pytest.fixture
def wind_project():
    """Select the wind binding for the duration of one test, and put it back."""
    import src.sra.cpm.bindings.synthetic_wind  # noqa: F401  self-registers

    saved = os.environ.get("QRA_PROJECT")
    os.environ["QRA_PROJECT"] = "synthetic_wind"
    try:
        yield
    finally:
        if saved is None:
            os.environ.pop("QRA_PROJECT", None)
        else:
            os.environ["QRA_PROJECT"] = saved


@pytest.fixture
def showcase_at_test_iterations(monkeypatch):
    """The showcase module forced to this test's iteration count, and put back afterwards.

    Without the teardown the cached module keeps n=2000 for every later consumer in the same
    process. The golden adapter re-forces both values itself so it is not at risk, but a test
    that leaves module state behind is one refactor away from being the reason another test
    is quietly running at the wrong iteration count.
    """
    import scripts.run_wind_showcase as wind

    monkeypatch.setenv("WIND_ITERATIONS", ITERATIONS)
    monkeypatch.setattr(wind, "N_ITER", int(ITERATIONS))
    monkeypatch.setattr(wind, "STAGE3_ITERATIONS", int(ITERATIONS))
    return wind


@pytest.mark.slow
def test_the_deliverable_reproduces_the_showcase_headline_exactly(
    wind_env, wind_project, showcase_at_test_iterations, tmp_path
):
    """A-5. One shared construction, so this is equality by structure, not coincidence."""
    wind = showcase_at_test_iterations
    wind._activate_wind_binding()
    showcase = np.asarray(wind.run_integrated()["integrated"], dtype=float)

    engines = run_engines(resolve_inputs(wind_env, output_dir=tmp_path))
    deliverable = engines.headline_integrated_array

    assert deliverable.shape == showcase.shape
    mismatches = [
        (i, a.hex(), b.hex())
        for i, (a, b) in enumerate(zip(showcase, deliverable))
        if a.hex() != b.hex()
    ]
    assert mismatches == [], (
        f"{len(mismatches)} of {showcase.size} elements differ from the showcase headline; "
        f"first: {mismatches[0] if mismatches else None}"
    )


@pytest.mark.slow
def test_the_sequence_produces_the_outputs_of_6a_and_no_pdf(wind_env, wind_project, tmp_path):
    """A-3 and D-12: HTML is nbconvert's; here, Word + Stage-3 standalone + state cache.

    REPORTDOC-01 P-5: and the guide's §7 table names every file on disk, and nothing that is
    not. Run with both headline metrics so the MEAN row is exercised by a real file too.
    """
    from tests.guide_artefacts import artefact_rows, disagreement, html_row_from_runner, written_files

    run = run_deliverable(env={**wind_env, "FINAL_REPORT_HEADLINE_METRIC": "both"}, output_dir=tmp_path)

    assert run.all_pass is True, run.cross_validation.get("interpretation")
    for artefact in run.artefacts:
        assert Path(artefact).exists(), artefact

    suffixes = {Path(a).suffix.lower() for a in run.artefacts}
    assert ".pdf" not in suffixes, "D-12: the deliverable produces no PDF"
    assert ".docx" in suffixes
    assert not list(tmp_path.rglob("*.pdf")), "a PDF was written somewhere under the output dir"

    rows = artefact_rows((REPO / "docs" / "SOFTWARE_USER_GUIDE.md").read_text(encoding="utf-8"))
    html = html_row_from_runner((REPO / "RUN_FINAL.bat").read_text(encoding="utf-8", errors="ignore"))
    found = disagreement(rows, written_files(tmp_path), {html}, headline="both")
    assert found == {"unnamed": [], "unwritten": [], "ambiguous": []}, (
        f"docs/SOFTWARE_USER_GUIDE.md §7 and this run's files disagree: {found}")

    # DIAGFIX-01 P-6: the run diagnostic against this run's own provenance and cross-validation
    # result. REPORT01's tests checked it against dicts in the shapes their author expected, and
    # five lines were wrong on every real page while they passed.
    diagnostic = next(Path(a) for a in run.artefacts if Path(a).name.startswith("run_diagnostic_"))
    page = diagnostic.read_text(encoding="utf-8")
    words = [ln.split() for ln in page.splitlines()]
    prov, cv = run.provenance, run.cross_validation
    assert ["Seed:", str(prov["params"]["seed"])] in words, "F-2: the seed is not the run's"
    for name, digest in prov["input_sha256"].items():
        assert [f"{name}:", digest["value"], f"({digest['method']})"] in words, f"F-3: {name}"
    assert "{'method'" not in page, "F-3: a digest printed as a raw dict"
    for label, key in (("1 cost calibration", "cost_passes_calibration"),
                       ("2 schedule", "schedule_passes_calibration"),
                       ("3 top-3 cost drivers", "top_3_cost_contributors_consistent"),
                       ("4 top-3 schedule drivers", "top_3_schedule_contributors_consistent")):
        assert cv[key] is True, key
        assert any(ln.strip().startswith(f"PASS  {label}") for ln in page.splitlines()), f"F-1: {label}"
    schedule_block = page.split("Schedule, by")[1].split("GATE")[0]
    for row in run.composed["page4_schedule_drivers"]["drivers_table"][:3]:
        line = next(ln for ln in schedule_block.splitlines() if ln.split()[:1] == [row["Risk"]])
        assert line.split()[-1] == str(row["Probability x mean impact (days)"]), f"F-4: {line}"

    # P-7: re-rendering from the cache reproduces the page the run wrote, byte for byte.
    import importlib.util
    spec = importlib.util.spec_from_file_location("run_diagnostic_script", REPO / "scripts" / "run_diagnostic.py")
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    rerendered = tmp_path / "rerendered_diagnostic.txt"
    assert script.main(["--project", prov["project"], "--run-date", prov["run_date"],
                        "--output-dir", str(tmp_path), "--out", str(rerendered)]) == 0
    assert rerendered.read_bytes() == diagnostic.read_bytes(), "P-7: the two writers disagree"


@pytest.mark.slow
def test_every_cross_validation_check_is_reported_and_the_gate_can_be_read(wind_env, wind_project, tmp_path):
    run = run_deliverable(env=wind_env, output_dir=tmp_path)
    cv = run.cross_validation
    for key in ("cost_passes_calibration", "schedule_passes_calibration",
                "top_3_cost_contributors_consistent",
                "top_3_schedule_contributors_consistent", "all_pass"):
        assert key in cv, key
    # STAGE3HEADLINE-01 (D-3, 10/09/2026): run_deliverable always supplies stage3_results to
    # cross_validate_reports now, so check 4 ranks the Stage 3 cost arm's additive per-risk
    # draws, not the register's own Mode basis, on every project. Was "register_mode_baseline"
    # under 3c Part A, before Stage 3 became the headline.
    assert cv["schedule_top3_basis"] == "stage3_cost_arm_additive"


@pytest.mark.slow
def test_d11_writes_every_artefact_bannered_and_only_then_fails(wind_env, wind_project, tmp_path):
    """D-11, both halves, in the order that makes them compatible.

    The two requirements the second review found contradictory - nothing withheld, nothing
    unbannered reaching a client - are met by the order: write, then fail. So the test asserts
    the artefacts EXIST after a failing gate, that the Word document carries the banner, and
    that the refusal comes afterwards.
    """
    from src.final.deliverable import compose, cross_validate, write_artefacts
    from src.final.word_generator import FAIL_BANNER_LINES

    inputs = resolve_inputs(wind_env, output_dir=tmp_path)
    engines = run_engines(inputs)
    cv = dict(cross_validate(inputs, engines))
    assert cv["all_pass"] is True, "fixture no longer passes; this test needs a passing base"

    # Fail the gate the way a real divergence would: one check false, all_pass false.
    cv["top_3_schedule_contributors_consistent"] = False
    cv["all_pass"] = False

    composed = compose(inputs, engines, cv)
    artefacts, stage3_error = write_artefacts(inputs, engines, composed, cv)

    assert stage3_error is None, f"the Stage 3 pair failed to emit: {stage3_error}"
    assert artefacts, "a failing gate wrote nothing - D-11 requires every artefact written"
    for artefact in artefacts:
        assert Path(artefact).exists(), f"missing on a FAIL run: {artefact}"

    docx = [Path(a) for a in artefacts if Path(a).suffix.lower() == ".docx"]
    assert docx, "no Word document on a FAIL run"
    from docx import Document

    text = "\n".join(p.text for p in Document(str(docx[0])).paragraphs)
    tables = " ".join(
        cell.text for t in Document(str(docx[0])).tables for row in t.rows for cell in row.cells
    )
    banner_line = FAIL_BANNER_LINES[0]
    assert banner_line in text or banner_line in tables, (
        "the Word document on a failing run does not carry the FAIL banner"
    )

    run = DeliverableRun(inputs=inputs, engines=engines, cross_validation=cv,
                         composed=composed, artefacts=artefacts, provenance={})
    assert run.all_pass is False
    with pytest.raises(RuntimeError, match="CROSS-VALIDATION FAILED"):
        run.raise_if_failed()


def test_a_stage3_emit_failure_fails_the_run_rather_than_shipping_stage25_numbers():
    """DELIVERABLE-SWALLOW-01: the Stage 3 artefact is absent and loud, never wrong.

    The composer used to catch any exception from the Stage 3 rewiring and substitute Stage
    2.5 numbers into an artefact still titled Stage 3, with a log line as the only signal.
    It now refuses, and the run carries that refusal into its verdict - so a deliverable
    missing one of §6a's outputs cannot report a pass.
    """
    run = DeliverableRun(
        inputs=None, engines=None,
        cross_validation={"all_pass": True, "interpretation": "all four checks passed"},
        composed={}, artefacts=[Path("outputs/executive_summary.docx")], provenance={},
        stage3_standalone_error="RuntimeError: Stage 3 schedule-arm rewiring failed",
    )
    assert run.all_pass is False, (
        "a run missing its Stage 3 pair reported a pass because the four checks passed")
    with pytest.raises(RuntimeError, match="INCOMPLETE"):
        run.raise_if_failed()


def test_a_clean_run_with_no_stage3_error_still_passes():
    """The control: the new failure path must not fail a run that is actually fine."""
    run = DeliverableRun(
        inputs=None, engines=None,
        cross_validation={"all_pass": True, "interpretation": ""},
        composed={}, artefacts=[], provenance={}, stage3_standalone_error=None,
    )
    assert run.all_pass is True
    run.raise_if_failed()


@pytest.mark.slow
def test_a_passing_run_does_not_raise_and_carries_no_banner(wind_env, wind_project, tmp_path):
    """The other half of D-11's red-first pair: the wind fixture passes and exits cleanly."""
    run = run_deliverable(env=wind_env, output_dir=tmp_path)
    assert run.all_pass is True
    run.raise_if_failed()  # must not raise

    from docx import Document
    from src.final.word_generator import FAIL_BANNER_LINES

    for docx in (Path(a) for a in run.artefacts if Path(a).suffix.lower() == ".docx"):
        text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
        assert FAIL_BANNER_LINES[0] not in text, f"FAIL banner on a passing run: {docx}"


@pytest.mark.slow
def test_the_provenance_manifest_carries_what_claude_md_requires(wind_env, wind_project, tmp_path):
    """A-5 / SPEC §3 limb 3: seed, input hashes, engine SHA, params, timestamp - one record."""
    run = run_deliverable(env=wind_env, output_dir=tmp_path)
    prov = run.provenance

    assert set(prov["input_sha256"]) == {"register", "xer", "strawman"}
    for digest in prov["input_sha256"].values():
        assert digest["value"], digest
    assert "engine_sha" in prov and "engine_tree_dirty" in prov
    assert prov["params"]["seed"] == 42
    assert prov["params"]["seed_source"] == "override"
    assert prov["params"]["iterations"] == int(ITERATIONS)
    assert prov["timestamp"]
    assert prov["project"] == "synthetic_wind"
    assert prov["scenarios_run"] == ["All_Series", "DP_Review"]

    manifest = [Path(a) for a in run.artefacts if Path(a).name == "provenance.json"]
    assert manifest and manifest[0].exists()


def test_the_deliverable_refuses_before_computing_anything_when_unconfigured(tmp_path):
    """A-1 / §6c: no default project and no literal path. The refusal is the pre-flight's."""
    with pytest.raises(DeliverablePreflightError) as excinfo:
        run_deliverable(env={}, output_dir=tmp_path)
    assert "QRA_PROJECT" in str(excinfo.value)
    assert not list(tmp_path.rglob("*")), "a refused run wrote something"


@pytest.mark.slow
def test_the_title_comes_from_the_binding_not_a_literal(wind_env, wind_project, tmp_path):
    """A-1's portability rule (RS-A2): no project name is baked into the deliverable."""
    inputs = resolve_inputs(wind_env, output_dir=tmp_path)
    from src.sra.cpm.bindings import get_active_binding

    binding = get_active_binding()
    assert inputs.project_display_name in {
        binding.narrative("report_cover_title"), binding.project_name,
    }
    assert inputs.project_display_name
