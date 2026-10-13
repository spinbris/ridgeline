"""STATUSWB-01 (stage 4 F-31): a period status workbook can be built from any register.

Stage 4 step 7 could not be started by the analyst because the period cycle's input had no
template, no page and no generator; the Builder wrote a Kestrel-only script instead. This suite
holds the generic builder to three things: the period engine's reader accepts what it writes;
it refuses by name when the register does not carry what the rule needs; and the shipped
template is what the builder produces (the register template's own drift check, applied here).
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path

import openpyxl
import pytest

from scripts.build_status_workbook import (
    build_status_workbook, build_template, main, parse_milestone, tracked_risks,
)
from src.am.workbook_reader import read_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
ROAD_REGISTER = REPO_ROOT / "tests" / "fixtures" / "synthetic_road" / "road_register.xlsx"
TRACKED_TEMPLATE = REPO_ROOT / "templates" / "status_workbook_template.xlsx"


def _road_time_yes_ids() -> list[str]:
    """The rule, read independently of the builder: Open and Time Impact = Yes, by header."""
    ws = openpyxl.load_workbook(ROAD_REGISTER, data_only=True, read_only=True)["Risk Register"]
    header = [str(c or "").strip() for c in next(ws.iter_rows(min_row=4, max_row=4, values_only=True))]
    i_id, i_status, i_time = header.index("Risk ID"), header.index("Risk Status"), header.index("Time Impact")
    ids = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[i_id] and row[i_status] == "Open" and row[i_time] == "Yes":
            ids.append(str(row[i_id]))
    return ids


def test_the_reader_accepts_a_period_built_from_the_road_register(tmp_path):
    out = build_status_workbook(
        ROAD_REGISTER, tmp_path / "period_01.xlsx", period_end=_dt.date(2026, 8, 31),
        phase="Establishment", analyst="A", reviewer="R", notes="baseline",
        milestones=(parse_milestone("KM-01|Site establishment complete|2026-09-30"),),
    )
    data = read_workbook(out, register_path=ROAD_REGISTER)
    expected = _road_time_yes_ids()
    assert expected, "the fixture must carry Time=Yes risks for this test to mean anything"
    assert data.am_risk_ids == expected
    assert len(data.risk_status) == len(expected)
    assert len(data.milestones) == 1
    assert data.mitigations == []


def test_tracked_risks_carry_the_register_probability_mode_and_owner():
    risks = tracked_risks(ROAD_REGISTER)
    assert [r.risk_id for r in risks] == _road_time_yes_ids()
    assert all(0.0 < r.probability <= 1.0 for r in risks)
    assert all(r.mode in {"Series", "Hybrid", "Parallel"} for r in risks), {r.mode for r in risks}
    assert all(r.owner for r in risks)


def test_refuses_a_register_with_nothing_to_track(tmp_path):
    wb = openpyxl.load_workbook(ROAD_REGISTER)
    ws = wb["Risk Register"]
    header = [str(c.value or "").strip() for c in ws[4]]
    col = header.index("Time Impact") + 1
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == "Yes":
            ws.cell(row=r, column=col).value = "No"
    p = tmp_path / "no_time.xlsx"
    wb.save(p)
    with pytest.raises(SystemExit, match="no Open, Time Impact = Yes risks"):
        tracked_risks(p)


def test_refuses_a_register_missing_a_header_by_name(tmp_path):
    wb = openpyxl.load_workbook(ROAD_REGISTER)
    ws = wb["Risk Register"]
    for c in ws[4]:
        if str(c.value or "").strip() == "Probability":
            c.value = None
    p = tmp_path / "no_header.xlsx"
    wb.save(p)
    # "Probability (Time)" is still there and must not stand in for it (SPEC R-10).
    with pytest.raises(SystemExit, match=r"no header \['Probability'\]"):
        tracked_risks(p)


def _fixture_with(tmp_path, edit):
    wb = openpyxl.load_workbook(ROAD_REGISTER)
    ws = wb["Risk Register"]
    header = [str(c.value or "").strip() for c in ws[4]]
    edit(ws, header)
    p = tmp_path / "edited.xlsx"
    wb.save(p)
    return p


def test_refuses_more_tracked_risks_than_the_reader_reads(tmp_path):
    """Numerics review of 5a: the reader scans AM_Period_Status rows 15-30 and no further, so an
    overflowing period would be read short with no error. Refused by name instead."""
    from scripts.build_status_workbook import MAX_TRACKED_RISKS

    def edit(ws, header):
        col = header.index("Time Impact") + 1
        for r in range(5, ws.max_row + 1):
            if ws.cell(row=r, column=1).value:
                ws.cell(row=r, column=col).value = "Yes"   # all 30 rows tracked
    p = _fixture_with(tmp_path, edit)
    with pytest.raises(SystemExit, match=f"reads at most {MAX_TRACKED_RISKS}"):
        tracked_risks(p)


def test_refuses_a_tracked_risk_with_a_blank_probability(tmp_path):
    """A blank is not zero: a tracked risk that never fires is a wrong period number."""
    def edit(ws, header):
        col = header.index("Probability") + 1
        time_col = header.index("Time Impact") + 1
        first = next(r for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=time_col).value == "Yes")
        ws.cell(row=first, column=col).value = None
    p = _fixture_with(tmp_path, edit)
    with pytest.raises(SystemExit, match="'Probability' is blank"):
        tracked_risks(p)


def test_refuses_more_milestones_than_the_reader_reads(tmp_path):
    from scripts.build_status_workbook import MAX_MILESTONES

    too_many = tuple(parse_milestone(f"M-{i:02d}|m|2027-01-01") for i in range(MAX_MILESTONES + 1))
    with pytest.raises(SystemExit, match=f"reads at most {MAX_MILESTONES}"):
        build_status_workbook(ROAD_REGISTER, tmp_path / "p.xlsx", period_end=_dt.date(2026, 8, 31),
                              milestones=too_many)


def test_a_bad_milestone_spec_is_refused():
    with pytest.raises(SystemExit, match="ID\\|Name\\|YYYY-MM-DD"):
        parse_milestone("KM-01|no date")


def test_cli_builds_in_process(tmp_path, capsys):
    out = tmp_path / "p.xlsx"
    rc = main(["--register", str(ROAD_REGISTER), "--period-end", "2026-08-31", "--out", str(out)])
    assert rc == 0 and out.is_file()
    assert "tracked risks, period end 2026-08-31" in capsys.readouterr().out


def test_cli_refuses_without_its_inputs():
    with pytest.raises(SystemExit):
        main(["--period-end", "2026-08-31"])


# --- the shipped template is what the builder produces ---------------------------------------

def _structure(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    names = {name: sorted(str(d) for d in dn.destinations) for name, dn in wb.defined_names.items()}
    cells = {
        ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
        for ws in wb.worksheets
    }
    return {"sheets": wb.sheetnames, "names": names, "cells": cells}


def test_tracked_template_is_what_the_builder_produces(tmp_path):
    assert TRACKED_TEMPLATE.is_file(), "templates/status_workbook_template.xlsx is a tracked artefact"
    fresh = build_template(tmp_path / "t.xlsx")
    assert _structure(TRACKED_TEMPLATE) == _structure(fresh)


def test_template_carries_the_contract_and_no_rows(tmp_path):
    p = build_template(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Risk Register", "AM_Period_Status", "AM_Milestones", "AM_Mitigation_Log"]
    assert set(wb.defined_names) >= {"am_period_end", "am_previous_period", "am_period_notes"}
    ps = wb["AM_Period_Status"]
    assert ps.cell(row=14, column=1).value == "Risk_ID"
    assert ps.cell(row=15, column=1).value is None
