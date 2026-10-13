"""Stage 4 F-30: the cross-validation gate's schedule top-3 check treats a within-noise tie as
consistent - same risks in the same order, except where a position's two members tie within the
band in both arms - and the band is a register value.

Found on 08/09/2026: KH-002 and KH-017 tie exactly at 7.35 probability-weighted mean days on the
register's basis; each arm broke the tie by its own sampled means and the deliverable was stamped
DO NOT DISTRIBUTE on a register that was internally consistent. The check was also ``list ==
list``, so the same three risks in another order would have failed it.

The Director ruled a within-noise tie consistent at 2%, and "do not hard code the 2%" - so it is the register's
``Schedule_Top3_Tie_Tolerance_Pct`` row, read with the four calibration tolerances, refused by name
when absent, and scaled for iteration count exactly as the absolute tolerances are.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pytest

from src.final.cross_validator import _top_ranks_consistent, cross_validate_reports
from src.sra.methodology_defaults_reader import MethodologyDefaultsError, read_calibration_tolerances
from tests.test_sr9_cross_validator import _build_results, _make_config
from tools.register_scaffold import build_template


# --- the rule, on its own -----------------------------------------------------------------

def test_order_of_clearly_different_risks_is_a_disagreement():
    """A reversal is the disagreement the check exists to catch (test_3c_a_check4); the band
    must not turn it into a set comparison."""
    means = {"A": 10.0, "B": 9.0, "C": 7.0}
    ok, note = _top_ranks_consistent(["A", "B", "C"], ["C", "B", "A"], means, means, 0.02)
    assert not ok and "beyond the band" in note


def test_order_of_within_band_risks_is_not_a_disagreement():
    means_a = {"A": 10.0, "B": 9.95, "C": 7.0}
    means_b = {"B": 10.0, "A": 9.9, "C": 7.0}
    ok, note = _top_ranks_consistent(["A", "B", "C"], ["B", "A", "C"], means_a, means_b, 0.02)
    assert ok, note


def test_an_exact_tie_broken_two_ways_is_consistent_within_the_band():
    """The Kestrel case: third place is a coin toss between two equals."""
    sr5 = {"A": 10.0, "B": 9.0, "C": 7.35, "D": 7.34}
    integ = {"A": 10.1, "B": 8.9, "D": 7.36, "C": 7.33}
    ok, note = _top_ranks_consistent(["A", "B", "C"], ["A", "B", "D"], sr5, integ, 0.02)
    assert ok, note
    assert "within the band" in note


def test_the_same_tie_fails_with_no_band():
    """The band is what does it - with the register saying 0, the old behaviour stands."""
    sr5 = {"A": 10.0, "B": 9.0, "C": 7.35, "D": 7.34}
    integ = {"A": 10.1, "B": 8.9, "D": 7.36, "C": 7.33}
    ok, _ = _top_ranks_consistent(["A", "B", "C"], ["A", "B", "D"], sr5, integ, 0.0)
    assert not ok


def test_a_real_disagreement_still_fails():
    """A risk the arms rank apart by far more than the band is the disagreement the check exists
    to catch; the band must not swallow it."""
    sr5 = {"A": 10.0, "B": 9.0, "C": 7.0, "D": 2.0}
    integ = {"A": 10.0, "B": 9.0, "D": 7.0, "C": 2.0}
    ok, note = _top_ranks_consistent(["A", "B", "C"], ["A", "B", "D"], sr5, integ, 0.02)
    assert not ok
    assert "beyond the band" in note


def test_within_band_in_one_arm_only_is_a_disagreement():
    """The conjunction is the rule's conservatism: a pair that ties in the SR5 arm but sits far
    apart in the integrated arm is the integrated arm disagreeing, and must fail."""
    sr5 = {"A": 10.0, "B": 9.0, "C": 7.35, "D": 7.34}
    integ = {"A": 10.0, "B": 9.0, "D": 7.35, "C": 5.00}
    ok, note = _top_ranks_consistent(["A", "B", "C"], ["A", "B", "D"], sr5, integ, 0.02)
    assert not ok and "beyond the band" in note


def test_a_swapped_risk_the_other_arm_never_priced_fails():
    ok, _ = _top_ranks_consistent(["A", "B", "C"], ["A", "B", "D"], {"A": 10, "B": 9, "C": 7},
                                 {"A": 10, "B": 9, "D": 7}, 0.5)
    assert not ok


# --- through the gate ----------------------------------------------------------------------

@pytest.fixture
def strawman_path(tmp_path) -> Path:
    p = tmp_path / "strawman.xlsx"
    p.write_bytes(b"")
    return p


def _tied_case(seed: int = 42):
    """Five risks; the schedule arms' third and fourth are equal on the deterministic basis."""
    return _build_results(
        n_iter=4000, cost_mean_dollars=11_000_000.0, sched_mean_days=1.5,
        risk_ids=["R001", "R002", "R003", "R004", "R005"],
        cost_per_risk_means=[5_000_000, 3_000_000, 2_000_000, 800_000, 200_000],
        sched_per_risk_means=[1.0, 0.4, 0.1, 0.1],
        sched_risk_ids=["R001", "R002", "R003", "R004"], seed=seed,
    )


def test_the_gate_reads_the_band_from_the_config_and_reports_it(strawman_path):
    cost, sched, integ = _tied_case()
    integ = dict(integ)
    arr = np.asarray(integ["per_risk_schedule_array"]).copy()
    arr[:, [2, 3]] = arr[:, [3, 2]]                      # the integrated arm breaks the tie the other way
    integ["per_risk_schedule_array"] = arr
    integ["risk_ids_dual"] = list(integ["risk_ids_dual"])
    cfg = _make_config(strawman_path)
    res = cross_validate_reports(cost, sched, integ, strawman_path=strawman_path, config=cfg)
    assert res["schedule_top3_tie_tolerance_pct_base"] == 0.02
    assert res["schedule_top3_tie_tolerance_pct_effective"] >= 0.02   # scaled up below 100k
    assert "tie band" in res["interpretation"]
    # whatever the sampled ranks did, two equal risks swapping third place is consistent
    assert res["top_3_schedule_contributors_consistent"] is True


def test_the_gate_still_fails_a_reversed_ranking(strawman_path):
    """The existing red case (test_3c_a_check4) restated: a wholesale reversal is not a tie."""
    cost, sched, integ = _tied_case()
    integ = dict(integ)
    integ["per_risk_schedule_array"] = np.asarray(integ["per_risk_schedule_array"])[:, ::-1]
    cfg = _make_config(strawman_path)
    res = cross_validate_reports(cost, sched, integ, strawman_path=strawman_path, config=cfg)
    assert res["top_3_schedule_contributors_consistent"] is False


# --- the register carries it ---------------------------------------------------------------

def test_the_template_carries_the_row_and_the_reader_returns_it(tmp_path):
    p = build_template(tmp_path / "t.xlsx")
    tol = read_calibration_tolerances(p, caller_module="test")
    assert tol["schedule_top3_tie_tolerance_pct"] == 0.02
    ws = openpyxl.load_workbook(p)["Methodology"]
    row = next(r for r in range(2, 30) if ws.cell(row=r, column=1).value == "Schedule_Top3_Tie_Tolerance_Pct")
    assert ws.cell(row=row, column=2).protection.locked is False   # an input a specialist types


def test_a_register_without_the_row_is_refused_by_name(tmp_path):
    p = build_template(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb["Methodology"]
    for r in range(2, 30):
        if ws.cell(row=r, column=1).value == "Schedule_Top3_Tie_Tolerance_Pct":
            ws.cell(row=r, column=1).value = None
            ws.cell(row=r, column=2).value = None
    q = tmp_path / "without.xlsx"
    wb.save(q)
    with pytest.raises(MethodologyDefaultsError, match="schedule_top3_tie_tolerance_pct"):
        read_calibration_tolerances(q, caller_module="test")
