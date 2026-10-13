"""P6 schedule reader for the SRA workstream.

Reads a Primavera P6 export xlsx and returns activities as a pandas
DataFrame. Header-driven column lookup; no hardcoded positional indices.

Note on the strawman reader
---------------------------
A separate P6 reader exists at ``src/sra_strawman/p6_reader.py``, shipped
in commit eac9a99 as part of the SRA Strawman Phase 2 build. That reader
returns ``list[dict]`` and is consumed by the strawman pipeline.

SR1 deliberately ships a clean, self-contained DataFrame-shaped reader
here rather than wrapping the strawman, for three reasons:
    1. The SRA workstream needs a DataFrame API; wrapping list[dict] hides
       type-coercion choices (NaT vs None, Int64 vs int) that are easier
       to reason about in fresh code.
    2. The strawman is frozen at one shipped commit; coupling SR1 to its
       internals would create maintenance friction.
    3. Self-containment keeps SR2 debugging single-layer.

Two pieces of logic were copied across from the strawman with attribution:
    - Asterisk handling on date cells (``_parse_date``).
    - Duplicate-"WBS"-header disambiguation (``_find_columns``).
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


# ----------------------------------------------------------------------------
# public API
# ----------------------------------------------------------------------------

def read_p6_schedule(
    path: str | Path, sheet: str | None = None
) -> pd.DataFrame:
    """Read a P6 schedule export xlsx and return activities as a DataFrame.

    Parameters
    ----------
    path : str or Path
        Path to the P6 xlsx file.
    sheet : str or None
        Sheet name. If None, the function auto-detects by scanning each
        sheet's first 10 rows for an "Activity ID" header cell.

    Returns
    -------
    pd.DataFrame
        Columns (header-driven, resolved at read time):
          - activity_id (str)
          - activity_name (str)
          - wbs (str)
          - duration_days (float, NaN allowed)
          - start_date (datetime64[ns], NaT allowed)
          - finish_date (datetime64[ns], NaT allowed)
          - total_float_days (Int64, pd.NA allowed)
          - critical_flag (bool, True where P6 Critical = "Yes")
          - primary_constraint (str, empty string if none)
          - row_type (one of: activity, milestone, wbs_summary, other)

    Raises
    ------
    FileNotFoundError
        If path does not exist.
    ValueError
        If the sheet cannot be auto-detected or required columns are missing.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"P6 file not found: {p}")

    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        ws_name = sheet if sheet is not None else _auto_detect_sheet(wb)
        if ws_name not in wb.sheetnames:
            raise ValueError(
                f"Sheet {ws_name!r} not found. Available: {wb.sheetnames}"
            )
        ws = wb[ws_name]
        header_row = _find_p6_header_row(ws)
        col_map = _find_columns(ws, header_row)
        rows = _read_rows(ws, col_map, header_row)
    finally:
        wb.close()

    df = pd.DataFrame(
        rows,
        columns=[
            "activity_id", "activity_name", "wbs", "duration_days",
            "start_date", "finish_date", "total_float_days",
            "critical_flag", "primary_constraint", "row_type",
        ],
    )
    for c in ("start_date", "finish_date"):
        df[c] = pd.to_datetime(df[c], errors="coerce")
    df["duration_days"] = pd.to_numeric(df["duration_days"], errors="coerce")
    df["total_float_days"] = (
        pd.to_numeric(df["total_float_days"], errors="coerce")
        .round()
        .astype("Int64")
    )
    df["critical_flag"] = df["critical_flag"].astype(bool)
    return df


# ----------------------------------------------------------------------------
# header / column resolution
# ----------------------------------------------------------------------------

# Field -> ranked list of acceptable header strings (post-normalisation).
# Multiple variants are tried in order; first present wins.
_FIELD_HEADERS: dict[str, tuple[str, ...]] = {
    "activity_id": ("activity id",),
    "activity_name": ("activity name",),
    # 'wbs' is duplicate-prone; handled separately in _find_columns.
    "duration_days": ("duration", "original duration", "remaining duration"),
    "start_date": ("start", "start date"),
    "finish_date": ("finish", "finish date"),
    "total_float_days": ("total float", "tf", "free float"),
    "critical_flag": ("critical",),
    "primary_constraint": (
        "primary constraint", "constraint type", "constraint",
    ),
}

_REQUIRED_FIELDS: tuple[str, ...] = (
    "activity_id", "activity_name", "duration_days",
    "start_date", "finish_date", "total_float_days",
)


def _normalise_header(s: object) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _auto_detect_sheet(wb) -> str:
    target = "activity id"
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            if any(_normalise_header(c) == target for c in row):
                return name
    raise ValueError(
        f"Could not auto-detect P6 sheet (no 'Activity ID' header in "
        f"rows 1-10 of any sheet). Available: {wb.sheetnames}"
    )


def _find_p6_header_row(ws) -> int:
    """Return 1-based row index of the row containing 'Activity ID'."""
    target = "activity id"
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1
    ):
        if any(_normalise_header(c) == target for c in row):
            return row_idx
    raise ValueError(
        f"Could not find P6 header row: no 'Activity ID' cell in rows "
        f"1-10 of sheet {ws.title!r}."
    )


def _find_columns(ws, header_row: int) -> dict[str, int]:
    """Return mapping {field_name: 1-based column index} via header lookup.

    P6 exports often carry two columns headed "WBS": one with a YES/NO
    summary flag and one with the WBS path string. Disambiguation logic
    adapted from src/sra_strawman/p6_reader.py (commit eac9a99): leftmost
    "WBS" is the flag, rightmost is the path.
    """
    headers = list(
        ws.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True
        )
    )[0]
    norm = [_normalise_header(h) for h in headers]

    col_map: dict[str, int] = {}

    for field, variants in _FIELD_HEADERS.items():
        for variant in variants:
            for i, h in enumerate(norm, start=1):
                if h == variant:
                    col_map[field] = i
                    break
            if field in col_map:
                break

    wbs_path_idx = next(
        (i for i, h in enumerate(norm, start=1) if h == "wbs path"), None
    )
    wbs_code_idx = next(
        (i for i, h in enumerate(norm, start=1) if h == "wbs code"), None
    )
    plain_wbs_indices = [
        i for i, h in enumerate(norm, start=1) if h == "wbs"
    ]

    if wbs_path_idx is not None:
        col_map["wbs"] = wbs_path_idx
    elif wbs_code_idx is not None:
        col_map["wbs"] = wbs_code_idx
    elif plain_wbs_indices:
        col_map["wbs"] = plain_wbs_indices[-1]

    if len(plain_wbs_indices) >= 2:
        col_map["_wbs_flag"] = plain_wbs_indices[0]
    elif (
        len(plain_wbs_indices) == 1
        and col_map.get("wbs") != plain_wbs_indices[0]
    ):
        col_map["_wbs_flag"] = plain_wbs_indices[0]

    missing = [f for f in _REQUIRED_FIELDS if f not in col_map]
    if missing:
        raise ValueError(
            f"P6 header is missing required columns: {missing}. "
            f"Headers found: {[h for h in headers if h is not None]}"
        )

    return col_map


# ----------------------------------------------------------------------------
# cell parsers
# ----------------------------------------------------------------------------

_DURATION_RE = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*d?\s*$", re.IGNORECASE)


def _parse_duration(value: object) -> float | None:
    """Parse '950.0d' / '15.0d' / '0' style cells. None if blank/unparseable."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    m = _DURATION_RE.match(s)
    if m:
        return float(m.group(1))
    try:
        return float(s)
    except ValueError:
        return None


def _parse_date(value: object) -> object:
    """Return a datetime, ISO-ish string, or None for pandas to coerce.

    P6 sometimes appends a trailing asterisk on constrained dates; strip it.
    Asterisk handling adapted from src/sra_strawman/p6_reader.py
    (commit eac9a99). The 22 April P6 file uses a dedicated Primary
    Constraint column, so asterisk parsing is a legacy concern only.
    TODO: an older 17 April P6 file uses asterisks for the constraint
    convention; primary_constraint extraction from asterisks is out of
    scope at SR1.
    """
    if value is None:
        return None
    if hasattr(value, "year"):
        return value
    s = str(value).strip().rstrip("*").strip()
    if not s:
        return None
    return s


def _str_or_empty(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _coerce_critical(v: object) -> bool:
    if v is None:
        return False
    return str(v).strip().lower() == "yes"


# ----------------------------------------------------------------------------
# row classification + read
# ----------------------------------------------------------------------------

def _classify_row_type(
    activity_id: str,
    duration: float | None,
    wbs_flag_value: object,
) -> str:
    """Row-type classification (rule order locked at SR1).

    1. WBS flag column == 'YES' -> wbs_summary
    2. duration == 0 and activity_id non-empty -> milestone
    3. activity_id non-empty -> activity
    4. else -> other
    """
    if wbs_flag_value is not None:
        if str(wbs_flag_value).strip().upper() == "YES":
            return "wbs_summary"
    if activity_id:
        if duration is not None and duration == 0:
            return "milestone"
        return "activity"
    return "other"


def _read_rows(ws, col_map: dict[str, int], header_row: int) -> list[dict]:
    out: list[dict] = []
    flag_idx = col_map.get("_wbs_flag")

    def cell_at(raw_row, idx):
        if idx is None or idx > len(raw_row):
            return None
        return raw_row[idx - 1]

    for raw_row in ws.iter_rows(
        min_row=header_row + 1, values_only=True
    ):
        if raw_row is None or all(v is None for v in raw_row):
            continue

        activity_id = _str_or_empty(cell_at(raw_row, col_map.get("activity_id")))
        duration = _parse_duration(cell_at(raw_row, col_map.get("duration_days")))
        wbs_flag_value = cell_at(raw_row, flag_idx) if flag_idx else None

        if not activity_id and (
            wbs_flag_value is None or str(wbs_flag_value).strip() == ""
        ):
            continue

        out.append(
            {
                "activity_id": activity_id,
                "activity_name": _str_or_empty(
                    cell_at(raw_row, col_map.get("activity_name"))
                ),
                "wbs": _str_or_empty(cell_at(raw_row, col_map.get("wbs"))),
                "duration_days": duration,
                "start_date": _parse_date(
                    cell_at(raw_row, col_map.get("start_date"))
                ),
                "finish_date": _parse_date(
                    cell_at(raw_row, col_map.get("finish_date"))
                ),
                "total_float_days": _parse_duration(
                    cell_at(raw_row, col_map.get("total_float_days"))
                ),
                "critical_flag": _coerce_critical(
                    cell_at(raw_row, col_map.get("critical_flag"))
                ),
                "primary_constraint": _str_or_empty(
                    cell_at(raw_row, col_map.get("primary_constraint"))
                ),
                "row_type": _classify_row_type(
                    activity_id, duration, wbs_flag_value
                ),
            }
        )

    return out
