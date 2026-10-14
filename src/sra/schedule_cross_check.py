"""SRA-02 S-E: the deterministic schedule cross-check, Python side.

The Python half of the two independent computations the stage reconciles. The
Excel half is ``tools.register_scaffold.calibration_formula_strings``, written into
the register's ``Summary`` tab; this module computes the same four figures from the
consolidated ``read_risk_scope`` frame, and
``tests/test_sra02_se_reconciliation.py`` proves the two agree.

Two comparators, both over the Time=Yes rows:

* **Comparator 1** - total EMV days: the PERT mean of the three-point duration,
  weighted by the risk's ``probability`` (SPEC R-10), with **no** Mode multiplier. Total exposure
  before any judgement about how much of it reaches the finish date.
* **Comparator 2** - aggregate finish-date slippage days: the same per-risk EMV
  days multiplied by the register's Mode multiplier. Series / Hybrid / Parallel is
  the deterministic proxy for how much of a delay lands on the critical path, which
  is what the CPM schedule arm derives properly from the network. That is what makes
  the pair a cross-check: a crude closed-form estimate against a full network MC.

Dollars are days x day rate in both cases.

Which one a Monte Carlo run is judged against is the run's Mode basis (SCHEDCAL-01): Comparator 2
for a schedule arm that applies the Mode multipliers (Stage 2.5), Comparator 1 for one that does
not (Stage 3, the headline). See :func:`evaluate_schedule_calibration`.

**The Mode multipliers are the register's own** (D-12): they come in as a parameter,
read from the register's ``Methodology`` grid by :func:`read_register_mode_multipliers`,
never from a cross-project constant. This is deliberate and is the reason the
reconciliation is meaningful rather than two copies of one literal - the Excel side
resolves them with ``INDEX/MATCH`` over the same grid.

Note the divergence this does **not** follow: the strawman builder computes
a float-aware Hybrid multiplier (``src/sra_strawman/strawman_builder.py:346-350``),
where the register and the CPM engine both use the flat grid value. The consolidated
path is self-consistent on the register's constant; the strawman-side divergence is
recorded, not touched (SRA-02 Amendment 10, A10.1 / A10.7).

Pure-data layer: dicts in, dicts out, no rendering.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Any, Iterable, Mapping

import pandas as pd

from src.sra.mode_multipliers import resolve_mode_multiplier

#: The four figures, in ``CALIBRATION_LABELS`` order. These key names are the
#: contract with ``strawman_reader.get_strawman_calibration_target`` - the two sides
#: are tied together by ``test_comparator_keys_match_the_reader_contract`` rather
#: than by one importing the other's private.
COMPARATOR_KEYS: tuple[str, ...] = (
    "comparator_1_emv_days_pert",
    "comparator_1_emv_dollars",
    "comparator_2_finish_delay_days_pert",
    "comparator_2_finish_delay_dollars",
)

#: Keys of :func:`evaluate_schedule_calibration`'s result that a report surface
#: renders. Named here so a consumer's table and this function cannot drift apart
#: silently.
CALIBRATION_REPORT_KEYS: tuple[str, ...] = (
    "mc_schedule_mean_days",
    "comparator_1_emv_days_pert",
    "comparator_2_finish_delay_days_pert",
    "mode_multipliers_applied",
    "judged_against",
    "absolute_gap_days",
    "absolute_tolerance_days",
    "passes_calibration",
    "diagnostic_comparator",
    "diagnostic_gap_days",
    "interpretation",
)

#: The register's own methodology grid sheet, and the label column's mode names.
_METHODOLOGY_SHEET_TITLE = "Methodology"
_TIME_IMPACT_YES = "Yes"
#: The scaffolder's day-rate named range (``tools.register_scaffold.DAY_RATE_NAME``).
#: Kept as this module's own literal rather than imported - ``src/`` does not depend
#: on ``tools/`` - and tied to the scaffolder's by a test.
_DAY_RATE_NAME = "MC_day_rate"


def _as_float(value: Any) -> float:
    """Blank / missing / non-numeric -> 0.0, matching Excel's treatment of an empty
    cell in an arithmetic expression. The Excel side reaches the same result by the
    same route, so a half-filled workshop row contributes zero on both sides rather
    than diverging."""
    if value is None:
        return 0.0
    try:
        out = float(value)
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if pd.isna(out) else out


# RS-D-3 (D-15): the resolver is shared by every multiplier lookup and lives in
# src/sra/mode_multipliers.py; this module keeps its private name for its callers.
_resolve_mode_multiplier = resolve_mode_multiplier


def compute_schedule_comparators(
    risk_scope_df: pd.DataFrame,
    mode_multipliers: Mapping[str, float],
    day_rate: float,
) -> dict[str, float]:
    """Compute the four deterministic comparators from a ``read_risk_scope`` frame.

    Parameters
    ----------
    risk_scope_df
        The consolidated scope frame. Only ``time_impact``, ``mode_aggregated``,
        ``probability`` and the three ``*_days`` columns are read.
    mode_multipliers
        ``{mode_name: multiplier}`` from the register's own Methodology grid.
    day_rate
        Cost of delay per working day, from the register.

    Returns
    -------
    dict keyed by :data:`COMPARATOR_KEYS`. An empty frame - or one with no Time=Yes
    rows - returns 0.0 for all four rather than raising or returning None: a register
    with no schedule risks has a cross-check of zero, which is a fact, not an absence.

    Raises ``ValueError`` naming the risk when a Time=Yes row carries a Mode the
    grid does not define - blank, missing, or simply absent from the table. See
    :func:`_resolve_mode_multiplier`: this is D-14, and it is the one place the
    Python twin deliberately diverges from the Excel formula's
    ``IFERROR(INDEX(MATCH(...)), 0)``. A cost-only register never reaches it.
    """
    required = ("time_impact", "mode_aggregated", "probability",
                "opt_days", "ml_days", "pess_days")
    missing = [c for c in required if c not in risk_scope_df.columns]
    if missing:
        raise KeyError(
            f"risk scope frame is missing {missing}; "
            f"got columns {list(risk_scope_df.columns)}"
        )

    comparator_1_days = 0.0
    comparator_2_days = 0.0
    for _, row in risk_scope_df.iterrows():
        if str(row["time_impact"]).strip().casefold() != _TIME_IMPACT_YES.casefold():  # RS-D-4
            continue
        opt = _as_float(row["opt_days"])
        ml = _as_float(row["ml_days"])
        pess = _as_float(row["pess_days"])
        emv_days = _as_float(row["probability"]) * (opt + 4.0 * ml + pess) / 6.0
        comparator_1_days += emv_days
        comparator_2_days += emv_days * _resolve_mode_multiplier(
            row["mode_aggregated"], mode_multipliers, row.get("risk_id"))

    rate = _as_float(day_rate)
    return {
        COMPARATOR_KEYS[0]: comparator_1_days,
        COMPARATOR_KEYS[1]: comparator_1_days * rate,
        COMPARATOR_KEYS[2]: comparator_2_days,
        COMPARATOR_KEYS[3]: comparator_2_days * rate,
    }


#: The one injection method under which Comparator 1 bounds a Mode-off mean from above: one sampled
#: delay per risk, written once, to the gating activity (SCHEDCAL-01 Gate 3 round 2, Director
#: 14/10/2026). ``additive_pert_days`` writes the same delay to every activity a risk binds, so a
#: serial cohort moves the finish by a multiple of it; the multiplicative methods scale a factor by
#: each activity's duration, which Comparator 1 sums as if it were days.
BOUNDED_INJECTION_METHOD = "discrete_event_additive_days"


def evaluate_schedule_calibration(
    mc_schedule_mean_days: float,
    comparators: Mapping[str, float],
    tolerance_days: float,
    *,
    mode_multipliers_applied: bool,
    injection_methods: Iterable[str],
) -> dict[str, Any]:
    """Compare an MC schedule arm's mean finish delay against the comparators, on its own basis.

    The consolidated path's calibration check (P-E4), judged on the run's Mode basis
    (SCHEDCAL-01, Director 14/10/2026):

    * **Mode multipliers applied** (Stage 2.5): the mean estimates **Comparator 2**, the
      Mode-weighted EMV sum, and must agree with it within the tolerance either side. Unchanged.
    * **Mode multipliers off** (Stage 3, the headline: ``STAGE3_APPLY_MODE_MULTIPLIERS = False``),
      with every simulated risk injected by ``discrete_event_additive_days``: each risk's sampled
      delay is written once, to its gating activity, and the finish is the longest path, so no
      network can push the mean above **Comparator 1**, the unweighted EMV sum. The check is
      one-sided (Gate 3, Director 14/10/2026): it fails only on a mean above Comparator 1 by more
      than the tolerance. A mean below it is delay that did not reach the finish - float absorbed,
      or a risk bound to no activity - which is always reported and never decides.
    * **Mode multipliers off with any other injection method**: Comparator 1 is no bound (see
      :data:`BOUNDED_INJECTION_METHOD`), so the run is **not assessed** (``passes_calibration`` is
      ``None``); the comparators and the gap are still reported (Gate 3 round 2, Director 14/10/2026).

    ``mode_multipliers_applied`` and ``injection_methods`` have no default: they are the run's own,
    taken from the config and the drivers it ran on, and a call that does not state them is
    refused. A mean that is missing or not finite is refused too, rather than read as zero.
    ``tolerance_days`` comes from the caller's config
    (``IntegratedConfig.schedule_calibration_tolerance_days``) rather than a literal here.

    Consumes the **Python** comparators only. The register's ``Summary`` cells are an
    audit surface proven equal by the reconciliation gate; they are never a run-time
    input, because a cached value is only as fresh as the last Excel save (P-E8).

    Returns a dict; never raises on a failing check and never gates - reporting a
    disagreement is the job, and a check that raised would take the deliverable down
    over a methodology question.
    """
    if not isinstance(mode_multipliers_applied, bool):
        raise TypeError("mode_multipliers_applied must be the run's own bool setting, "
                        f"not {mode_multipliers_applied!r}")
    if isinstance(injection_methods, str):
        raise TypeError("injection_methods must be the collection of the run's drivers' methods, "
                        "not one string")
    methods = sorted(set(injection_methods))
    if not methods or not all(isinstance(m, str) and m for m in methods):
        raise TypeError("injection_methods must name the run's drivers' injection methods, "
                        f"not {injection_methods!r}")
    if mc_schedule_mean_days is None or not math.isfinite(float(mc_schedule_mean_days)):
        raise ValueError(f"the MC schedule mean is {mc_schedule_mean_days!r}; a missing or "
                         "non-finite mean is a broken run, not a calibration result")
    comparator_1 = _as_float(comparators.get(COMPARATOR_KEYS[0]))
    comparator_2 = _as_float(comparators.get(COMPARATOR_KEYS[2]))
    mc_mean = _as_float(mc_schedule_mean_days)
    tolerance = _as_float(tolerance_days)
    # Where the MC mean sits relative to the two comparators: Comparator 1 assumes every delay
    # reaches the finish date and Comparator 2 assumes the Mode weighting is exact, so on the
    # Mode-on basis they bracket what a network should produce.
    bracketed = min(comparator_1, comparator_2) <= mc_mean <= max(comparator_1, comparator_2)

    if mode_multipliers_applied:
        gap = abs(mc_mean - comparator_2)
        passes = gap <= tolerance
        interpretation = (
            f"MC schedule mean = {mc_mean:.3f} days vs deterministic Comparator 2 = "
            f"{comparator_2:.3f} days. Absolute gap = {gap:.3f} days "
            f"(tolerance {tolerance:.1f} days) - "
            f"{'PASS' if passes else 'FAIL'}. Comparator 2 is the register's "
            f"Mode-weighted expected-value sum, a closed-form estimate of "
            f"finish-date slippage; the MC mean is the same quantity derived from "
            f"the CPM network. Agreement inside tolerance means the Mode "
            f"classification and the network price float the same way. "
            + (
                f"The MC mean sits between Comparator 1 ({comparator_1:.3f} days, "
                f"no Mode weighting) and Comparator 2, so the network is absorbing "
                f"float - but less of it than the Mode weighting assumed."
                if bracketed else
                f"The MC mean sits outside the Comparator 1 ({comparator_1:.3f} "
                f"days) to Comparator 2 bracket, which the two comparators should "
                f"span; treat the Mode classification or the network binding as "
                f"suspect before the tolerance."
            )
        )
        return {
            "mc_schedule_mean_days": mc_mean,
            "comparator_1_emv_days_pert": comparator_1,
            "comparator_2_finish_delay_days_pert": comparator_2,
            "mode_multipliers_applied": True,
            "injection_methods": methods,
            "judged_against": "Comparator 2",
            "absolute_gap_days": gap,
            "absolute_tolerance_days": tolerance,
            "passes_calibration": bool(passes),
            "diagnostic_comparator": "Comparator 1",
            "diagnostic_gap_days": abs(mc_mean - comparator_1),
            "mc_mean_within_comparator_bracket": bool(bracketed),
            "interpretation": interpretation,
        }

    excess = max(mc_mean - comparator_1, 0.0)
    shortfall = max(comparator_1 - mc_mean, 0.0)
    expected = comparator_1 - comparator_2
    bounded = methods == [BOUNDED_INJECTION_METHOD]
    passes = (excess <= tolerance) if bounded else None
    head = (f"MC schedule mean = {mc_mean:.3f} days against Comparator 1 = {comparator_1:.3f} days "
            f"(no Mode weighting). ")
    if not bounded:
        interpretation = head + (
            f"NOT ASSESSED: the run injects by {', '.join(methods)}, and Comparator 1 bounds a "
            f"run without Mode multipliers only when every risk injects by "
            f"{BOUNDED_INJECTION_METHOD} (one delay, written once). Gap to Comparator 1 = "
            f"{mc_mean - comparator_1:+.3f} days; Comparator 2 = {comparator_2:.3f} days. "
            f"Reported, not judged."
        )
    else:
        interpretation = head + (
            f"Comparator 1 is the upper bound for this run: every risk injects one delay, written "
            f"once, so no path can carry more than the sum. Excess over the bound = {excess:.3f} "
            f"days (tolerance {tolerance:.1f} days) - {'PASS' if passes else 'FAIL'}. "
            + (
                f"Delay that did not reach the finish: {shortfall:.1f} days (float absorbed, or "
                f"risks bound to no activity); the register's Mode labels expect about "
                f"{expected:.1f} (Comparator 2 = {comparator_2:.3f} days). Diagnostic, not part "
                f"of the check."
                if passes else
                "The mean exceeds the sum of every risk's expected delay, which this injection "
                "cannot add; treat the network binding or the risk inputs as suspect before the "
                "tolerance."
            )
        )
    return {
        "mc_schedule_mean_days": mc_mean,
        "comparator_1_emv_days_pert": comparator_1,
        "comparator_2_finish_delay_days_pert": comparator_2,
        "mode_multipliers_applied": False,
        "injection_methods": methods,
        "assessed": bounded,
        "judged_against": "Comparator 1 (upper bound)" if bounded else "not assessed",
        "absolute_gap_days": excess,
        "absolute_tolerance_days": tolerance,
        "passes_calibration": passes,
        "diagnostic_comparator": "Comparator 2",
        "diagnostic_gap_days": abs(mc_mean - comparator_2),
        "delay_not_reaching_finish_days": shortfall,
        "mc_mean_within_comparator_bracket": bool(bracketed),
        "interpretation": interpretation,
    }


def read_register_day_rate(register_path: str | Path) -> float:
    """Read the ``MC_day_rate`` named range - the same cell the Excel dollar
    formulas multiply by, so both sides dollarise off one source."""
    from openpyxl import load_workbook

    wb = load_workbook(str(register_path), data_only=True)
    try:
        if _DAY_RATE_NAME not in wb.defined_names:
            raise ValueError(
                f"register '{register_path}' does not define the "
                f"{_DAY_RATE_NAME!r} named range"
            )
        destinations = list(wb.defined_names[_DAY_RATE_NAME].destinations)
        if len(destinations) != 1:
            raise ValueError(
                f"{_DAY_RATE_NAME} resolved to {len(destinations)} destinations; "
                "expected exactly one cell"
            )
        sheet, addr = destinations[0]
        return _as_float(wb[sheet][addr].value)
    finally:
        wb.close()


def read_register_mode_multipliers(register_path: str | Path) -> dict[str, float]:
    """Read ``{mode: multiplier}`` from the register's own Methodology grid.

    The same grid the ``Summary`` formulas resolve through ``INDEX/MATCH``, read
    through the shared parser so both sides see one source. Fail-loud, per the
    no-defaults directive: a register without a readable Methodology grid raises
    rather than substituting a cross-project constant.
    """
    from openpyxl import load_workbook

    from src.qra_register_parse import resolve_methodology_from_grid

    wb = load_workbook(str(register_path), read_only=True, data_only=True)
    try:
        if _METHODOLOGY_SHEET_TITLE not in wb.sheetnames:
            raise ValueError(
                f"register '{register_path}' has no {_METHODOLOGY_SHEET_TITLE} sheet; "
                "the Mode multipliers the schedule cross-check needs live there"
            )
        grid = [
            list(r)
            for r in wb[_METHODOLOGY_SHEET_TITLE].iter_rows(values_only=True)
        ]
    finally:
        wb.close()
    return resolve_methodology_from_grid(grid, source=_METHODOLOGY_SHEET_TITLE)[
        "mode_multipliers"
    ]
