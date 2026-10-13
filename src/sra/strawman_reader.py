"""Strawman workbook reader for the SRA workstream (SR5).

Reads the project's Phase 2 strawman workbook
and exposes its content as DataFrames and dicts. The strawman is the
canonical input source for the SR5 schedule MC engine: it carries the
risk-to-pattern assignments, the per-risk PERT parameters, the run
configuration (day rate, programme dates, LD parameters), and the
calibration targets (Comparator 1 and Comparator 2).

Pure-data layer. No rendering. Header-driven column lookup throughout
(no hardcoded row indices for data, only for clearly-labelled section
headers in the strawman).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


# ----------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------

def _normalise(value: Any) -> str:
    """Lowercase and strip a header cell value for matching."""
    if value is None:
        return ""
    return str(value).strip().lower().replace("\n", " ")


def _find_header_row(ws, expected_first_header: str, max_scan_rows: int = 20) -> int:
    """Scan top rows to find the row containing expected_first_header in column A.

    Returns 1-based row index. Raises ValueError if not found within
    max_scan_rows.
    """
    target = expected_first_header.strip().lower()
    for r_idx in range(1, max_scan_rows + 1):
        cell = ws.cell(row=r_idx, column=1).value
        if cell is None:
            continue
        if str(cell).strip().lower() == target:
            return r_idx
    raise ValueError(
        f"Header row with '{expected_first_header}' not found in first "
        f"{max_scan_rows} rows of sheet '{ws.title}'"
    )


#: Header-map fields matched exactly, never by prefix.
_EXACT_HEADER_FIELDS: frozenset[str] = frozenset({"probability"})


def _build_col_map(ws, header_row: int, expected_headers: dict[str, list[str]]) -> dict[str, int]:
    """Return {field_name: 0-based column index} from header_row.

    expected_headers maps each canonical field name to a list of acceptable
    header strings (case-insensitive, prefix-matched).
    """
    raw = list(ws.iter_rows(
        min_row=header_row, max_row=header_row, values_only=True,
    ))[0]
    norm = [_normalise(v) for v in raw]

    col_map: dict[str, int] = {}
    for field, variants in expected_headers.items():
        for variant in variants:
            v = variant.strip().lower()
            for idx, h in enumerate(norm):
                if h == v:
                    col_map[field] = idx
                    break
            if field in col_map:
                break
        # prefix match fallback - never for an exact-only field, where a prefix would find a
        # different column ("Probability" is a prefix of "Probability_Time").
        if field not in col_map and field not in _EXACT_HEADER_FIELDS:
            for variant in variants:
                v = variant.strip().lower()
                for idx, h in enumerate(norm):
                    if h.startswith(v):
                        col_map[field] = idx
                        break
                if field in col_map:
                    break
    return col_map


def _coerce_bool(value: Any) -> bool:
    """Coerce strawman bool/text representations to Python bool."""
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    s = str(value).strip().lower()
    return s in ("true", "yes", "y", "1")


def _coerce_str(value: Any) -> str:
    """Coerce a cell value to a non-None string (empty string for None)."""
    if value is None:
        return ""
    return str(value).strip()


def canonical_time_impact(value: Any) -> str:
    """Canonical ``Time_Impact`` cell: any case/whitespace spelling of yes becomes ``"Yes"``.

    RS-D-4 (D-16): Excel counts ``<>"Yes"`` case-insensitively, so the Python side must
    not let ``yes``/``YES``/`` YeS `` decide differently which rows simulate. Canonicalised
    once here, at the reader, so every consumer compares a canonical value. Anything that is
    not a spelling of yes or no is returned stripped, unchanged, so an odd value stays
    visible rather than being folded into ``No``.
    """
    text = _coerce_str(value)
    folded = text.casefold()
    if folded == "yes":
        return "Yes"
    if folded == "no":
        return "No"
    return text


def _coerce_float(value: Any) -> float | None:
    """Coerce to float; None stays None."""
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _coerce_int_or_none(value: Any) -> int | None:
    """Coerce to int; unparseable or empty stays None."""
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _check_path(path: str | Path) -> Path:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Strawman file not found: {p}")
    return p


# ----------------------------------------------------------------------------
# public API: Risk_Activity_Patterns
# ----------------------------------------------------------------------------

_PATTERN_HEADERS: dict[str, list[str]] = {
    "pattern_id": ["Pattern_ID", "PatternID"],
    "description": ["Description"],
    "prefix_list": ["Prefix_List", "PrefixList"],
    "is_cost_only": ["Is_Cost_Only", "IsCostOnly"],
    "default_mapping_confidence": ["Default_Mapping_Confidence", "DefaultMappingConfidence"],
    "reconciliation_treatment": ["Reconciliation_Treatment", "ReconciliationTreatment"],
    "example_risk_ids": ["Example_Risk_IDs", "ExampleRiskIDs"],
    "notes": ["Notes"],
}


def read_risk_activity_patterns(strawman_path: str | Path) -> pd.DataFrame:
    """Read the Risk_Activity_Patterns sheet.

    Returns one row per pattern with the canonical column set in
    _PATTERN_HEADERS. Headers are in row 1 of the sheet.
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        ws = wb["Risk_Activity_Patterns"]
        header_row = _find_header_row(ws, "Pattern_ID")
        col_map = _build_col_map(ws, header_row, _PATTERN_HEADERS)
        for required in ("pattern_id", "description"):
            if required not in col_map:
                raise ValueError(
                    f"Required column '{required}' missing from "
                    f"Risk_Activity_Patterns header row"
                )

        rows: list[dict[str, Any]] = []
        for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
            pid_idx = col_map["pattern_id"]
            if pid_idx >= len(raw) or raw[pid_idx] is None:
                continue
            row: dict[str, Any] = {
                "pattern_id": _coerce_str(raw[col_map["pattern_id"]]),
                "description": _coerce_str(raw[col_map["description"]]),
            }
            row["prefix_list"] = (
                _coerce_str(raw[col_map["prefix_list"]])
                if "prefix_list" in col_map else ""
            )
            row["is_cost_only"] = (
                _coerce_bool(raw[col_map["is_cost_only"]])
                if "is_cost_only" in col_map else False
            )
            row["default_mapping_confidence"] = (
                _coerce_str(raw[col_map["default_mapping_confidence"]])
                if "default_mapping_confidence" in col_map else ""
            )
            row["reconciliation_treatment"] = (
                _coerce_str(raw[col_map["reconciliation_treatment"]])
                if "reconciliation_treatment" in col_map else "A"
            )
            row["example_risk_ids"] = (
                _coerce_str(raw[col_map["example_risk_ids"]])
                if "example_risk_ids" in col_map else ""
            )
            row["notes"] = (
                _coerce_str(raw[col_map["notes"]]) if "notes" in col_map else ""
            )
            rows.append(row)
        return pd.DataFrame(rows, columns=list(_PATTERN_HEADERS.keys()))
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# public API: Risk_Pattern_Assignments
# ----------------------------------------------------------------------------

_ASSIGNMENT_HEADERS: dict[str, list[str]] = {
    "risk_id": ["Risk_ID", "Risk ID"],
    "pattern_id": ["Pattern_ID", "Pattern ID"],
    "confidence_override": ["Assignment_Confidence", "Confidence_Override", "Confidence Override"],
    "notes": ["Notes"],
}


def read_risk_pattern_assignments(strawman_path: str | Path) -> pd.DataFrame:
    """Read the Risk_Pattern_Assignments sheet.

    Returns one row per Open risk mapping risk_id to pattern_id.
    Headers are in row 1 of the sheet.
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        ws = wb["Risk_Pattern_Assignments"]
        header_row = _find_header_row(ws, "Risk_ID")
        col_map = _build_col_map(ws, header_row, _ASSIGNMENT_HEADERS)
        for required in ("risk_id", "pattern_id"):
            if required not in col_map:
                raise ValueError(
                    f"Required column '{required}' missing from "
                    f"Risk_Pattern_Assignments header row"
                )

        rows: list[dict[str, Any]] = []
        for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
            rid_idx = col_map["risk_id"]
            if rid_idx >= len(raw) or raw[rid_idx] is None:
                continue
            rid = _coerce_str(raw[rid_idx])
            if not rid:
                continue
            row = {
                "risk_id": rid,
                "pattern_id": _coerce_str(raw[col_map["pattern_id"]]),
                "confidence_override": (
                    _coerce_str(raw[col_map["confidence_override"]])
                    if "confidence_override" in col_map else ""
                ),
                "notes": (
                    _coerce_str(raw[col_map["notes"]]) if "notes" in col_map else ""
                ),
            }
            rows.append(row)
        return pd.DataFrame(rows, columns=list(_ASSIGNMENT_HEADERS.keys()))
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# public API: Risk_Scope (the SR5 simulation source of truth)
# ----------------------------------------------------------------------------

# The variant lists carry BOTH the strawman header (underscore form, matched
# first/exact) AND the consolidated showcase register's themed header (REG-01).
# The register's time-section day headers are suffixed "(days)" so they read
# unambiguously here while still prefix-matching v5_register_reader's Pass-3
# ("optimistic"/"most likely"/"pessimistic"/"probability") - cost-side
# bit-identity is preserved. Register-only variants are inert on a strawman
# read (the exact underscore header is found first).
_SCOPE_HEADERS: dict[str, list[str]] = {
    "risk_id": ["Risk_ID", "Risk ID"],
    "risk_title": ["Risk_Title", "Risk Description"],
    "risk_owner": ["Risk_Owner", "Risk Owner"],
    "rbs_l1": ["RBS_L1", "RBS LVL 1"],
    "time_impact": ["Time_Impact", "Time Impact"],
    "dual_coded": ["Dual_Coded"],
    "pattern_id": ["Pattern_ID"],
    "pattern_description": ["Pattern_Description"],
    "prefix_list": ["Prefix_List"],
    "is_cost_only": ["Is_Cost_Only"],
    "activity_count": ["Activity_Count"],
    "min_total_float": ["Min_Total_Float"],
    "mode_aggregated": ["Mode_Aggregated", "Mode"],
    "series_source": ["Series_Source"],
    # SPEC R-10 (METHFIRE-01): a risk has one probability, and both arms fire on it.
    # Exact-match only (_EXACT_HEADER_FIELDS).
    "probability": ["Probability"],
    # Retired as a firing probability by R-10. Still read, for one purpose: a register
    # whose Time Impact = Yes row disagrees with Probability is refused
    # (_refuse_unequal_probabilities) until PROBTIME-01 removes the column.
    "probability_time": ["Probability_Time", "Probability (Time)"],
    "opt_days": ["Opt_Days", "Optimistic (days)"],
    "ml_days": ["ML_Days", "Most Likely (days)"],
    "pess_days": ["Pess_Days", "Pessimistic (days)"],
    "ld_considered": ["LD_Considered"],
    "day_rate": ["Day_Rate"],
    "reconciliation_treatment": ["Reconciliation_Treatment"],
    "emv_days_simple": ["EMV_days_Simple"],
    "emv_days_pert": ["EMV_days_PERT"],
    "emv_cost_simple": ["EMV_cost_Simple"],
    "emv_cost_pert": ["EMV_cost_PERT"],
    "finish_delay_contribution": ["Finish_delay_contribution"],
    "finish_delay_days": ["Finish_delay_days"],
    "finish_delay_dollars": ["Finish_delay_dollars"],
    "review_flag": ["Review_Flag"],
    # SR12 — per-risk injection-method dispatch. All three columns are
    # optional from the strawman's perspective: legacy strawmans without
    # them produce ``""`` (read as null/blank by the binding layer and
    # mapped to the pre-SR12 ``additive_pert_days`` fallback).
    "injection_method": ["Injection_Method", "InjectionMethod"],
    "gating_activity_id": ["Gating_Activity_ID", "GatingActivityID"],
    "injection_method_notes": [
        "Injection_Method_Notes", "InjectionMethodNotes",
    ],
}


# REG-01 §3: CPM binding columns the consolidated register carries on its
# machine-facing "Mapping" sheet (dissolved from the strawman). Joined onto
# the scope frame by Risk ID in the consolidated read path.
_MAPPING_HEADERS: dict[str, list[str]] = {
    "injection_method": ["Injection_Method", "InjectionMethod"],  # RS-D-6
    "risk_id": ["Risk_ID", "Risk ID"],
    "pattern_id": ["Pattern_ID", "Pattern ID"],
    "prefix_list": ["Prefix_List", "Prefix List"],
    "activity_count": ["Activity_Count", "Activity Count"],
    # SRA-02 S-C2. The scaffolder has emitted this column since REG03
    # (tools/register_scaffold.py:182-183) and nothing read it, so every risk on
    # a consolidated register resolved gating to None and fell to
    # _auto_pick_gating_activity. Declared here so the workshop's explicit gate
    # is the authoritative one, per SRA-01 D-A.
    "gating_activity_id": ["Gating_Activity_ID", "Gating Activity ID"],
}


def _find_header_row_any(ws, candidates: list[str], max_scan_rows: int = 20) -> int:
    """Like _find_header_row but accepts any of several column-A header spellings."""
    targets = {c.strip().lower() for c in candidates}
    for r_idx in range(1, max_scan_rows + 1):
        cell = ws.cell(row=r_idx, column=1).value
        if cell is not None and str(cell).strip().lower() in targets:
            return r_idx
    raise ValueError(
        f"None of {candidates} found in first {max_scan_rows} rows of "
        f"sheet '{ws.title}'"
    )


def _read_mapping_sheet(wb) -> dict[str, dict[str, Any]]:
    """Read the consolidated register's Mapping sheet -> {risk_id: {...}}."""
    ws = wb["Mapping"]
    header_row = _find_header_row_any(ws, ["Risk_ID", "Risk ID"])
    col_map = _build_col_map(ws, header_row, _MAPPING_HEADERS)
    if "risk_id" not in col_map:
        raise ValueError("Mapping sheet missing a Risk ID column")
    out: dict[str, dict[str, Any]] = {}
    for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rid_idx = col_map["risk_id"]
        if rid_idx >= len(raw) or raw[rid_idx] is None:
            continue
        rid = _coerce_str(raw[rid_idx])
        if not rid:
            continue
        out[rid] = {
            "pattern_id": (_coerce_str(raw[col_map["pattern_id"]])
                           if "pattern_id" in col_map else ""),
            "prefix_list": (_coerce_str(raw[col_map["prefix_list"]])
                            if "prefix_list" in col_map else ""),
            "activity_count": (raw[col_map["activity_count"]]
                               if "activity_count" in col_map else None),
            # "" when the column is absent or the cell blank; the join below
            # converts that to None so an absent gate is absent rather than an
            # empty string masquerading as a value.
            "gating_activity_id": (_coerce_str(raw[col_map["gating_activity_id"]])
                                   if "gating_activity_id" in col_map else ""),
            "injection_method": (_coerce_str(raw[col_map["injection_method"]])
                                 if "injection_method" in col_map else ""),
        }
    return out


def _refuse_unequal_probabilities(scope: pd.DataFrame, sheet: str) -> pd.DataFrame:
    """Refuse a scope frame whose Time Impact = Yes rows give one risk two probabilities.

    SPEC R-10 (METHFIRE-01, D-2): a risk is one event with one probability, ``Probability``,
    and both arms fire on it. ``Probability (Time)`` is no longer read as a probability, so a
    row where it holds a different value would silently change meaning. Refused instead, every
    offending row named. Transitional: PROBTIME-01 removes the column and this with it.
    """
    time_yes = scope[scope["time_impact"] == "Yes"]
    blank = [str(r) for r, p in zip(time_yes["risk_id"], time_yes["probability"]) if p is None or pd.isna(p)]
    if blank:
        raise ValueError(
            f"sheet '{sheet}': Time Impact = Yes risks {blank} have no Probability. A risk has one "
            f"probability, Probability, and both arms fire on it (SPEC R-10) - including a risk that "
            f"only delays: give it its Probability and leave its dollar estimates at zero."
        )
    unequal = [
        f"{rid} (Probability={p!r}, Probability (Time)={pt!r})"
        for rid, p, pt in zip(time_yes["risk_id"], time_yes["probability"], time_yes["probability_time"])
        if pt is not None and not pd.isna(pt) and float(pt) != float(p)
    ]
    if unequal:
        raise ValueError(
            f"sheet '{sheet}': these Time Impact = Yes risks give Probability (Time) a different "
            f"value from Probability: {'; '.join(unequal)}. A risk has one probability and both arms "
            f"fire on it (SPEC R-10). If the second value is the chance a risk that occurs also "
            f"delays, put that in the delay estimate; if it describes a different event, give that "
            f"event its own row. Then set Probability (Time) equal to Probability, or clear it."
        )
    return scope


def _read_risk_scope_from_register(wb, sheet_name: str) -> pd.DataFrame:
    """REG-01 consolidated path: assemble the scope frame from the themed
    "Risk Register" sheet (schedule three-points + Mode) joined with the
    "Mapping" sheet (Pattern_ID / Prefix_List / Activity_Count).

    Bit-identity note: the schedule three-points read here are the register's
    time-section columns (headers suffixed "(days)"), the SAME values the
    dissolved strawman's Risk_Scope carried - both were written from one
    source (build_synthetic_wind_fixture.WIND_RISKS).
    """
    ws = wb[sheet_name]
    header_row = _find_header_row(ws, "Risk ID")
    col_map = _build_col_map(ws, header_row, _SCOPE_HEADERS)
    for required in ("risk_id", "time_impact", "mode_aggregated",
                     "probability", "opt_days", "ml_days", "pess_days"):
        if required not in col_map:
            raise ValueError(
                f"Required column '{required}' missing from consolidated "
                f"register sheet '{sheet_name}' header row {header_row}"
            )

    mapping = _read_mapping_sheet(wb) if "Mapping" in wb.sheetnames else {}

    from src.sra.cpm.bindings import active_risk_id_prefix
    _prefix = active_risk_id_prefix().upper()

    rows: list[dict[str, Any]] = []
    for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rid_idx = col_map["risk_id"]
        if rid_idx >= len(raw) or raw[rid_idx] is None:
            continue
        rid = _coerce_str(raw[rid_idx])
        if not rid or not rid.upper().startswith(_prefix):
            continue

        row: dict[str, Any] = {}
        for field in _SCOPE_HEADERS:
            if field not in col_map:
                row[field] = None
                continue
            value = raw[col_map[field]] if col_map[field] < len(raw) else None
            if field in ("is_cost_only", "dual_coded"):
                row[field] = _coerce_bool(value)
            elif field in ("activity_count", "min_total_float"):
                row[field] = None if value in (None, "") else _coerce_int_or_none(value)
            elif field in ("probability", "probability_time", "opt_days", "ml_days", "pess_days",
                           "day_rate", "emv_days_simple", "emv_days_pert",
                           "emv_cost_simple", "emv_cost_pert",
                           "finish_delay_contribution", "finish_delay_days",
                           "finish_delay_dollars"):
                row[field] = _coerce_float(value)
            elif field == "time_impact":
                row[field] = canonical_time_impact(value)
            else:
                row[field] = _coerce_str(value)

        # Join CPM binding columns from the Mapping sheet.
        m = mapping.get(rid, {})
        if m:
            row["pattern_id"] = m.get("pattern_id") or None
            row["prefix_list"] = m.get("prefix_list") or None
            ac = m.get("activity_count")
            row["activity_count"] = None if ac in (None, "") else _coerce_int_or_none(ac)
            # SRA-02 S-C2. Mapping is the authoritative binding surface (D-A), so
            # it wins over whatever the Risk Register sheet carried - which on a
            # scaffolded register is nothing, since col 31 "Schedule Activity ID
            # (P6)" is reserved-inert and matches no _SCOPE_HEADERS variant.
            # Absent stays absent: "" becomes None, never a fabricated gate.
            row["gating_activity_id"] = m.get("gating_activity_id") or None
            # RS-D-6: a per-risk method on the Mapping sheet; blank means the register's default.
            if m.get("injection_method"):
                row["injection_method"] = m["injection_method"]
        rows.append(row)
    return _refuse_unequal_probabilities(
        pd.DataFrame(rows, columns=list(_SCOPE_HEADERS.keys())), sheet_name)


def read_risk_scope(strawman_path: str | Path) -> pd.DataFrame:
    """Read the risk scope frame (one row per Open risk).

    Two layouts (REG-01 §3):

    * **Strawman**: a ``Risk_Scope`` sheet. Read as before.
    * **Consolidated register** (showcase): no ``Risk_Scope`` sheet; schedule
      three-points + Mode come from the themed ``Risk Register`` sheet and the
      CPM binding columns from the ``Mapping`` sheet.

    Returns the canonical columns in _SCOPE_HEADERS. Time_Impact filtering is
    applied at simulation time, not here.
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        if "Risk_Scope" not in wb.sheetnames:
            for candidate in ("Risk Register", "Risk_Register"):
                if candidate in wb.sheetnames:
                    return _read_risk_scope_from_register(wb, candidate)
            raise ValueError(
                f"workbook '{p}' has neither a Risk_Scope sheet nor a themed "
                "Risk Register sheet to read the risk scope from."
            )
        ws = wb["Risk_Scope"]
        header_row = _find_header_row(ws, "Risk_ID")
        col_map = _build_col_map(ws, header_row, _SCOPE_HEADERS)

        for required in ("risk_id", "time_impact", "mode_aggregated",
                         "probability", "opt_days", "ml_days", "pess_days"):
            if required not in col_map:
                raise ValueError(
                    f"Required column '{required}' missing from Risk_Scope "
                    f"header row {header_row}"
                )

        # Risk-ID prefix is project-specific (synthetic road "KH", synthetic wind
        # "WT"); resolve from the active binding so a non-R-prefixed strawman
        # loads without an engine edit. Lazy import avoids an import cycle.
        from src.sra.cpm.bindings import active_risk_id_prefix
        _prefix = active_risk_id_prefix().upper()

        rows: list[dict[str, Any]] = []
        for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
            rid_idx = col_map["risk_id"]
            if rid_idx >= len(raw) or raw[rid_idx] is None:
                continue
            rid = _coerce_str(raw[rid_idx])
            if not rid or not rid.upper().startswith(_prefix):
                continue

            row: dict[str, Any] = {}
            for field in _SCOPE_HEADERS:
                if field not in col_map:
                    row[field] = None
                    continue
                value = raw[col_map[field]] if col_map[field] < len(raw) else None
                if field in ("is_cost_only", "dual_coded"):
                    row[field] = _coerce_bool(value)
                elif field in ("activity_count", "min_total_float"):
                    if value is None or value == "":
                        row[field] = None
                    else:
                        try:
                            row[field] = int(value)
                        except (TypeError, ValueError):
                            row[field] = None
                elif field in ("probability", "probability_time", "opt_days", "ml_days", "pess_days",
                               "day_rate", "emv_days_simple", "emv_days_pert",
                               "emv_cost_simple", "emv_cost_pert",
                               "finish_delay_contribution", "finish_delay_days",
                               "finish_delay_dollars"):
                    row[field] = _coerce_float(value)
                elif field == "time_impact":
                    row[field] = canonical_time_impact(value)
                else:
                    row[field] = _coerce_str(value)
            rows.append(row)
        return _refuse_unequal_probabilities(
            pd.DataFrame(rows, columns=list(_SCOPE_HEADERS.keys())), "Risk_Scope")
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# public API: Run_Config_Preview
# ----------------------------------------------------------------------------

_RUN_CONFIG_KEY_MAP: dict[str, str] = {
    "programme_seasons": "programme_seasons",
    "cascading": "cascading",
    "chains": "chains",
    "iterations": "iterations",
    "p-level_reporting": "p_level_reporting",
    "reconciliation_default": "reconciliation_default",
    "cp_rule": "cp_rule",
    "cp_hybrid_float_threshold_days": "cp_hybrid_float_threshold_days",
    "day_rate_default ($/wd)": "day_rate_default",
    "ld_layer_status": "ld_layer_status",
    "expected_time_yes_count": "expected_time_yes_count",
    "avg prelims per week ($)": "avg_prelims_weekly",
    "schedule duration (weeks)": "schedule_duration_weeks",
    "schedule start": "schedule_start",
    "schedule finish": "schedule_finish",
    "toc for lds ($)": "toc_for_lds",
    "ld burn rate ($/day)": "ld_burn_rate",
    "ld cap ($)": "ld_cap",
}


def _parse_on_off(value: Any) -> bool:
    """Convert strawman OFF/ON/True/False to a bool."""
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    s = str(value).strip().lower()
    if s in ("on", "true", "yes", "y", "1"):
        return True
    if s in ("off", "false", "no", "n", "0"):
        return False
    return False


# REG-01 §12.4: when the strawman is dissolved into the consolidated
# showcase register, its Run_Config_Preview folds into MC_Params. The
# config keys then live in MC_Params (label col B, value col C) under these
# labels. A separate strawman keeps its Run_Config_Preview sheet
# and is read by the branch above - byte-unchanged. Only workbooks lacking
# Run_Config_Preview fall through to this MC_Params map.
_MC_PARAMS_CONFIG_KEY_MAP: dict[str, str] = {
    "iterations (n)": "iterations",
    "reporting p-level (pct)": "p_level_reporting",
    "day_rate ($/wd)": "day_rate_default",
    "programme_seasons": "programme_seasons",
    "cp_hybrid_float_threshold_days": "cp_hybrid_float_threshold_days",
    "schedule duration (weeks)": "schedule_duration_weeks",
    "expected_time_yes_count": "expected_time_yes_count",
}


def _coerce_run_config_value(target_key: str, value: Any) -> Any:
    """Type-coerce one run-config value by its target key."""
    if target_key in ("cascading", "chains"):
        return _parse_on_off(value)
    if target_key in ("iterations", "p_level_reporting"):
        # CONFIGTRUTH-01: passed as read, so the run-config contract can refuse a cell that is not a
        # whole number by name; int() here raised unnamed on "abc" and turned 2.5 into 2.
        return value
    if target_key in ("programme_seasons",
                      "cp_hybrid_float_threshold_days", "expected_time_yes_count"):
        return int(value) if value is not None else None
    if target_key in ("day_rate_default", "avg_prelims_weekly",
                      "schedule_duration_weeks", "toc_for_lds",
                      "ld_burn_rate", "ld_cap"):
        return _coerce_float(value)
    return value


def _walk_config_rows(ws, key_map: dict[str, str], label_idx: int,
                      value_idx: int) -> dict[str, Any]:
    """Walk a label/value sheet, mapping known labels to typed config values."""
    out: dict[str, Any] = {}
    for raw in ws.iter_rows(values_only=True):
        if not raw or len(raw) <= label_idx or raw[label_idx] is None:
            continue
        label_norm = _normalise(raw[label_idx])
        if label_norm in key_map:
            target_key = key_map[label_norm]
            value = raw[value_idx] if len(raw) > value_idx else None
            out[target_key] = _coerce_run_config_value(target_key, value)
    return out


def read_run_config(strawman_path: str | Path) -> dict:
    """Read run configuration as a flat dict.

    Two workbook layouts are supported:

    * **Strawman** (a separate workbook): a ``Run_Config_Preview`` sheet with
      Parameter/Value tables in columns A/B. Read byte-for-byte as before.
    * **Consolidated register** (REG-01 showcase): no ``Run_Config_Preview``;
      the config folded into ``MC_Params`` (label col B, value col C). The
      duplicated iterations / P-level values were reconciled to MC_Params at
      build time (REG-01 §12.4).

    Returns a dict with keys matching the config-key maps, plus defaults for
    any keys not present except ``iterations`` and ``p_level_reporting``, which
    stay absent (CONFIGTRUTH-01), and ``run_config_sheet``: the sheet read.
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        if "Run_Config_Preview" in wb.sheetnames:
            sheet = "Run_Config_Preview"
            out = _walk_config_rows(
                wb[sheet], _RUN_CONFIG_KEY_MAP,
                label_idx=0, value_idx=1,
            )
        elif "MC_Params" in wb.sheetnames:
            sheet = "MC_Params"
            out = _walk_config_rows(
                wb[sheet], _MC_PARAMS_CONFIG_KEY_MAP,
                label_idx=1, value_idx=2,
            )
        else:
            raise ValueError(
                f"workbook '{p}' has neither a Run_Config_Preview nor an "
                "MC_Params sheet to read run configuration from."
            )

        # A refusal names the sheet the row was looked for on (CONFIGTRUTH-01).
        out["run_config_sheet"] = sheet

        # apply defaults if missing. Not iterations or the reporting P-level (CONFIGTRUTH-01):
        # a missing row is left absent, so the run-config contract can resolve it from the
        # project's binding and record that it did, or refuse - never take a number from here.
        defaults = {
            "programme_seasons": 1,
            "cascading": False,
            "chains": False,
            "reconciliation_default": "A",
            "cp_rule": "Option_Z_Hybrid",
            "cp_hybrid_float_threshold_days": 20,
            "day_rate_default": None,
            "ld_layer_status": "OFF",
            "expected_time_yes_count": None,
            "avg_prelims_weekly": None,
            "schedule_duration_weeks": None,
            "schedule_start": None,
            "schedule_finish": None,
            "toc_for_lds": None,
            "ld_burn_rate": None,
            "ld_cap": None,
        }
        for k, v in defaults.items():
            if k not in out:
                out[k] = v
        return out
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# public API: Summary -> Comparator targets
# ----------------------------------------------------------------------------

_COMPARATOR_LABELS: dict[str, str] = {
    "comparator_1_emv_days_pert": "total emv days (pert)",
    "comparator_1_emv_dollars": "total emv $ cost-of-delay (pert)",
    "comparator_2_finish_delay_days_pert": "aggregate finish-date slippage days (pert)",
    "comparator_2_finish_delay_dollars": "aggregate finish-date slippage $ (pert)",
}


def get_strawman_calibration_target(strawman_path: str | Path) -> dict:
    """Read Comparator 1 and Comparator 2 figures from the Summary sheet.

    Walks column A looking for the four labelled rows. Cell references
    are recorded for audit.
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p), data_only=True)
    try:
        ws = wb["Summary"]
        out: dict[str, Any] = {
            "comparator_1_emv_days_pert": None,
            "comparator_1_emv_dollars": None,
            "comparator_2_finish_delay_days_pert": None,
            "comparator_2_finish_delay_dollars": None,
            "source_summary_cells": {},
        }
        # Take the first occurrence only. Summary contains both the upper
        # Comparator 1/2 panel (rows 4-11) and a per-RBS breakdown lower down
        # which can repeat the same labels with different aggregation.
        for row in ws.iter_rows():
            if not row or row[0].value is None:
                continue
            label_norm = _normalise(row[0].value)
            for target_key, expected in _COMPARATOR_LABELS.items():
                if out.get(target_key) is not None:
                    continue
                if label_norm == expected:
                    if len(row) > 1 and row[1].value is not None:
                        out[target_key] = _coerce_float(row[1].value)
                        out["source_summary_cells"][target_key] = row[1].coordinate
        return out
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# public API: Treatments (Beta-19 §4 — action-implications input source)
# ----------------------------------------------------------------------------

_TREATMENTS_HEADERS: dict[str, list[str]] = {
    "risk_id": ["Risk_ID", "RiskID"],
    "treatment_description": ["Treatment_Description", "TreatmentDescription"],
    "effectiveness_pct": ["Effectiveness_Pct", "EffectivenessPct"],
    "effectiveness_confidence": [
        "Effectiveness_Confidence", "EffectivenessConfidence",
    ],
}


def read_treatments(strawman_path: str | Path) -> pd.DataFrame:
    """Read the Treatments sheet (Beta-19 §4).

    Returns one row per risk_id with treatment_description,
    effectiveness_pct, and effectiveness_confidence. Rows with empty
    effectiveness_pct render as "Workshop input required" downstream;
    populated rows feed into the action-implications table.

    Returns an empty DataFrame with the canonical four columns if the
    Treatments sheet does not exist (legacy strawman without the sheet).
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        if "Treatments" not in wb.sheetnames:
            return pd.DataFrame(columns=list(_TREATMENTS_HEADERS.keys()))
        ws = wb["Treatments"]
        try:
            header_row = _find_header_row(ws, "Risk_ID")
        except ValueError:
            return pd.DataFrame(columns=list(_TREATMENTS_HEADERS.keys()))
        col_map = _build_col_map(ws, header_row, _TREATMENTS_HEADERS)
        if "risk_id" not in col_map:
            return pd.DataFrame(columns=list(_TREATMENTS_HEADERS.keys()))

        rows: list[dict[str, Any]] = []
        for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
            rid_idx = col_map["risk_id"]
            if rid_idx >= len(raw) or raw[rid_idx] is None:
                continue
            rid = _coerce_str(raw[rid_idx])
            if not rid:
                continue

            def _cell(field: str, _raw=raw) -> Any:
                if field not in col_map:
                    return None
                idx = col_map[field]
                if idx >= len(_raw):
                    return None
                return _raw[idx]

            rows.append({
                "risk_id": rid,
                "treatment_description": _coerce_str(_cell("treatment_description")),
                "effectiveness_pct": _coerce_float(_cell("effectiveness_pct")),
                "effectiveness_confidence": _coerce_str(
                    _cell("effectiveness_confidence")
                ),
            })
        return pd.DataFrame(rows, columns=list(_TREATMENTS_HEADERS.keys()))
    finally:
        wb.close()


def write_treatments_placeholder_sheet(
    strawman_path: str | Path,
    risk_ids: list[str],
    overwrite: bool = False,
    recalc_via_excel: bool = True,
) -> dict[str, Any]:
    """Write a placeholder Treatments sheet to the strawman workbook.

    Used by the strawman builder to emit a per-risk placeholder row
    when the sheet is missing. Workshop output replaces the rows
    in-place; the schema (four columns) is the contract between the
    strawman and the action-implications composer.

    Parameters
    ----------
    strawman_path : Path
        Workbook to write into.
    risk_ids : list[str]
        One row per risk_id will be written.
    overwrite : bool
        If False (default) and a Treatments sheet already exists, this
        is a no-op and the existing sheet is preserved (workshop input
        is never destroyed). If True, the existing sheet is removed and
        rewritten in placeholder form.
    recalc_via_excel : bool
        If True (default), after writing the sheet with openpyxl this
        function opens the workbook in Excel via xlwings and forces a
        full recalc + save. This restores the cached formula values
        that openpyxl strips on save (Risk_Scope uses formulas that
        reference Risk_Register / Mapping_ARCHIVE / tblRiskPatternAssignments
        — without the cached values, downstream read_* helpers see
        None for every formula cell). Skip only on machines without
        Excel (recalc will need to happen manually on next open in
        Excel).

    Returns
    -------
    dict with keys ``rows_written`` (int), ``action`` (str), and
    optionally ``recalc`` ("excel" | "skipped" | "failed").
    """
    p = _check_path(strawman_path)
    wb = load_workbook(str(p))
    try:
        if "Treatments" in wb.sheetnames:
            if not overwrite:
                return {"rows_written": 0, "action": "skipped_existing"}
            del wb["Treatments"]
        ws = wb.create_sheet("Treatments")
        ws.cell(row=1, column=1, value="Risk_ID")
        ws.cell(row=1, column=2, value="Treatment_Description")
        ws.cell(row=1, column=3, value="Effectiveness_Pct")
        ws.cell(row=1, column=4, value="Effectiveness_Confidence")
        for i, rid in enumerate(risk_ids, start=2):
            ws.cell(row=i, column=1, value=str(rid))
            ws.cell(row=i, column=2, value="Workshop input required")
            ws.cell(row=i, column=3, value=None)
            ws.cell(row=i, column=4, value=None)
        wb.save(str(p))
    finally:
        wb.close()

    result: dict[str, Any] = {
        "rows_written": len(risk_ids),
        "action": "written",
    }
    if not recalc_via_excel:
        result["recalc"] = "skipped"
        return result
    try:
        import xlwings as xw  # noqa: WPS433  (local import; optional dep)

        app = xw.App(visible=False)
        try:
            xbook = app.books.open(str(p.resolve()))
            app.calculate()
            xbook.save()
            xbook.close()
            result["recalc"] = "excel"
        finally:
            app.quit()
    except Exception as exc:  # pragma: no cover
        result["recalc"] = "failed"
        result["recalc_error"] = str(exc)
    return result
