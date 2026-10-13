"""The earlier reference register's Risk Register reader for the SRA strawman.

Reads the `Risk Register` sheet and returns a list of normalised risk
dicts. Uses prefix-match header lookup because the earlier reference register headers have
trailing colons and help text (e.g. "Risk Description:\nRisk is an
effect or ...").

Does NOT key any behaviour on Risk ID literals - everything is driven
by narrative content, RBS L1 / L2, and recurrence type.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

RISK_REGISTER_SHEET = "Risk Register"
HEADER_ROW = 4
DATA_START_ROW = 5

# Canonical field name -> list of header-name prefixes to match against,
# case-insensitive. First-wins.
HEADER_PREFIXES: dict[str, tuple[str, ...]] = {
    "risk_id": ("risk id",),
    "risk_status": ("risk status",),
    "package": ("package",),
    "rbs_l1": ("rbs lvl 1", "rbs l1", "rbs level 1"),
    "rbs_l2": ("rbs lvl 2", "rbs l2", "rbs level 2"),
    "description": ("risk description",),
    "causes": ("causes",),
    "consequences": ("consequences",),
    "assumptions": ("assumptions",),
    "optimistic_scenario": ("optimistic scenario",),
    "likely_scenario": ("likely scenario",),
    "pessimistic_scenario": ("pessimistic scenario",),
    "recurrence_type": ("recurrence type",),
    "risk_owner": ("risk owner",),
    "control_effectiveness": ("control effectiveness",),
    "residual_rating": ("residual rating (mitigated)", "residual rating"),
    "probability": ("probability",),
    # The earlier reference register has three "Basis of Estimate" columns (Inherent/Contingent/Time).
    # The Contingent one (cost ML driver, 30/34 populated) is disambiguated
    # in the sheet as "Basis of Estimate7" - we match its full header so the
    # first-wins loop picks it and not the sparse Inherent column.
    "cost_basis_of_estimate": ("basis of estimate7",),
    "time_impact_flag": ("time impact",),
}


def _find_cols(header_row: tuple) -> dict[str, int]:
    """Return {canonical_field_key: 0-based column index}."""
    out: dict[str, int] = {}
    for idx, raw in enumerate(header_row):
        if raw is None:
            continue
        hl = str(raw).strip().lower()
        if not hl:
            continue
        for key, prefixes in HEADER_PREFIXES.items():
            if key in out:
                continue
            for pref in prefixes:
                if hl.startswith(pref):
                    out[key] = idx
                    break
    return out


def _str_or_none(v: object) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def read_risks(path: str | Path, open_only: bool = True) -> list[dict]:
    """Read the earlier reference register's Risk Register -> list of risk dicts.

    Parameters
    ----------
    path : the earlier reference register's xlsx path.
    open_only : if True (default), filter to Risk Status == 'Open'.

    Returns
    -------
    list[dict] with keys matching HEADER_PREFIXES plus any derived fields.
    """
    wb = load_workbook(str(path), read_only=True, data_only=True)
    if RISK_REGISTER_SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Sheet {RISK_REGISTER_SHEET!r} not found. Sheets: {wb.sheetnames}"
        )
    sht = wb[RISK_REGISTER_SHEET]
    rows = list(sht.iter_rows(values_only=True))
    wb.close()
    if len(rows) < DATA_START_ROW:
        return []

    header = rows[HEADER_ROW - 1]
    col = _find_cols(header)
    if "risk_id" not in col or "risk_status" not in col:
        raise ValueError(
            f"Required columns missing. Found: {sorted(col.keys())}"
        )

    out: list[dict] = []
    for raw in rows[DATA_START_ROW - 1:]:
        if raw is None or all(v is None for v in raw):
            continue
        rid = raw[col["risk_id"]] if col["risk_id"] < len(raw) else None
        if rid is None or not str(rid).strip():
            continue
        status = raw[col["risk_status"]] if col["risk_status"] < len(raw) else None
        if open_only and (not status or str(status).strip().lower() != "open"):
            continue

        risk = {}
        for key, idx in col.items():
            val = raw[idx] if idx < len(raw) else None
            risk[key] = _str_or_none(val) if isinstance(val, str) else val
        # Coerce Risk ID to a stripped string.
        risk["risk_id"] = str(risk["risk_id"]).strip()
        out.append(risk)

    return out
