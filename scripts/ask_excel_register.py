"""Ask Excel, not openpyxl, whether a scaffolded register behaves (REGUX-01 §6a).

Why this exists. ``insertRows`` was once set correctly by the scaffolder and Excel ignored it
entirely. A test proves an attribute was written; only Excel proves it means anything. This script
opens a throwaway register in a separate, hidden Excel instance and does what a person would do -
types into cells, sorts through the Table, asks the ribbon whether a button is enabled - and reports
Excel's answer. It never saves.

What it can and cannot establish (08/09/2026, found by control, not reasoning):

* A write to a locked cell on a protected sheet raises through COM exactly as it refuses in the
  UI. Typing evidence here is real evidence.
* ``ListObject.Sort.Apply`` goes through ``AllowSorting``; ``CommandBars.GetEnabledMso`` reports
  whether the ribbon's own Sort / Filter buttons are enabled for the selected cell. Both discriminate
  correctly against a control (enabled on the Risk Register, disabled on Settings).
* ``Range.AutoFilter(criteria)`` through COM is REFUSED on every protected sheet - including the
  Risk Register, whose filter the Director used by hand at 3c C-2. So that call is not a user's
  dropdown click and is not probed. **Clicking a filter dropdown is the one residual only a person
  can close.** The script says so in its output rather than pretending.

Run from the repository root, with Excel installed::

    python scripts/ask_excel_register.py                 # builds outputs/regux01_verify/ask_excel.xlsx
    python scripts/ask_excel_register.py path/to/copy.xlsx
    python scripts/ask_excel_register.py --dropdown     # also drive the filter dropdown via the UI

Exit status is 1 on any FAIL. ``outputs/`` is git-ignored; the throwaway is left for inspection.
"""
from __future__ import annotations

import sys
import time
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402
from tools.register_scaffold import build_template  # noqa: E402

RESULTS: list[tuple[str, str, str, str, str]] = []


def rec(sheet: str, probe: str, expected, actual, ok: bool) -> None:
    RESULTS.append((sheet, probe, str(expected), str(actual), "PASS" if ok else "FAIL"))


def try_write(sh, addr: str, value) -> tuple[bool, str]:
    try:
        sh.range(addr).value = value
        return True, f"wrote {value!r}"
    except Exception as e:  # noqa: BLE001 - pywintypes.com_error is what Excel says no with
        return False, "refused: " + str(e).replace("\n", " ")[60:150]


def probe_table_sheet(app, sh, name: str, last_col: str) -> None:
    """One list-shaped input sheet: a Table over the body, sort, the fill boundary."""
    api = sh.api
    n_lo = api.ListObjects.Count
    rec(name, "ListObjects.Count", 1, n_lo, n_lo == 1)
    if n_lo != 1:
        return
    lo = api.ListObjects(1)
    hdr = lo.HeaderRowRange.Row                 # part 3: a note block may sit above the header
    last_row = hdr + lo.Range.Rows.Count - 1
    ref = lo.Range.Address.replace("$", "")
    rec(name, "Table range", f"A{hdr}:{last_col}{last_row}", ref, ref == f"A{hdr}:{last_col}{last_row}")
    rec(name, "banded rows", True, lo.ShowTableStyleRowStripes, bool(lo.ShowTableStyleRowStripes))
    rec(name, "filter dropdowns shown", True, lo.ShowAutoFilter, bool(lo.ShowAutoFilter))
    p = api.Protection
    rec(name, "Protection.AllowFiltering", True, p.AllowFiltering, bool(p.AllowFiltering))
    rec(name, "Protection.AllowSorting", True, p.AllowSorting, bool(p.AllowSorting))
    hcell = sh.range(f"A{hdr}").api
    rec(name, "header bold", True, hcell.Font.Bold, bool(hcell.Font.Bold))
    sh.activate()
    win = app.api.ActiveWindow
    rec(name, "header row frozen", f"SplitRow {hdr}", f"SplitRow {win.SplitRow}",
        bool(win.FreezePanes) and win.SplitRow == hdr)
    rec(name, "UsedRange", f"A1:{last_col}{last_row}", api.UsedRange.Address.replace("$", ""), True)
    if hdr > 1:
        ok, msg = try_write(sh, "A1", "x")
        rec(name, "type A1 (the note above the header)", "refuse", msg, not ok)
        ok, msg = try_write(sh, f"A{hdr}", "x")
        rec(name, f"type A{hdr} (the header)", "refuse", msg, not ok)

    # the ribbon's own answer, for the selected body cell
    sh.range(f"A{hdr + 1}").select()
    cb = app.api.CommandBars
    rec(name, "ribbon Sort A-Z enabled", True, cb.GetEnabledMso("SortAscendingExcel"),
        bool(cb.GetEnabledMso("SortAscendingExcel")))
    allow_ins = bool(p.AllowInsertingRows)   # A-0.6: the expectation follows the sheet's own flag
    rec(name, f"ribbon Insert Sheet Rows enabled (flag {allow_ins})", allow_ins,
        cb.GetEnabledMso("SheetRowsInsert"), bool(cb.GetEnabledMso("SheetRowsInsert")) == allow_ins)

    # typing: three IDs out of order, the last column, the last Table row, then past both edges
    first = hdr + 1
    for row, v in zip(range(first, first + 3), ["RSK-003", "RSK-001", "RSK-002"]):
        ok, msg = try_write(sh, f"A{row}", v)
        rec(name, f"type A{row}", "accept", msg, ok)
    ok, msg = try_write(sh, f"{last_col}{first}", "x")
    rec(name, f"type {last_col}{first} (last column)", "accept", msg, ok)
    ok, msg = try_write(sh, f"A{last_row}", "RSK-END")
    rec(name, f"type A{last_row} (last Table row)", "accept", msg, ok)
    ok, msg = try_write(sh, f"A{last_row + 1}", "RSK-PAST")
    rec(name, f"type A{last_row + 1} (below Table)", "refuse", msg, not ok)
    nxt = get_column_letter(column_index_from_string(last_col) + 1)
    ok, msg = try_write(sh, f"{nxt}{first}", "x")
    rec(name, f"type {nxt}{first} (past last column)", "refuse", msg, not ok)

    # sort through the Table's own sort - the path AllowSorting governs
    try:
        lo.Sort.SortFields.Clear()
        lo.Sort.SortFields.Add(Key=lo.ListColumns(1).Range, SortOn=0, Order=1)
        lo.Sort.Header = 1
        lo.Sort.Apply()
        got = [sh.range(f"A{r}").value for r in range(first, first + 3)]
        rec(name, "sort by column A", ["RSK-001", "RSK-002", "RSK-003"], got,
            got == ["RSK-001", "RSK-002", "RSK-003"])
    except Exception as e:  # noqa: BLE001
        rec(name, "sort by column A", "sorts", "refused: " + str(e)[60:150], False)

    # A-0.6: a whole-row insert inside the Table. Expected to succeed iff the flag is granted; when
    # it does, the Table must have grown by one row and the new row must accept typing.
    try:
        sh.range(f"A{first + 1}").api.EntireRow.Insert()
        grown = hdr + lo.Range.Rows.Count - 1 == last_row + 1
        ok2, msg2 = try_write(sh, f"A{first + 1}", "RSK-NEW")
        rec(name, f"insert row (flag {allow_ins})", "accept" if allow_ins else "refuse",
            f"inserted; Table rows {lo.Range.Rows.Count}; type A{first + 1} -> {msg2}",
            allow_ins and grown and ok2)
    except Exception as e:  # noqa: BLE001
        rec(name, f"insert row (flag {allow_ins})", "accept" if allow_ins else "refuse",
            "refused: " + str(e)[60:150], not allow_ins)


#: REGUX-01 A-0.4: the Layer-1 grids at their real shape. (sheet, cell, expected, value)
LAYER1_PROBES = [
    ("Layer1_Base", "A2", "accept", "Civils"),          # rename a placeholder WBS row
    ("Layer1_Base", "B2", "accept", 10.5),              # a $M value
    ("Layer1_Base", "I13", "accept", 1.0),              # last cell of the grid
    ("Layer1_Base", "B1", "accept", "Labour hire"),     # a resource column may be renamed
    ("Layer1_Base", "A1", "refuse", "x"),               # the corner label
    ("Layer1_Base", "B15", "accept", 25.0),             # Declared_Total value
    ("Layer1_Base", "A15", "refuse", "x"),              # its label
    ("Layer1_Base", "B16", "refuse", "x"),              # the Units text
    ("Layer1_Base", "J2", "refuse", 1.0),               # past the last column
    ("Layer1_Base", "B14", "refuse", 1.0),              # the blank spacer row
    ("Layer1_Allocation", "B1", "accept", "Civils"),    # the WBS header follows Layer1_Base
    ("Layer1_Allocation", "A1", "refuse", "x"),
    ("Layer1_Allocation", "A2", "accept", "Labour rate"),
    ("Layer1_Allocation", "B2", "accept", "x"),
    ("Layer1_Allocation", "M26", "accept", "x"),        # last cell
    ("Layer1_Allocation", "N2", "refuse", "x"),
    ("Layer1_Allocation", "A27", "refuse", "x"),
    ("Layer1_Correlation", "B1", "accept", "zero"),     # the mode toggle, unchanged
    ("Layer1_Correlation", "B2", "refuse", "x"),        # the PSD row the reader recomputes
    ("Layer1_Correlation", "B4", "accept", "Labour rate"),   # a name across the header
    ("Layer1_Correlation", "A5", "accept", "Labour rate"),   # and down column A
    ("Layer1_Correlation", "B6", "accept", 0.2),        # strictly below the diagonal
    ("Layer1_Correlation", "Y29", "accept", 0.1),       # the last lower-triangle cell
    ("Layer1_Correlation", "B5", "refuse", 1.0),        # the diagonal
    ("Layer1_Correlation", "Z29", "refuse", 1.0),       # the last diagonal cell
    ("Layer1_Correlation", "C5", "refuse", 0.3),        # above the diagonal - never read
    ("Layer1_Correlation", "A4", "refuse", "x"),        # the 'Factor' header label
    ("Layer1_Correlation", "A30", "refuse", "x"),       # below the triangle
]

#: Cells whose behaviour must not move, whatever part is under test. (sheet, cell, expected, value)
CONTROLS = [
    ("Risk Register", "A5", "accept", "RSK-001"),
    ("Risk Register", "AB5", "refuse", 1),          # EMV formula column
    ("MC_Params", "C11", "accept", 100000),         # the day rate
    ("MC_Params", "B11", "refuse", "x"),            # its label
    ("Settings", "B6", "refuse", "x"),              # contract version
    ("TOC & Instructions", "B61", "accept", "x"),   # project name
    ("TOC & Instructions", "A27", "refuse", "x"),   # legend swatch, knowingly locked
    ("Simulation", "D2", "accept", 1),              # write-back, unlocked by design (open question)
]


def drive_dropdown(app, book, sheet_names) -> None:
    """The residual that COM cannot reach: open a Table header's filter dropdown the way a person
    does (Alt+Down), jump to its search box ('e'), type a value, CLICK OK (Enter does not reach
    the popup), read back which rows Excel hid, then clear (Alt+Down, 'c'). Opt-in via
    ``--dropdown`` because it makes Excel visible, maximised and foreground; every key and click
    is sent only while the foreground window belongs to this Excel process, and the run aborts
    otherwise rather than typing into whatever is in front. Closed the 2b / A-0.4 residual on
    08/09/2026 on all five Table sheets."""
    import ctypes
    import win32api, win32con, win32gui, win32process  # pywin32
    ctypes.windll.user32.SetProcessDPIAware()
    app.visible = True
    app.api.WindowState = -4137          # xlMaximized: the popup's OK lands where OK_XY says
    OK_XY = (275, 1273)                  # physical pixels, 3840x2160; re-measure on another screen

    def foreground_is_ours():
        h = win32gui.GetForegroundWindow()
        _, pid = win32process.GetWindowThreadProcessId(h)
        return pid == app.pid, win32gui.GetWindowText(h)

    def click(x, y):
        win32api.SetCursorPos((x, y)); time.sleep(0.15)
        win32api.mouse_event(win32con.MOUSEEVENTF_LEFTDOWN, 0, 0, 0, 0); time.sleep(0.08)
        win32api.mouse_event(win32con.MOUSEEVENTF_LEFTUP, 0, 0, 0, 0)

    time.sleep(2.0)
    for name in sheet_names:
        sh = book.sheets[name]
        sh.activate()
        lo = sh.api.ListObjects(1)
        hdr = lo.HeaderRowRange.Row
        for r, v in zip(range(hdr + 1, hdr + 4), ["RSK-003", "RSK-001", "RSK-002"]):
            sh.range(f"A{r}").value = v
        sh.range(f"A{hdr}").select(); app.activate(steal_focus=True); time.sleep(1.5)
        ok, title = foreground_is_ours()
        if not ok:
            rec(name, "dropdown filter (UI)", "filters", f"ABORTED - foreground is {title!r}", False)
            return
        app.api.SendKeys("%{DOWN}", True); time.sleep(1.0)
        app.api.SendKeys("e", True); time.sleep(0.4)
        app.api.SendKeys("RSK-002", True); time.sleep(0.8)
        ok, title = foreground_is_ours()
        if not ok:
            rec(name, "dropdown filter (UI)", "filters", f"ABORTED before OK - foreground is {title!r}", False)
            return
        click(*OK_XY); time.sleep(1.0)
        rows = range(hdr + 1, hdr + 4)
        visible = [sh.range(f"A{r}").value for r in rows if not sh.range(f"A{r}").api.EntireRow.Hidden]
        rec(name, "dropdown filter (UI): rows left visible", ["RSK-002"], visible, visible == ["RSK-002"])
        sh.range(f"A{hdr}").select(); app.activate(steal_focus=True); time.sleep(0.5)
        app.api.SendKeys("%{DOWN}", True); time.sleep(1.0)
        app.api.SendKeys("c", True); time.sleep(1.0)
        cleared = not any(sh.range(f"A{r}").api.EntireRow.Hidden for r in rows)
        rec(name, "dropdown clear (UI)", "all rows visible", f"cleared={cleared}", cleared)


def main(argv: list[str]) -> int:
    import xlwings as xw

    dropdown = "--dropdown" in argv
    argv = [a for a in argv if a != "--dropdown"]
    if argv:
        path = Path(argv[0]).resolve()
    else:
        path = REPO / "outputs" / "regux01_verify" / "ask_excel.xlsx"
        if path.exists():
            path.unlink()
        build_template(path)
    if path == (REPO / "templates" / "register_template.xlsx").resolve():
        raise SystemExit("refusing to type into the tracked template; pass a copy")
    print(f"asking Excel about {path.relative_to(REPO) if path.is_relative_to(REPO) else path}")

    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    app.screen_updating = False
    try:
        book = app.books.open(str(path), update_links=False)
        # every sheet carrying a Table is a list-shaped input sheet (§5.4: five of them)
        for sh in book.sheets:
            if sh.api.ListObjects.Count:
                n_cols = sh.api.ListObjects(1).Range.Columns.Count
                probe_table_sheet(app, sh, sh.name, get_column_letter(n_cols))
        for sheet, addr, exp, v in LAYER1_PROBES:
            ok, msg = try_write(book.sheets[sheet], addr, v)
            rec(sheet, f"type {addr}", exp, msg, ok if exp == "accept" else not ok)
        for sheet, addr, exp, v in CONTROLS:
            ok, msg = try_write(book.sheets[sheet], addr, v)
            rec(sheet, f"control: type {addr}", exp, msg, ok if exp == "accept" else not ok)
        rr = book.sheets["Risk Register"]
        rr.activate()
        rr.range("A5").select()
        rec("Risk Register", "control: ribbon Sort A-Z enabled", True,
            app.api.CommandBars.GetEnabledMso("SortAscendingExcel"),
            bool(app.api.CommandBars.GetEnabledMso("SortAscendingExcel")))
        book.sheets["Settings"].activate()
        book.sheets["Settings"].range("A2").select()
        rec("Settings", "control: ribbon Sort A-Z DISABLED (not permitted)", False,
            app.api.CommandBars.GetEnabledMso("SortAscendingExcel"),
            not app.api.CommandBars.GetEnabledMso("SortAscendingExcel"))
        rec("workbook", "structure protected", True, book.api.ProtectStructure,
            bool(book.api.ProtectStructure))
        if dropdown:
            drive_dropdown(app, book, [sh.name for sh in book.sheets if sh.api.ListObjects.Count])
        book.close()
    finally:
        app.quit()

    w = max(len(r[1]) for r in RESULTS)
    print(f"\n{'sheet':26} {'probe':{w}}  {'expected':32} {'actual':44} verdict")
    for s, p, e, a, v in RESULTS:
        print(f"{s:26} {p:{w}}  {e[:32]:32} {a[:44]:44} {v}")
    fails = [r for r in RESULTS if r[4] == "FAIL"]
    print(f"\n{len(RESULTS)} probes, {len(RESULTS) - len(fails)} PASS, {len(fails)} FAIL")
    for s, p, e, a, _ in fails:
        print(f"  FAIL  {s} | {p} | expected {e} | got {a}")
    if not dropdown:
        print("\nNOT PROBED - the filter dropdown click. COM cannot do it on a protected sheet; pass "
              "--dropdown to drive it through the UI (Excel becomes visible and foreground).")
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
