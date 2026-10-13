"""Build a period status workbook from any register.

The period cycle's input is a four-sheet status workbook whose layout was stated only in
``src/am/workbook_reader.py`` and in two Builder scripts (the wind fixture's, and the Kestrel one
written during stage 4 because the analyst could not start step 7). This is the user-facing
generator: it reads a register by its header labels, takes every risk that is Open with
``Time Impact = Yes`` - the same rule as the register's ``AM_Include`` helper formula - and
writes a baseline period the engine accepts. It also writes the empty template that ships as
``templates/status_workbook_template.xlsx``.

Usage (see docs/PLAYBOOK.md Part A step 7 and Part B):

    python scripts/build_status_workbook.py --register outputs/<project>/register_v1.xlsx ^
        --period-end 2026-08-31 --out outputs/<project>/period_01.xlsx [--phase TEXT]
        [--previous-period YYYY-MM-DD] [--analyst TEXT] [--reviewer TEXT] [--notes TEXT]
        [--milestone "ID|Name|YYYY-MM-DD" ...]
    python scripts/build_status_workbook.py --template templates/status_workbook_template.xlsx

The reader contract is imported from ``scripts/build_wind_am_periods.py`` (its one statement in
``scripts/``) and is not restated here. The embedded ``Risk Register`` helper sheet is written
because a scaffolded register carries no ``AM_Include`` column, so the reader falls back to it.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from scripts.build_wind_am_periods import (  # noqa: E402 - the one statement of the contract
    AM_INCLUDE_COL, AM_INCLUDE_FIRST_ROW, AM_MILESTONES_SHEET, AM_MITIGATION_SHEET,
    AM_STATUS_FIRST_DATA_ROW, AM_STATUS_HEADER_ROW, AM_STATUS_SHEET, MILESTONE_COLUMNS,
    MITIGATION_COLUMNS, PERIOD_METADATA_RANGES, RISK_REGISTER_SHEET, RISK_STATUS_COLUMNS,
)
from src.am.workbook_reader import (  # noqa: E402 - the reader scans a bounded row range
    AM_MILESTONES_HEADER_ROW, AM_MILESTONES_MAX_DATA_ROW, AM_PERIOD_STATUS_MAX_DATA_ROW,
)
from src.qra_register_parse import DATA_START_ROW, HEADER_ROW, find_columns  # noqa: E402

#: The reader reads rows 15..30 of AM_Period_Status and 5..30 of AM_Milestones and no further, so
#: a period with more rows than that would be read short with no error. Refused here by name.
MAX_TRACKED_RISKS = AM_PERIOD_STATUS_MAX_DATA_ROW - AM_STATUS_FIRST_DATA_ROW + 1
MAX_MILESTONES = AM_MILESTONES_MAX_DATA_ROW - AM_MILESTONES_HEADER_ROW

#: Register header labels this builder needs, keyed as the parser keys them. ``risk_owner`` is
#: optional: the parser reads the owner from ``Package`` first and falls back to ``Risk Owner``
#: too, and this builder does the same so the period shows the owner the deliverable
#: shows.
_REQUIRED_HEADERS = {
    "risk_id": ["Risk ID"],
    "risk_status": ["Risk Status"],
    "description": ["Risk Description"],
    "time_impact": ["Time Impact"],
    # A tracked risk's one probability.
    "probability": ["Probability"],
    "mode": ["Mode"],
}
_OPTIONAL_HEADERS = {"package": ["Package"], "risk_owner": ["Risk Owner"]}

_METADATA_LABELS = (
    "Period end", "Previous period", "Project phase", "Engine run date",
    "Engine run hash", "Analyst", "Reviewer", "Period notes",
)


@dataclass(frozen=True)
class TrackedRisk:
    risk_id: str
    title: str
    probability: float
    mode: str
    owner: str


@dataclass(frozen=True)
class Milestone:
    milestone_id: str
    name: str
    planned_finish: _dt.date


def tracked_risks(register: Path) -> list[TrackedRisk]:
    """Every Open, ``Time Impact = Yes`` risk in the register, in row order.

    Refuses by name when a header the rule needs is absent, and when the rule selects nothing:
    a period that tracks no risk is not a period.
    """
    wb = openpyxl.load_workbook(register, data_only=True, read_only=True)
    try:
        grid = [tuple(row) for row in wb[RISK_REGISTER_SHEET].iter_rows(values_only=True)]
    finally:
        wb.close()
    col = find_columns(grid, HEADER_ROW, {**_REQUIRED_HEADERS, **_OPTIONAL_HEADERS})
    # Exact, never prefix: "Probability" is a prefix of "Probability (Time)", so a register
    # missing the one would otherwise be read from the other without a word.
    headers = grid[HEADER_ROW - 1] if 0 < HEADER_ROW <= len(grid) else ()
    idx = col.get("probability")
    if idx is not None and str(headers[idx] or "").strip().lower() != "probability":
        col["probability"] = None
    missing = [_REQUIRED_HEADERS[k][0] for k in _REQUIRED_HEADERS if col.get(k) is None]
    if missing:
        raise SystemExit(
            f"{register}: 'Risk Register' row {HEADER_ROW} has no header {missing}; "
            "the register is not one this product scaffolds"
        )

    def cell(row: tuple, key: str):
        idx = col.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    out: list[TrackedRisk] = []
    for row in grid[DATA_START_ROW - 1:]:
        rid = cell(row, "risk_id")
        if rid is None or str(rid).strip() == "":
            continue
        status = str(cell(row, "risk_status") or "").strip().lower()
        time_yes = str(cell(row, "time_impact") or "").strip().lower() == "yes"
        if status != "open" or not time_yes:
            continue
        owner = str(cell(row, "package") or "").strip() or str(cell(row, "risk_owner") or "").strip()
        prob = cell(row, "probability")
        if prob is None or str(prob).strip() == "":
            # a blank is not zero: a tracked risk that never fires is a wrong period number
            raise SystemExit(
                f"{register}: risk {str(rid).strip()} is Open with Time Impact = Yes but its "
                f"'Probability' is blank"
            )
        out.append(TrackedRisk(
            risk_id=str(rid).strip(),
            title=str(cell(row, "description") or "").strip(),
            probability=float(prob),
            mode=str(cell(row, "mode") or "").strip(),
            owner=owner,
        ))
    if not out:
        raise SystemExit(
            f"{register}: no Open, Time Impact = Yes risks - nothing for a period to track"
        )
    if len(out) > MAX_TRACKED_RISKS:
        raise SystemExit(
            f"{register}: {len(out)} Open, Time Impact = Yes risks, but the period reader reads at "
            f"most {MAX_TRACKED_RISKS} (AM_Period_Status rows {AM_STATUS_FIRST_DATA_ROW}-"
            f"{AM_PERIOD_STATUS_MAX_DATA_ROW}); a period would be read short without error"
        )
    return out


def parse_milestone(spec: str) -> Milestone:
    """``ID|Name|YYYY-MM-DD``."""
    parts = [p.strip() for p in spec.split("|")]
    if len(parts) != 3:
        raise SystemExit(f"--milestone {spec!r}: expected ID|Name|YYYY-MM-DD")
    return Milestone(parts[0], parts[1], _dt.date.fromisoformat(parts[2]))


def _write_frame(wb: Workbook) -> tuple:
    """The four sheets, headers and named ranges - the contract, with no period in it."""
    rr = wb.active
    rr.title = RISK_REGISTER_SHEET
    rr.cell(row=4, column=1, value="Risk ID")
    rr.cell(row=4, column=3, value="Risk Description")
    rr.cell(row=4, column=AM_INCLUDE_COL, value="AM_Include")

    ps = wb.create_sheet(AM_STATUS_SHEET)
    for i, (label, name) in enumerate(zip(_METADATA_LABELS, PERIOD_METADATA_RANGES), start=1):
        ps.cell(row=i, column=1, value=label)
        wb.defined_names[name] = DefinedName(name, attr_text=f"{AM_STATUS_SHEET}!$B${i}")
    for c, colname in enumerate(RISK_STATUS_COLUMNS, start=1):
        ps.cell(row=AM_STATUS_HEADER_ROW, column=c, value=colname)

    ms = wb.create_sheet(AM_MILESTONES_SHEET)
    for c, colname in enumerate(MILESTONE_COLUMNS, start=1):
        ms.cell(row=4, column=c, value=colname)

    mt = wb.create_sheet(AM_MITIGATION_SHEET)
    for c, colname in enumerate(MITIGATION_COLUMNS, start=1):
        mt.cell(row=4, column=c, value=colname)
    return rr, ps, ms, mt


def build_template(out: Path) -> Path:
    """The empty status workbook: headers and named ranges, no rows."""
    wb = Workbook()
    _write_frame(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_status_workbook(
    register: Path, out: Path, *, period_end: _dt.date, previous_period: _dt.date | None = None,
    phase: str = "", analyst: str = "", reviewer: str = "", notes: str = "",
    milestones: tuple[Milestone, ...] = (),
) -> Path:
    """A baseline period: every tracked risk Open at its register probability, no overrides."""
    risks = tracked_risks(register)
    if len(milestones) > MAX_MILESTONES:
        raise SystemExit(
            f"{len(milestones)} milestones, but the period reader reads at most {MAX_MILESTONES} "
            f"(AM_Milestones rows {AM_MILESTONES_HEADER_ROW + 1}-{AM_MILESTONES_MAX_DATA_ROW})"
        )
    wb = Workbook()
    rr, ps, ms, _mt = _write_frame(wb)

    for i, r in enumerate(risks):
        row = AM_INCLUDE_FIRST_ROW + i
        rr.cell(row=row, column=1, value=r.risk_id)
        rr.cell(row=row, column=3, value=r.title)
        rr.cell(row=row, column=AM_INCLUDE_COL, value=r.risk_id)

    ps["B1"] = period_end
    ps["B2"] = previous_period
    ps["B3"] = phase
    ps["B6"] = analyst
    ps["B7"] = reviewer
    ps["B8"] = notes
    for i, r in enumerate(risks):
        row = AM_STATUS_FIRST_DATA_ROW + i
        vals = {
            "Risk_ID": r.risk_id, "Risk_title": r.title, "Status_this_period": "Open",
            "Status_changed": None, "Probability_current": r.probability,
            "Probability_override": None, "Owner": r.owner, "Mode_current": r.mode,
            "Mode_override": None, "Fired_date": None, "Closed_date": None,
            "Notes_this_period": None,
        }
        for c, colname in enumerate(RISK_STATUS_COLUMNS, start=1):
            ps.cell(row=row, column=c, value=vals[colname])

    for i, m in enumerate(milestones):
        vals = {
            "Milestone_ID": m.milestone_id, "Milestone_name": m.name, "WBS_code": None,
            "Planned_start": None, "Planned_finish": m.planned_finish, "Status": "Not started",
            "Percent_complete": 0.0, "Actual_start": None, "Actual_finish": None, "Notes": None,
        }
        for c, colname in enumerate(MILESTONE_COLUMNS, start=1):
            ms.cell(row=5 + i, column=c, value=vals[colname])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_parser() -> argparse.ArgumentParser:
    """The command's arguments, built here so a test can read the help a user is given."""
    parser = argparse.ArgumentParser(
        prog="build_status_workbook",
        description="Build a period status workbook from a register (STATUSWB-01).",
    )
    parser.add_argument("--register", type=Path, help="the register .xlsx the period tracks")
    parser.add_argument("--period-end", type=_dt.date.fromisoformat, help="YYYY-MM-DD")
    parser.add_argument("--out", type=Path, help="where to write the status workbook")
    parser.add_argument("--previous-period", type=_dt.date.fromisoformat, default=None,
                        help="the previous period end, YYYY-MM-DD, for the trend and delta blocks")
    parser.add_argument("--phase", default="",
                        help="free text naming the phase the project is in, printed on the "
                             "period's metadata block for a reader (e.g. Construction). "
                             "Optional, and nothing validates it")
    parser.add_argument("--analyst", default="", help="who filled this period, for the record")
    parser.add_argument("--reviewer", default="", help="who reviewed it, for the record")
    parser.add_argument("--notes", default="", help="free text carried onto the period's notes cell")
    parser.add_argument("--milestone", action="append", default=[], metavar="ID|Name|YYYY-MM-DD",
                        help="a milestone row, repeatable")
    parser.add_argument("--template", type=Path, default=None,
                        help="write the empty template here instead (no register needed)")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    if args.template is not None:
        path = build_template(args.template)
        print(f"wrote {path}  (empty status workbook template)")
        return 0
    missing = [n for n, v in (("--register", args.register), ("--period-end", args.period_end),
                              ("--out", args.out)) if v is None]
    if missing:
        parser.error(f"required: {' '.join(missing)} (or --template)")
    path = build_status_workbook(
        args.register, args.out, period_end=args.period_end,
        previous_period=args.previous_period, phase=args.phase, analyst=args.analyst,
        reviewer=args.reviewer, notes=args.notes,
        milestones=tuple(parse_milestone(m) for m in args.milestone),
    )
    n = len(tracked_risks(args.register))
    print(f"wrote {path}  ({n} tracked risks, period end {args.period_end})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
