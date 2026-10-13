"""Stage 4 step 7 - a baseline status-period workbook for the Kestrel Highway Upgrade, authored
from the analyst's hand-filled register.

**Why this exists, and what it is evidence of.** Step 7 (period JSON, then the app) takes a status
workbook - the Active Management cycle's own input - whose layout is stated only in
``src/am/workbook_reader.py`` and the wind-only ``scripts/build_wind_am_periods.py`` (stage 4
finding F-31). An analyst on a third project could not produce one from anything the product
shipped - since STATUSWB-01 (stage 5a) they can: ``scripts/build_status_workbook.py`` is the generic
builder, and this script stays as stage 4's evidence.
The Director ruled 08/09/2026 that the Builder author one so the period engine and the app are
measured on a third project anyway; this script is that input, and the record labels it
**Builder-supplied**, in the same way as the workshop handout.

**What it reads.** The register the analyst filled (``outputs/kestrel_highway/register_v1.xlsx``):
the ten Time=Yes risks - ID, description, Probability (Time), Mode, and the column the parser reads
as owner - become the AM in-scope set, exactly as the wind cycle tracks its schedule-driving risks.
Milestones are the planner's six work groups from the workshop output. One period, the baseline: every
risk Open at its register probability, no overrides, no mitigations.

**Reader contract, copied from the wind script and not re-derived**: ``Risk Register`` with the
``AM_Include`` helper at column BQ from row 5; ``AM_Period_Status`` with eight ``am_*`` named ranges
at B1:B8 and the twelve-column header at row 14, rows from 15; ``AM_Milestones`` (ten columns, header
row 4); ``AM_Mitigation_Log`` (twelve columns, header row 4).

Run:  python scripts/build_kestrel_am_period.py [register.xlsx] [out.xlsx]
"""
from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from scripts.build_wind_am_periods import (  # noqa: E402 - the one statement of the contract
    AM_INCLUDE_COL, AM_INCLUDE_FIRST_ROW, AM_MILESTONES_SHEET, AM_MITIGATION_SHEET,
    AM_STATUS_FIRST_DATA_ROW, AM_STATUS_HEADER_ROW, AM_STATUS_SHEET, MILESTONE_COLUMNS,
    MITIGATION_COLUMNS, PERIOD_METADATA_RANGES, RISK_REGISTER_SHEET, RISK_STATUS_COLUMNS,
)

DEFAULT_REGISTER = REPO_ROOT / "outputs" / "kestrel_highway" / "register_v1.xlsx"
DEFAULT_OUT = REPO_ROOT / "outputs" / "kestrel_highway" / "kestrel_period_01.xlsx"

#: The planner's six work groups (workshop output §2), as milestones for display.
MILESTONES = [
    ("KM-01", "Site establishment & traffic staging complete", _dt.date(2026, 9, 30)),
    ("KM-02", "Earthworks & drainage complete", _dt.date(2027, 1, 31)),
    ("KM-03", "Viaduct complete", _dt.date(2027, 8, 31)),
    ("KM-04", "Bridge complete", _dt.date(2027, 11, 30)),
    ("KM-05", "Pavement & surfacing complete", _dt.date(2028, 4, 30)),
    ("KM-06", "Finishing & handover", _dt.date(2028, 6, 30)),
]

PERIOD = {
    "end": _dt.date(2026, 8, 31),
    "prev": None,
    "phase": "Establishment & Earthworks",
    "notes": "Baseline period, stage 4 step 7. All delaying risks Open at register probability; "
             "no overrides, no mitigations. Status workbook authored by the Builder (F-31).",
}


def tracked_risks(register: Path) -> list[tuple[str, str, float, str, str]]:
    """(id, title, P(time), Mode, owner) for every Time=Yes risk in the register, in row order."""
    ws = openpyxl.load_workbook(register, data_only=True)["Risk Register"]
    out = []
    for r in range(5, ws.max_row + 1):
        rid = ws.cell(row=r, column=1).value
        if rid is None or ws.cell(row=r, column=20).value != "Yes":
            continue
        out.append((
            str(rid).strip(),
            str(ws.cell(row=r, column=3).value or "").strip(),
            float(ws.cell(row=r, column=21).value or 0.0),
            str(ws.cell(row=r, column=26).value or "").strip(),
            str(ws.cell(row=r, column=4).value or "").strip(),   # the column the parser reads as owner
        ))
    if not out:
        raise SystemExit(f"{register}: no Time=Yes risks - nothing for the period to track")
    return out


def build(register: Path, out: Path) -> Path:
    risks = tracked_risks(register)
    wb = Workbook()

    rr = wb.active
    rr.title = RISK_REGISTER_SHEET
    rr.cell(row=4, column=1, value="Risk ID")
    rr.cell(row=4, column=3, value="Risk Description")
    rr.cell(row=4, column=AM_INCLUDE_COL, value="AM_Include")
    for i, (rid, title, *_rest) in enumerate(risks):
        r = AM_INCLUDE_FIRST_ROW + i
        rr.cell(row=r, column=1, value=rid)
        rr.cell(row=r, column=3, value=title)
        rr.cell(row=r, column=AM_INCLUDE_COL, value=rid)

    ps = wb.create_sheet(AM_STATUS_SHEET)
    labels = ["Period end", "Previous period", "Project phase", "Engine run date",
              "Engine run hash", "Analyst", "Reviewer", "Period notes"]
    for i, (label, name) in enumerate(zip(labels, PERIOD_METADATA_RANGES), start=1):
        ps.cell(row=i, column=1, value=label)
        wb.defined_names[name] = DefinedName(name, attr_text=f"{AM_STATUS_SHEET}!$B${i}")
    ps["B1"] = PERIOD["end"]
    ps["B2"] = PERIOD["prev"]
    ps["B3"] = PERIOD["phase"]
    ps["B6"] = "Kestrel Highway PM"
    ps["B7"] = "Kestrel Highway Reviewer"
    ps["B8"] = PERIOD["notes"]
    for c, col in enumerate(RISK_STATUS_COLUMNS, start=1):
        ps.cell(row=AM_STATUS_HEADER_ROW, column=c, value=col)
    for i, (rid, title, prob, mode, owner) in enumerate(risks):
        row = AM_STATUS_FIRST_DATA_ROW + i
        vals = {
            "Risk_ID": rid, "Risk_title": title, "Status_this_period": "Open",
            "Status_changed": None, "Probability_current": prob, "Probability_override": None,
            "Owner": owner, "Mode_current": mode, "Mode_override": None,
            "Fired_date": None, "Closed_date": None, "Notes_this_period": None,
        }
        for c, col in enumerate(RISK_STATUS_COLUMNS, start=1):
            ps.cell(row=row, column=c, value=vals[col])

    ms = wb.create_sheet(AM_MILESTONES_SHEET)
    for c, col in enumerate(MILESTONE_COLUMNS, start=1):
        ms.cell(row=4, column=c, value=col)
    for i, (mid, mname, planned) in enumerate(MILESTONES):
        vals = {
            "Milestone_ID": mid, "Milestone_name": mname, "WBS_code": None,
            "Planned_start": None, "Planned_finish": planned,
            "Status": "In progress" if i == 0 else "Not started",
            "Percent_complete": 0.3 if i == 0 else 0.0,
            "Actual_start": None, "Actual_finish": None, "Notes": None,
        }
        for c, col in enumerate(MILESTONE_COLUMNS, start=1):
            ms.cell(row=5 + i, column=c, value=vals[col])

    mt = wb.create_sheet(AM_MITIGATION_SHEET)
    for c, col in enumerate(MITIGATION_COLUMNS, start=1):
        mt.cell(row=4, column=c, value=col)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main(argv: list[str]) -> int:
    register = Path(argv[0]) if argv else DEFAULT_REGISTER
    out = Path(argv[1]) if len(argv) > 1 else DEFAULT_OUT
    path = build(register, out)
    print(f"wrote {path}  ({len(tracked_risks(register))} tracked risks, period end {PERIOD['end']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
