"""Showcase-P3 — section 4 pre-paste verification (checks 2, 3, 4).

Proves the engine readers + Stage-3 validator absorb the *larger* P3
dataset shapes (30-40 risks, 150-300 activities, the full CG01-CG13
cost-group surface) with NO engine edit, before any builder change is
written. Constructs throwaway in-memory artefacts at the target scale and
runs them through the production readers / validators.

Run:  python scripts/p3_prepaste_verify.py
Exit 0 = all checks pass (safe to proceed). Non-zero = an engine edit
would be required → pause and report per brief §4.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
from src.sra.methodology_defaults_reader import build_test_methodology_defaults  # noqa: E402

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

CG_CODES = [f"CG{n:02d}" for n in range(1, 14)]  # CG01..CG13


def _write_scaled_register(path: Path, n_risks: int) -> None:
    """Emit a V5-format register with n_risks rows spanning CG01..CG13."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Register"
    ws.cell(row=3, column=14, value="Contingent Risk Cost")
    ws.cell(row=3, column=20, value="Time (Working Days)")
    headers = [
        "Risk ID", "Risk Status", "Risk Description", "Package",
        "RBS LVL 1", "RBS LVL 2", "MC Include", "Distribution",
        "Correlation Group", "Recurrence Type", "Causes", "Consequences",
        "Existing Controls",
        "Cost Impact", "Probability", "Optimistic ($)", "Most Likely ($)",
        "Pessimistic ($)", "Basis of Estimate",
        "Time Impact", "Probability", "Optimistic", "Most Likely",
        "Pessimistic", "Basis of Estimate",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    for i in range(n_risks):
        row = 5 + i
        cg = CG_CODES[i % len(CG_CODES)]
        time_yes = (i % 4 == 0)
        vals = [
            f"WT-{i + 1:03d}", "Open", f"Synthetic risk {i + 1}", "EPC Contractor",
            "Civil", "Foundations", "Yes", "PERT", cg, "",
            "cause", "consequence", "controls",
            "Yes", 0.4, 100_000, 300_000, 600_000, "Synthetic",
            "Yes" if time_yes else "No", 0.4 if time_yes else 0.0,
            3 if time_yes else 0, 8 if time_yes else 0, 18 if time_yes else 0,
            "Synthetic",
        ]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=row, column=c, value=v)
    mp = wb.create_sheet("MC_Params")
    params = [("Iterations (n)", 5000), ("Random seed (0=random)", 42),
              ("TIC - Total Installed Cost ($)", 50_000_000)]
    for cg in CG_CODES:
        params.append((f"{cg} - synthetic", 0.5))
    for i, (label, value) in enumerate(params, start=2):
        mp.cell(row=i, column=2, value=label)
        mp.cell(row=i, column=3, value=value)
    wb.save(str(path))


def _write_scaled_strawman(path: Path, n_risks: int) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk_Scope"
    ws.cell(row=1, column=1, value="Scaled synthetic strawman")
    headers = ["Risk_ID", "Risk_Title", "Risk_Owner", "RBS_L1", "Time_Impact",
               "Mode_Aggregated", "Series_Source", "Probability_Time",
               "Opt_Days", "ML_Days", "Pess_Days", "Pattern_ID", "Prefix_List",
               "Activity_Count", "Day_Rate"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c, value=h)
    n_yes = 0
    for i in range(n_risks):
        row = 6 + i
        time_yes = (i % 4 == 0)
        if time_yes:
            n_yes += 1
        ws.cell(row=row, column=1, value=f"WT-{i + 1:03d}")
        ws.cell(row=row, column=2, value=f"Synthetic risk {i + 1}")
        ws.cell(row=row, column=3, value="EPC Contractor")
        ws.cell(row=row, column=4, value="Civil")
        ws.cell(row=row, column=5, value="Yes" if time_yes else "No")
        ws.cell(row=row, column=6, value="Series" if time_yes else "")
        ws.cell(row=row, column=8, value=0.4 if time_yes else 0.0)
        ws.cell(row=row, column=9, value=3 if time_yes else 0)
        ws.cell(row=row, column=10, value=8 if time_yes else 0)
        ws.cell(row=row, column=11, value=18 if time_yes else 0)
        ws.cell(row=row, column=14, value=50)
        ws.cell(row=row, column=15, value=90_000.0)
    rc = wb.create_sheet("Run_Config_Preview")
    for i, (label, value) in enumerate(
        [("iterations", 5000), ("day_rate_default ($/wd)", 90_000.0),
         ("expected_time_yes_count", n_yes)], start=1):
        rc.cell(row=i, column=1, value=label)
        rc.cell(row=i, column=2, value=value)
    wb.save(str(path))
    return n_yes


def _build_scaled_xer_tables(n_activities: int) -> dict:
    calendar = [{
        "clndr_id": 1, "clndr_name": "Standard", "default_flag": "Y",
        "clndr_data": "(0||CalendarData())", "day_hr_cnt": 8.0,
        "week_hr_cnt": 40.0, "month_hr_cnt": 172.0, "year_hr_cnt": 2000.0,
    }]
    wbs_defs = [(110, "10", "Civil"), (120, "20", "Foundations"),
                (130, "30", "Turbine Erection"), (140, "40", "Electrical"),
                (150, "50", "Grid"), (160, "60", "Commissioning")]
    projwbs = [{"wbs_id": 100, "proj_id": 1, "parent_wbs_id": None,
                "wbs_short_name": "WF", "wbs_name": "Wind Farm", "seq_num": 1}]
    for k, (wid, short, name) in enumerate(wbs_defs, start=2):
        projwbs.append({"wbs_id": wid, "proj_id": 1, "parent_wbs_id": 100,
                        "wbs_short_name": short, "wbs_name": name, "seq_num": k})
    start = pd.Timestamp("2026-02-02 08:00")
    task, taskpred = [], []
    cursor = start
    prev_tid = None
    for i in range(n_activities):
        tid = 1000 + i
        wid, short, _ = wbs_defs[i % len(wbs_defs)]
        hrs = 40.0 + (i % 5) * 20.0
        end = cursor + pd.Timedelta(hours=hrs)
        task.append({
            "task_id": tid, "proj_id": 1, "wbs_id": wid, "clndr_id": 1,
            "task_code": f"{short}.A.{i + 1:03d}", "task_name": f"Activity {i + 1}",
            "task_type": "TT_Task", "target_drtn_hr_cnt": hrs,
            "total_float_hr_cnt": 0.0, "free_float_hr_cnt": 0.0,
            "early_start_date": cursor, "early_end_date": end,
            "late_start_date": cursor, "late_end_date": end,
            "target_start_date": cursor, "target_end_date": end,
            "cstr_date": None, "cstr_type": "", "driving_path_flag": "Y",
        })
        if prev_tid is not None:
            taskpred.append({"task_pred_id": 5000 + i, "task_id": tid,
                             "pred_task_id": prev_tid, "pred_type": "PR_FS",
                             "lag_hr_cnt": 0.0})
        prev_tid = tid
        cursor = end
    return {"CALENDAR": calendar, "PROJWBS": projwbs, "TASK": task, "TASKPRED": taskpred}


def _activate_wind_binding() -> None:
    """Register + select the synthetic-wind binding (prefix WT) for readers
    that resolve their risk-ID prefix from the active binding."""
    import importlib.util
    import os
    path = REPO_ROOT / "tests" / "fixtures" / "synthetic_wind" / "wind_binding.py"
    spec = importlib.util.spec_from_file_location("wind_binding_probe", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # self-registers 'synthetic_wind'
    os.environ["QRA_PROJECT"] = "synthetic_wind"


def check2_reader_shapes() -> bool:
    _activate_wind_binding()
    from src.v5_register_reader import read_risks_v5
    from src.sra.strawman_reader import read_risk_scope
    print("\n[Check 2] header-driven reads at 35-risk scale")
    with tempfile.TemporaryDirectory() as td:
        reg = Path(td) / "scaled_register.xlsx"
        straw = Path(td) / "scaled_strawman.xlsx"
        _write_scaled_register(reg, 35)
        n_yes = _write_scaled_strawman(straw, 35)
        risks = read_risks_v5(str(reg), risk_id_prefix="WT")
        scope = read_risk_scope(str(straw))
        ok_r = len(risks) == 35 and all(r.risk_id.startswith("WT-") for r in risks)
        cgs = {r.correlation_group for r in risks}
        ok_cg = cgs == set(CG_CODES)
        ok_s = len(scope) == 35 and (scope["time_impact"] == "Yes").sum() == n_yes
        print(f"  register rows read : {len(risks)} (expect 35)  -> {ok_r}")
        print(f"  distinct CG codes  : {len(cgs)} == 13 CG01..CG13 -> {ok_cg}")
        print(f"  strawman rows      : {len(scope)} (expect 35), Time=Yes {int((scope['time_impact']=='Yes').sum())}=={n_yes} -> {ok_s}")
        return ok_r and ok_cg and ok_s


def check3_xer_scale() -> bool:
    from src.sra.cpm.xer_writer import write_xer
    from src.sra.cpm.xer_reader import parse_xer, validate_xer_tables
    from src.sra.cpm.graph_builder import build_schedule_graph
    print("\n[Check 3] scaled XER (220 activities) validate + graph build + small MC")
    with tempfile.TemporaryDirectory() as td:
        xer = Path(td) / "scaled.xer"
        write_xer(_build_scaled_xer_tables(220), xer)
        tables = parse_xer(str(xer))
        validate_xer_tables(tables)  # raises on failure
        graph = build_schedule_graph(tables)
        n_act = len(getattr(graph, "activities", []) or getattr(graph, "nodes", []) or [])
        print(f"  validate_xer_tables: PASS")
        print(f"  graph activities   : {n_act} (expect ~220)")
        # quick MC to confirm a few-hundred-activity schedule runs without surprises
        from src.sra.cpm.cpm_simulation import ScheduleMCConfig, run_schedule_mc
        from src.sra.cpm.auto_mapper import expand_auto_rules
        from src.sra.cpm.risk_driver_schema import AutoMapperRule
        rules = {"WT-001": AutoMapperRule(
            risk_id="WT-001", pattern_id="P10", pattern_prefixes=("10",),
            float_threshold_hours=480.0, injection_method="additive_pert_days",
            opt_pert=3.0, ml_pert=8.0, pess_pert=18.0, probability=0.5, mode="Series")}
        drivers = expand_auto_rules(rules, graph)
        res = run_schedule_mc(graph=graph, risk_drivers=drivers,
                              config=ScheduleMCConfig(n_iter=2000, batch_size=1000, seed=42,
                                                      mode_multipliers=build_test_methodology_defaults().mode_multipliers))  # RS-D-3b: a verify script, the documented test factory
        import numpy as np
        finish = np.asarray(res.project_finish_hours)
        ok_mc = finish.shape[0] == 2000 and np.isfinite(finish).all()
        print(f"  2k-iter MC         : shape {finish.shape} finite -> {ok_mc}")
        return n_act >= 200 and ok_mc


def check4_cg_codes() -> bool:
    from src.v5_register_reader import _extract_cg_code
    print("\n[Check 4] CG-code extraction across CG01..CG13 (+ labelled form)")
    ok = True
    for cg in CG_CODES:
        if _extract_cg_code(cg) != cg:
            print(f"  FAIL bare {cg!r} -> {_extract_cg_code(cg)!r}"); ok = False
        labelled = f"{cg} - Wind cost group"
        if _extract_cg_code(labelled) != cg:
            print(f"  FAIL labelled {labelled!r} -> {_extract_cg_code(labelled)!r}"); ok = False
    print(f"  all 13 CG codes extract bare + labelled -> {ok}")
    return ok


def main() -> int:
    results = {
        "check2_reader_shapes": check2_reader_shapes(),
        "check3_xer_scale": check3_xer_scale(),
        "check4_cg_codes": check4_cg_codes(),
    }
    print("\n" + "=" * 60)
    for k, v in results.items():
        print(f"  {k:24s}: {'PASS' if v else 'FAIL'}")
    print("=" * 60)
    return 0 if all(results.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
